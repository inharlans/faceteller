"""Export saved, actual VGGT pilot outputs as a readable Excel sample."""
import hashlib
import json
import math
import sys
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / 'data/derived/pilot-20261014'
DEST = Path.home() / 'Downloads/FaceTeller_파일럿_공유자료'
NAMES = ['정면', '좌측사선', '우측사선']
FEATURES = {
    'nose_projection': ('코끝 돌출 비율', '코끝 − 양쪽 콧방울 기준점의 평균'),
    'nose_bridge': ('콧뿌리 돌출 비율', '콧뿌리 − 양쪽 눈 안쪽 기준점의 평균'),
    'lip_projection': ('입술 돌출 비율', '위·아래 입술 기준점의 평균 − 양쪽 입꼬리의 평균'),
    'chin_projection': ('턱 돌출 비율', '턱 기준점 − 코 아래 기준점'),
}


def main():
    DEST.mkdir(exist_ok=True, parents=True)
    target = DEST / 'FaceTeller_VGGT_실제출력_측정값샘플.xlsx'
    with np.load(SRC / 'private_depth_example.npz') as archive:
        arrays = {key: archive[key] for key in archive.files}
    sample = json.loads((SRC / 'smoke_518.jsonl').read_text(encoding='utf-8').splitlines()[0])
    config = json.loads((SRC / 'boundary_config.json').read_text(encoding='utf-8'))['features']
    book = Workbook()
    overview = book.active
    overview.title = '먼저보기'
    header_fill = PatternFill('solid', fgColor='17365D')
    numeric_checks = []

    def sheet(name, headers):
        ws = book.create_sheet(name)
        ws.append(headers)
        ws.freeze_panes = 'A2'
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = Font(name='맑은 고딕', color='FFFFFF', bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[1].height = 32
        return ws

    def pixel_values(view, x, y):
        return [x, y, float(arrays['depth'][view,y,x,0]),
                float(arrays['confidence'][view,y,x]),
                *map(float, arrays['points'][view,y,x])]

    overview.append(['VGGT 실제 출력값 샘플', '2026-10-14 파일럿에서 저장한 실제 결과'])
    notes = [
        ('대상', '실제 K-FACE 한 사람의 정면·좌측 사선·우측 사선 3장. 새 추론이나 가상 수치를 사용하지 않음.'),
        ('어디부터 볼까', '정면_픽셀값에서 x,y에 따른 깊이·신뢰도·X,Y,Z를 확인. 정면_깊이격자는 같은 깊이를 사진 배열처럼 배치.'),
        ('픽셀 좌표', '리사이즈·패딩 후 518×518 입력 기준. 좌상단 (0,0), x는 오른쪽, y는 아래쪽. 원본 사진 좌표가 아님.'),
        ('추출 범위', '각 방향 중앙 21×21 영역(x,y=249~269)의 441개 픽셀. 전체 지도에서 뽑은 표본이며 얼굴 부위로 지정한 영역은 아님.'),
        ('전체영역표본', '사진 전체에서 약 32픽셀 간격으로 추출한 실제 픽셀값도 별도 시트에 수록. 보간·평균·가상 데이터 없음.'),
        ('깊이 값', 'depth head의 실제 출력. 해당 카메라 좌표계의 깊이이며 mm나 검증된 물리 거리가 아님.'),
        ('신뢰도 값', 'depth head의 내부 confidence 출력. 1보다 클 수 있고 정답 확률·퍼센트가 아님.'),
        ('3D 좌표 X,Y,Z', 'VGGT 깊이+예측 카메라 행렬로 역투영한 공통 좌표계의 점. 이번 실행은 point head를 사용하지 않음.'),
        ('깊이와 Z의 차이', '깊이는 각 촬영 카메라 기준, XYZ는 공통 좌표계 기준이므로 측면에서는 깊이와 Z가 다를 수 있음.'),
        ('부위별 측정값', 'MediaPipe의 위치와 위 3D 정보를 합쳐 계산한 저장 결과. 눈 바깥꼬리 사이의 3D 거리로 나눈 무단위 비율.'),
        ('분류·점수', '400명 전체의 25%·75% 경계와 학습군 반복 측정 σ 사용. 낮음·보통·높음 점수는 정답 확률이 아님.'),
        ('엑셀 숫자', '셀에는 원본 숫자를 입력. 보기에는 소수점 9자리까지 표시하며 수식 입력줄에서는 더 많은 자릿수를 확인할 수 있음.'),
    ]
    for row in notes:
        overview.append(row)
    overview.append([])
    overview.append(['바로 보는 실제 값', '픽셀 (x,y)', '깊이', '신뢰도', 'X', 'Y', 'Z'])
    for v, name in enumerate(NAMES):
        vals = pixel_values(v, 259, 259)
        overview.append([name, '(259,259)', *vals[2:]])
    overview.append([])
    overview.append(['시트 바로가기', '설명'])
    links = [('부위별_측정값', '저장된 네 피처의 측정값·경계·분류·점수'),
             ('정면_픽셀값', '정면 픽셀별 실제 숫자'), ('좌측사선_픽셀값', '좌측 사선 픽셀별 실제 숫자'),
             ('우측사선_픽셀값', '우측 사선 픽셀별 실제 숫자'),
             ('전체영역표본', '배경을 포함한 전체 영역의 표본'), ('카메라행렬', '실제 extrinsics·intrinsics'),
             ('400명_측정값', '400명의 네 피처와 분류'), ('깊이그림', '같은 저장 결과를 색으로 표시한 그림')]
    for name, desc in links:
        overview.append([name, desc])
        overview.cell(overview.max_row, 1).hyperlink = f"#'{name}'!A1"
        overview.cell(overview.max_row, 1).style = 'Hyperlink'

    ws = sheet('부위별_측정값', ['특징','실제 측정값 r','하한 b1','상한 b2','σ(학습군 반복)','분류','낮음 점수','보통 점수','높음 점수','계산 기준(모두 눈 사이 거리로 나눔)'])
    for key, (name, desc) in FEATURES.items():
        r = sample['values'][key]
        c = config[key]
        lower, upper, sigma = c['lower'], c['upper'], c['sigma_repeat_train']
        phi = lambda t: .5 * (1 + math.erf(t / math.sqrt(2)))
        p1, p2 = phi((r-lower)/sigma), phi((r-upper)/sigma)
        label = '낮음' if r < lower else ('보통' if r < upper else '높음')
        ws.append([name,r,lower,upper,sigma,label,1-p1,p1-p2,p2,desc])
        assert abs((1-p1)+(p1-p2)+p2-1) < 1e-12
    ws.append([])
    ws.append(['눈 바깥꼬리 거리(모델 단위)', sample['eye_distance_model_units']])
    ws.append(['분류 규칙', 'r < b1: 낮음 / b1 ≤ r < b2: 보통 / r ≥ b2: 높음'])
    ws.append(['점수 식', 'A=Φ((r-b1)/σ), B=Φ((r-b2)/σ); 낮음=1-A, 보통=A-B, 높음=B'])
    ws.append(['턱 음수의 의미', '코 아래 기준점보다 뒤쪽. 낮음·높음은 이 비율의 크기 기준.'])
    ws.append(['성능 의미', '모델이 추정한 값. 실제 3D 스캔과의 정확도 검증값은 아님.'])

    headers = ['픽셀 x','픽셀 y','깊이(모델 단위)','깊이 신뢰도(확률 아님)','X(공통 좌표)','Y(공통 좌표)','Z(공통 좌표)']
    for v, name in enumerate(NAMES):
        ws = sheet(name+'_픽셀값', headers)
        for y in range(249, 270):
            for x in range(249, 270):
                vals = pixel_values(v,x,y)
                ws.append(vals)
                numeric_checks.append((ws.title,ws.max_row,vals))
        ws.auto_filter.ref = ws.dimensions
        matrix = sheet(name+'_깊이격자', ['y / x', *range(249,270)])
        for y in range(249,270):
            matrix.append([y,*map(float, arrays['depth'][v,y,249:270,0])])
        matrix.freeze_panes = 'B2'
        matrix.conditional_formatting.add('B2:V22', ColorScaleRule(start_type='min',start_color='D9EAF7',end_type='max',end_color='FFC857'))

    ws = sheet('전체영역표본', ['방향', *headers])
    indices = sorted(set([*range(0,518,32),517]))
    for v,name in enumerate(NAMES):
        for y in indices:
            for x in indices:
                ws.append([name,*pixel_values(v,x,y)])
    ws.auto_filter.ref = ws.dimensions
    ws = sheet('카메라행렬', ['방향','행렬','행(0부터)','열0','열1','열2','열3','설명'])
    for v,name in enumerate(NAMES):
        for key in ['extrinsics','intrinsics']:
            for row,values in enumerate(arrays[key][v]):
                vals = list(map(float,values))
                if len(vals)==3: vals.append(None)
                ws.append([name,key,row,*vals,'world → camera [R|t]' if key=='extrinsics' else '518×518 입력의 카메라 내부 행렬 K'])

    ws = sheet('400명_측정값', ['사례','분할',*[n for name,_ in FEATURES.values() for n in [name,name+' 분류']]])
    rows = {}
    for line in (SRC/'cohort_518.jsonl').read_text(encoding='utf-8').splitlines():
        row=json.loads(line)
        if row['status']=='measured':rows[row['subject_id']]=row
    assert len(rows)==400
    for index,(_,row) in enumerate(sorted(rows.items()),1):
        vals=[f'사례{index:03d}',row['split']]
        for key in FEATURES:
            r=row['values'][key]; c=config[key]
            vals.extend([r,'낮음' if r<c['lower'] else ('보통' if r<c['upper'] else '높음')])
        ws.append(vals)
    ws.auto_filter.ref=ws.dimensions
    ws = book.create_sheet('깊이그림')
    ws.append(['같은 npz 파일의 실제 깊이 지도. 원본 얼굴 사진은 포함하지 않음.'])
    ws.append(['각 패널 색 범위가 다름. 보라색: 가까움, 노란색: 멂. 정확한 수치는 픽셀값 시트 참고.'])
    image=Image(str(SRC/'private_depth_example.png'))
    image.width=1056;image.height=352
    ws.add_image(image,'A4')
    ws = sheet('출처', ['항목','값'])
    ws.append(['샘플 대응','smoke_518.jsonl 첫 번째 기록과 private_depth_example.npz. 원래 참여자 식별자는 공유 파일에서 생략.'])
    for name in ['private_depth_example.npz','smoke_518.jsonl','cohort_518.jsonl','boundary_config.json']:
        ws.append([name+' SHA256',hashlib.sha256((SRC/name).read_bytes()).hexdigest()])
    for key in arrays:
        a=arrays[key]
        ws.append([key, f'shape={a.shape}; dtype={a.dtype}; min={a.min():.12g}; max={a.max():.12g}'])
    ws.append(['생성 코드','scripts/pilot/export_vggt_sample.py'])

    for ws in book:
        ws.sheet_view.showGridLines=False
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        for row in ws:
            for cell in row:
                if cell.row>1 or ws==overview:
                    cell.font=Font(name='맑은 고딕',size=11,color='17365D' if cell.hyperlink else '222222')
                if isinstance(cell.value,float):cell.number_format='0.000000000'
        for col in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(col)].width=23
        if '깊이격자' in ws.title:
            for col in range(2,23):ws.column_dimensions[get_column_letter(col)].width=16
        if ws.title in ['먼저보기','출처']:
            ws.column_dimensions['A'].width=32
            ws.column_dimensions['B'].width=105
            for row in ws:
                for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='center')
                ws.row_dimensions[row[0].row].height=34
        if ws.title=='부위별_측정값':ws.column_dimensions['J'].width=70
    overview.freeze_panes='A2'
    book.save(target)
    print('SAVED',target,flush=True)

    # Compare every exported central pixel with the saved raw arrays after reopening.
    check=load_workbook(target,data_only=True,read_only=True)
    count=0
    for name in NAMES:
        v=NAMES.index(name)
        for row in check[name+'_픽셀값'].iter_rows(min_row=2,values_only=True):
            x,y=int(row[0]),int(row[1])
            assert np.allclose(row,pixel_values(v,x,y),rtol=1e-14,atol=1e-15)
            count+=1
        for row in check[name+'_깊이격자'].iter_rows(min_row=2,values_only=True):
            y=int(row[0]);assert np.allclose(row[1:],arrays['depth'][v,y,249:270,0],rtol=1e-14)
    assert count==1323
    assert check['400명_측정값'].max_row==401
    check.close()
    report={'file':str(target),'sheets':book.sheetnames,'central_pixel_rows_verified':count,
            'overview_sample_pixels':{name:pixel_values(v,259,259) for v,name in enumerate(NAMES)},
            'cohort_people':400,'bytes':target.stat().st_size,'status':'passed'}
    (ROOT/'outputs/pilot-20261014/excel_sample_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False),flush=True)


if __name__=='__main__':main()
