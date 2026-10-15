"""Offline paired audit of existing CNN probabilities and VGGT boundary scores.

No inference, fitting, label changes, or running-service changes. Rank against
the actual draft literature repository, separately from score-only diagnostics.
"""
import csv
import hashlib
import itertools
import json
import math
import sys
from collections import Counter, defaultdict
from dataclasses import asdict
from pathlib import Path

import numpy as np
from scipy.special import softmax
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))
from facework.pilot.scoring import boundary_scores
from facework.pilot_contract import Observation, DecisionPolicy, PassageSelector, load_real_repository

OUT=ROOT/'outputs/pilot-20261014'
D=ROOT/'data/derived/pilot-20261014'
DEPTH=['nose_projection','nose_bridge','lip_projection','chin_projection']

def read(p):return json.loads(p.read_text(encoding='utf-8'))

def summary(values):
    a=np.asarray(values,float)
    if not len(a):return {'n':0}
    return {'n':len(a),'mean':float(a.mean()),'median':float(np.median(a)),
            'p10':float(np.quantile(a,.1)),'p90':float(np.quantile(a,.9)),
            'ge_060':float((a>=.6).mean()),'ge_090':float((a>=.9).mean()),
            'ge_099':float((a>=.99).mean()),'mean_negative_log':float(-np.log(np.clip(a,1e-15,1)).mean())}

def gmean(xs):return float(np.exp(np.mean(np.log(xs))))

def main():
    cache=read(OUT/'cnn_raw_predictions.json')
    calibration=read(OUT/'cnn_fold0_calibration.json')['features']
    schema=read(ROOT/'data/label_schema_kface_v3.json')['features']
    cuts=read(D/'boundary_config.json')['features']
    cohort={}
    for line in (D/'cohort_518.jsonl').read_text(encoding='utf-8').splitlines():
        row=json.loads(line)
        if row['status']=='measured':cohort[row['subject_id']]=row
    repo,audit=load_real_repository(ROOT,allow_draft=True)
    refs=list(csv.DictReader((ROOT/'data/derived/kface/reference_labels.csv').open(encoding='utf-8-sig')))
    ids=sorted(r['subject_id'] for r in refs if r['split']=='test')
    assert len(ids)==62
    records=[];by_feature=defaultdict(list); scenarios={};all_rank_rows=[]
    for boundary_scope in ['deployed_400','train_276']:
        all_c=[];all_d=[];accepted_c=[];accepted_d=[];paired=[];pair_g=defaultdict(list)
        common_gate_pairs=[];common_gate_wins=[]
        depth_by_label=defaultdict(list);person_rows=[];rank_stats=defaultdict(Counter)
        eligible_details=[];per_feature=defaultdict(list)
        for index,sid in enumerate(ids,1):
            saved=cache['test'][hashlib.sha256(sid.encode()).hexdigest()]
            cnns=[];depths=[];observations=[]
            for k,entry in saved['fold0'].items():
                fit=calibration[k];assert fit['calibrated']
                p=softmax(np.log(np.clip(entry['prob'],1e-12,1))/fit['temperature'])
                j=int(p.argmax());score=float(p[j]);f=schema[k]
                label=fit['classes'][j]
                blocked='all' in saved['blocked'] or f['part'] in saved['blocked']
                cnns.append((k,score,not blocked, j==entry['truth_index']))
                observations.append(Observation(k,f['part'],'classification',k+':'+label,label,
                    confidence=score,class_count=len(p),calibrated=True,basis='saved fold0 prediction + frozen validation temperature',
                    hold_reasons=('RECORDED_QUALITY_GATE',) if blocked else ()))
                if boundary_scope=='deployed_400':
                    records.append({'case':index,'kind':'CNN','feature':k,'label':label,'score':score,
                                    'gate_pass':not blocked,'accepted':not blocked and score>=.6,
                                    'reference_correct':j==entry['truth_index']})
            for k in DEPTH:
                c=cuts[k];bc=c if boundary_scope=='deployed_400' else c['train_only_calibration']
                z=boundary_scores(cohort[sid]['values'][k],bc['lower'],bc['upper'],c['sigma_repeat_train'])
                j=['low','medium','high'].index(z['label']);f=schema[k];label=f['classes'][j]
                score=z['score'];depths.append((k,score));depth_by_label[z['label']].append(score)
                observations.append(Observation(k,f['part'],'depth',k+':'+label,label,
                    measurement={'indicator':k,'value':cohort[sid]['values'][k],'unit':'ratio'},
                    measurement_score=score,quality_passed=True,boundaries_validated=True,
                    basis='saved three-view geometry; exploratory finite-geometry gate'))
                if boundary_scope=='deployed_400':
                    records.append({'case':index,'kind':'Depth','feature':k,'label':label,'score':score,
                                    'gate_pass':True,'accepted':True,'reference_correct':None})
            cs=[s for _,s,_,_ in cnns];ds=[s for _,s in depths]
            ca=[s for _,s,ok,_ in cnns if ok and s>=.6];da=[s for s in ds if s>=.6]
            cg=[s for s in cs if s>=.6]
            if cg and da:
                common_gate_pairs.append(float(np.mean(da)-np.mean(cg)))
                common_gate_wins.append(float(np.mean(np.asarray(da)[:,None]>np.asarray(cg)[None,:])))
            all_c.extend(cs);all_d.extend(ds);accepted_c.extend(ca);accepted_d.extend(da)
            for k,s,_,_ in cnns:per_feature[k].append(s)
            for k,s in depths:per_feature[k].append(s)
            probability=float(np.mean(np.asarray(ds)[:,None]>np.asarray(ca)[None,:])) if ca else None
            matched=float(np.mean(np.asarray(da)[:,None]>np.asarray(ca)[None,:])) if ca and da else None
            row={'case':index,'cnn_n':len(cs),'cnn_accepted_n':len(ca),'cnn_mean':float(np.mean(cs)),
                 'cnn_accepted_mean':float(np.mean(ca)) if ca else None,'depth_mean':float(np.mean(ds)),
                 'depth_ge060_mean':float(np.mean(da)) if da else None,
                 'depth_beats_accepted_cnn_pair_rate':probability,
                 'equal_gate_depth_beats_cnn_pair_rate':matched}
            person_rows.append(row)
            if ca and da:paired.append(float(np.mean(da)-np.mean(ca)))
            # Fixed two-feature compositions: average within person so the much
            # larger number of CNN feature combinations does not dominate.
            for name,vals in [('CNN2',list(itertools.combinations(ca,2))),
                              ('CNN1_Depth1',list(itertools.product(ca,da))),
                              ('Depth2',list(itertools.combinations(da,2)))]:
                if vals:pair_g[name].append(float(np.mean([gmean(x) for x in vals])))
            decisions=DecisionPolicy().apply(observations)
            for gate in ['current_depth_no_score_gate','equal_min_060']:
                items=[d for d in decisions if gate!='equal_min_060' or d.kind!='depth' or d.measurement_score>=.6]
                available={d.label_id:d for d in items if d.status=='CONFIRMED'}
                candidates=repo.load_candidates(set(available))
                eligible=[p for p in candidates if p.required_labels<=available.keys()]
                selected={mode:PassageSelector(mode).select(items,candidates) for mode in ['classification_only','include_measurement_scores']}
                old,new=selected['classification_only'],selected['include_measurement_scores']
                old_ids=[r['passage_id'] for r in old];new_ids=[r['passage_id'] for r in new]
                count=rank_stats[gate];count['people']+=1;count['eligible_total']+=len(eligible)
                count['people_with_candidates']+=bool(eligible);count['people_with_more_than5_candidates']+=len(eligible)>5
                count['people_with_5_selected']+=len(new)==5
                count['people_order_changed']+=old_ids!=new_ids;count['people_selection_changed']+=set(old_ids)!=set(new_ids)
                count['people_top1_changed']+=bool(old_ids and new_ids and old_ids[0]!=new_ids[0])
                count['max_eligible']=max(count['max_eligible'],len(eligible))
                for p in eligible:
                    ev=[available[k] for k in sorted(p.required_labels)]
                    cc=[x.confidence for x in ev if x.kind=='classification'];dd=[x.measurement_score for x in ev if x.kind=='depth']
                    group='mixed' if cc and dd else 'depth_only' if dd else 'cnn_only'
                    count['eligible_'+group]+=1
                    if cc and dd:
                        count['mixed_increased_G']+=gmean(cc+dd)>gmean(cc)
                    detail={'scope':boundary_scope,'gate':gate,'case':index,'passage_id':p.passage_id,'group':group,
                            'cnn_scores':cc,'depth_scores':dd,'old_G':gmean(cc) if cc else None,'mixed_G':gmean(cc+dd),
                            'old_rank':old_ids.index(p.passage_id)+1 if p.passage_id in old_ids else None,
                            'mixed_rank':new_ids.index(p.passage_id)+1 if p.passage_id in new_ids else None}
                    eligible_details.append(detail)
        rng=np.random.default_rng(20261014)
        boot=np.mean(rng.choice(paired,size=(10000,len(paired)),replace=True),axis=1)
        scenarios[boundary_scope]={'CNN_all':summary(all_c),'depth_all':summary(all_d),
            'CNN_ge060_score_only':summary([s for s in all_c if s>=.6]),
            'common_gate_score_only_people':len(common_gate_pairs),
            'common_gate_score_only_mean_gap':float(np.mean(common_gate_pairs)),
            'common_gate_score_only_depth_win_rate':float(np.mean(common_gate_wins)),
            'CNN_accepted':summary(accepted_c),'depth_ge060':summary(accepted_d),
            'feature_scores':{k:summary(v) for k,v in per_feature.items()},
            'depth_by_label':{k:summary(v) for k,v in depth_by_label.items()},
            'medium_score_theoretical_max':{k:float(math.erf(((cuts[k] if boundary_scope=='deployed_400' else cuts[k]['train_only_calibration'])['upper']-(cuts[k] if boundary_scope=='deployed_400' else cuts[k]['train_only_calibration'])['lower'])/(2*cuts[k]['sigma_repeat_train']*math.sqrt(2)))) for k in DEPTH},
            'paired_people_mean_gap_equal_gate':float(np.mean(paired)),
            'paired_people_bootstrap95_gap_equal_gate':np.quantile(boot,[.025,.975]).tolist(),
            'depth_beats_accepted_cnn_pair_rate_mean':float(np.mean([r['depth_beats_accepted_cnn_pair_rate'] for r in person_rows if r['depth_beats_accepted_cnn_pair_rate'] is not None])),
            'equal_gate_depth_beats_cnn_pair_rate_mean':float(np.mean([r['equal_gate_depth_beats_cnn_pair_rate'] for r in person_rows if r['equal_gate_depth_beats_cnn_pair_rate'] is not None])),
            'two_feature_G_person_mean':{k:summary(v) for k,v in pair_g.items()},
            'ranking':{k:dict(v) for k,v in rank_stats.items()},'person_rows':person_rows}
        all_rank_rows.extend(eligible_details)
    output={'people':62,'new_inference':False,'source':'existing paired participant predictions; CNN original frontal pipeline; VGGT guide-cropped triplets',
        'CNN_missing_heads':'cache contains only heads with available reference labels; not a complete replay of live requests',
        'runtime_scoring_mode':'classification_only; depth excluded from G in current live service',
        'ranking_experiment_mode':'include_measurement_scores; offline only; running service unchanged',
        'depth_gate':'finite geometry exploratory acceptance, not independent 3D quality validation',
        'repository_audit':audit,'scenarios':scenarios,
        'limitations':['Different score definitions; no common independently human-verified correctness target.',
                       'Real reference pool only14 eligible draft passages; missing regression heads excluded.',
                       'Score distribution equality is not evidence of equal accuracy; composition diagnostic is not a real literature quality score.'],
        'feature_records':records,'rank_records':all_rank_rows}
    (OUT/'joint_score_audit.json').write_text(json.dumps(output,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:{a:b for a,b in v.items() if a not in ['person_rows','feature_scores']} for k,v in scenarios.items()},ensure_ascii=False,indent=2))
    dest=Path.home()/'Downloads/FaceTeller_파일럿_공유자료';dest.mkdir(exist_ok=True,parents=True)
    book=Workbook();book.remove(book.active)
    tables={
      '점수비교':[['경계 설정','종류','개수','평균','중앙값','90점 이상 비율','99점 이상 비율']]+[
          [scope,kind,v['n'],v['mean'],v['median'],v['ge_090'],v['ge_099']]
          for scope,s in scenarios.items() for kind in ['CNN_all','depth_all','CNN_ge060_score_only','CNN_accepted','depth_ge060'] for v in [s[kind]]],
      '실제특징점수':[['사례','종류','피처','레이블','점수','품질통과','현재기준 채택','CNN 초벌정답 일치']]+[
          [r[k] for k in ['case','kind','feature','label','score','gate_pass','accepted','reference_correct']] for r in records],
      '피처별비교':[['경계 설정','피처','개수','평균','중앙값','90점 이상 비율']]+[
          [scope,k,v['n'],v['mean'],v['median'],v['ge_090']] for scope,s in scenarios.items() for k,v in s['feature_scores'].items()],
      '문헌순위비교':[['경계 설정','통과조건','사례','문헌ID','구성','CNN 점수','깊이 점수','기존 G','혼합 G','기존 순위','혼합 순위']]+[
          [r[k] if k not in ['cnn_scores','depth_scores'] else ', '.join(f'{x:.6f}' for x in r[k]) for k in ['scope','gate','case','passage_id','group','cnn_scores','depth_scores','old_G','mixed_G','old_rank','mixed_rank']] for r in all_rank_rows],
      '실험범위':[['항목','내용'],['대상','동일한 테스트 대상 62명의 저장된 실제 CNN·VGGT 출력'],
          ['실제 서비스','현재 G는 CNN만 사용. 혼합 계산은 이번 오프라인 비교 실험에서 실행.'],
          ['경계 설정','현재 전체400명 경계 + 학습276명 경계 두 경우 비교. σ는 학습군26명 반복 측정값 고정.'],
          ['CNN 통과','보정된 점수≥0.6 + 기록된 품질조건. 초벌정답이 있는 피처만 캐시에 존재.'],
          ['깊이 통과','현재는 수치 유효성 기준. 공통0.6 조건도 별도 비교.'],
          ['점수의 의미','CNN 보정 확률과 깊이 경계점수는 동일한 정답률이 아님. 분포 비교만으로 정확도 공평성을 입증할 수 없음.'],
          ['문헌','실제 기존 초안14건으로 비교. 새 문헌이나 허구의 판정조건을 추가하지 않음.'],
          ['원본','outputs/pilot-20261014/joint_score_audit.json 및 scripts/pilot/check_joint_scores.py']]}
    for name,rows in tables.items():
        ws=book.create_sheet(name)
        for row in rows:ws.append(row)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.fill=PatternFill('solid',fgColor='17365D');c.font=Font(name='맑은 고딕',bold=True,color='FFFFFF')
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font=Font(name='맑은 고딕',size=10)
                if isinstance(c.value,float):c.number_format='0.000000'
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=25
        if name=='실험범위':ws.column_dimensions['B'].width=125
    ws=book.create_sheet('결론과주의점',0)
    findings=[
      ['항목','확인 결과'],
      ['결론','점수 범위는 대체로 비슷하지만 완전히 대등하지는 않음. 깊이 점수가 조금 더 높은 경향.'],
      ['같은 점수 조건','양쪽 모두0.6 이상만 비교: CNN 평균0.8383(741개), 깊이 평균0.8573(221개). 품질 게이트는 이 비교에서 제외.'],
      ['현재 전체값','CNN908개 평균0.7829, 깊이248개 평균0.8242. 동일 테스트62명 저장 출력.'],
      ['기존 CNN 품질 게이트','CNN은 품질 및0.6 조건을 모두 적용하면645개 평균0.8373. 깊이는 독립 검증된 품질 게이트가 없어서 조건이 완전히 같지 않음.'],
      ['보통 점수 상한','코끝0.8516, 콧뿌리0.9098, 입술0.9388, 턱0.9849. 낮음·높음 점수는1에 접근 가능.'],
      ['상한의 해석','보통은 양쪽 경계를 모두 고려한 유한 구간이라서 발생. 단순 계산 오류가 아니며, 순위에서 극단 범주가 더 높은 점수를 받을 여지가 있음.'],
      ['실제 문헌 비교','62명 중24명에서 후보 발생. 총30개 후보 발생, 사람당 최대2개. 2명에서1위 변경, 선택 집합 변경0명. 상위5개 경쟁 검증 불가.'],
      ['기하평균 작동 사례','사례42: CNN0.9653+깊이0.6141 후보는 혼합G0.7700으로 내려가고, 다른 후보G0.8716이1위.'],
      ['현재 서비스 상태','실행 중인 서비스는 classification_only. 이 파일의 혼합순위는 오프라인 비교 결과이며 서비스 설정을 변경하지 않았음.'],
      ['판단','혼합 점수를 파일럿 순위에 사용하는 것은 타당. 깊이를 무조건 깎거나 모든 피처 점수 분포를 같게 만드는 보정은 이번 결과만으로 정당화되지 않음.'],
      ['검증 한계','깊이 점수와CNN 점수의 동일 정답률은 입증되지 않음. 초벌정답이 있는CNN 헤드만 포함, 실제문헌 초안14건, 입력 전처리도 모듈별로 다름.'],
    ]
    for row in findings:ws.append(row)
    ws.freeze_panes='A2';ws.column_dimensions['A'].width=30;ws.column_dimensions['B'].width=130
    from openpyxl.styles import Alignment
    for row in ws:
        for cell in row:
            cell.font=Font(name='맑은 고딕',size=11)
            cell.alignment=Alignment(wrap_text=True,vertical='center')
        ws.row_dimensions[row[0].row].height=42
    book.active=0
    file=dest/'FaceTeller_CNN_깊이점수_비교검증.xlsx';book.save(file)
    report='CNN·VGGT 깊이 점수 혼합 비교 검증\n\n'+'\n\n'.join(a+'\n'+b for a,b in findings[1:])
    report+='\n\n실험 코드: scripts/pilot/check_joint_scores.py\n수치 원본: outputs/pilot-20261014/joint_score_audit.json\n'
    (dest/'FaceTeller_CNN_깊이점수_비교검증.txt').write_text(report,encoding='utf-8-sig')
    (OUT/'joint_score_review.txt').write_text(report,encoding='utf-8')
    reopened=load_workbook(file,read_only=True,data_only=True)
    assert reopened['실제특징점수'].max_row==len(records)+1
    reopened.close();print('EXCEL',file)

if __name__=='__main__':main()
