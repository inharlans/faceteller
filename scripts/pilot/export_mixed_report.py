"""Readable exports of the real HTTP mixed-score service verification."""
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT=Path(__file__).resolve().parents[2]
SRC=ROOT/'outputs/pilot-20261014/mixed-service'
DEST=Path.home()/'Downloads/FaceTeller_파일럿_공유자료'

def main():
    report=json.loads((SRC/'live_verification.json').read_text(encoding='utf-8'))
    results=[json.loads(x) for x in (SRC/'actual_responses.jsonl').read_text(encoding='utf-8').splitlines()]
    t=report['timing'];c=report['counts']
    findings=[
        ('적용 상태','실행 중인 /api/v1/analysis 기본 선정 방식에 CNN+깊이 기하평균 적용 완료.'),
        ('정책 버전','cnn_depth_geomean_v1 / include_measurement_scores'),
        ('실제 요청',f"{report['actual_people']}명 ×3장 = {report['actual_photos']}장. 실제 로컬 TCP HTTP 요청으로 검증."),
        ('검증한 문헌',f"{c['verified_candidate_formulas']}건 모두 실제 반영 항목 수와 G를 별도 재계산하여 일치 확인."),
        ('문헌 구성',f"CNN만 {c['selected_cnn_only']}건 / CNN+깊이 {c['selected_mixed']}건 / 깊이만 {c['selected_depth_only']}건."),
        ('순위 변화',f"같은 응답의 CNN만 사용하는 기존 방식과 비교하여 {c['order_changed_people']}명의 순서가 달라짐. 선택 집합 변화 {c['selection_changed_people']}명."),
        ('후보 확보',f"{c['people_with_passages']}명에서 문헌이 선정됨. 나머지는 가능한 형태 결과만 반환. 후보를 임의로 채워5개를 만들지 않음."),
        ('실제 처리 시간',f"평균 {t['mean']:.3f}초 / P95 {t['p95']:.3f}초 / 최대 {t['max']:.3f}초 / 5초 이내 {c['within5_seconds']}명."),
        ('시간 범위','모델 로딩·예열 제외. ZIP사진 읽기·가이드 자르기·JPEG 인코딩 이후 HTTP 송수신+전처리+CNN+VGGT+메모리 문헌조회 포함.'),
        ('오류 처리','동일 정면3장 입력 시 깊이만 보류하고 CNN 결과 유지. 빈 측면 이미지 입력 시 해당 방향을 지정한422 응답 확인.'),
        ('단위·계약 검사','계약·깊이점수 pytest47개 통과. 실제 격리 PostgreSQL 검사34개 통과. DB 검사와GPU 실서비스 요청은 별도 수행.'),
        ('선정 규칙','모든 필수 조건 충족 / 보류 제외 / 중복 제외 / 부위별 최대3개 / 전체 최대5개 / 기존 동점 정렬 유지.'),
        ('숫자 의미','이 엑셀 점수는0~100점 환산. 후보의 종합 일치점수이며 정답일 확률이 아님.'),
        ('현 단계 한계','실제3D 정답 검증, 새7개 부위 모델, 문헌 독립 검수, 실물 카메라 촬영 검증은 아직 미완료.'),
        ('원본 변경','제출 PDF 수정 없음. 추가 결정 적용 문서는 FaceTeller_혼합점수_서비스적용.txt.'),
        ('실행 주소','http://127.0.0.1:8767/pilot'),
    ]
    tables={'결과요약':[['항목','확인 결과'],*findings],
            '62명_HTTP결과':[['사례','HTTP 초','상태','확정 항목','보류 항목','선정 문헌 수','기존 순서','혼합 순서','순서 변경']],
            '실제선정문헌':[['사례','순위','문헌ID','문헌명','종합점수(100점)','반영 항목 수','구성']],
            '점수구성':[['사례','문헌 순위','문헌ID','항목','레이블','종류','항목점수(100점)','해당 문헌 종합점수(100점)']]}
    for row in report['records']:
        tables['62명_HTTP결과'].append([row['case'],row['http_seconds'],row['status'],row['confirmed'],row['held'],row['selected_count'],', '.join(row['previous_ids']),', '.join(row['mixed_ids']),row['order_changed']])
    for row in results:
        for p in row['result']['passages']:
            kinds={x['kind'] for x in p['score_components']}
            group='CNN+깊이' if len(kinds)==2 else ('깊이' if 'depth' in kinds else 'CNN')
            tables['실제선정문헌'].append([row['case'],p['rank'],p['passage_id'],p['source']['passage_no'],p['G']*100,p['score_feature_count'],group])
            for x in p['score_components']:
                tables['점수구성'].append([row['case'],p['rank'],p['passage_id'],x['item_id'],x['label_name'],'깊이' if x['kind']=='depth' else 'CNN',x['score']*100,p['G']*100])
    book=Workbook();book.remove(book.active)
    for name,rows in tables.items():
        ws=book.create_sheet(name)
        for row in rows:ws.append(row)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for row in ws:
            for cell in row:
                cell.font=Font(name='맑은 고딕',size=11)
                cell.alignment=Alignment(wrap_text=True,vertical='center')
                if isinstance(cell.value,float):cell.number_format='0.000'
            ws.row_dimensions[row[0].row].height=34
        for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='17365D');cell.font=Font(name='맑은 고딕',bold=True,color='FFFFFF')
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=25
        if name=='결과요약':ws.column_dimensions['B'].width=118;ws.row_dimensions[10].height=48
        if name=='62명_HTTP결과':
            ws.column_dimensions['G'].width=50;ws.column_dimensions['H'].width=50
    DEST.mkdir(parents=True,exist_ok=True)
    target=DEST/'FaceTeller_혼합점수_실서비스_검증.xlsx';book.save(target)
    saved=load_workbook(target,read_only=True,data_only=True)
    assert saved['62명_HTTP결과'].max_row==63
    assert saved['실제선정문헌'].max_row==45
    saved.close()
    text='FaceTeller 혼합 점수 실서비스 적용·검증\n\n'+'\n\n'.join(k+'\n'+v for k,v in findings)
    (DEST/'FaceTeller_혼합점수_실서비스_검증.txt').write_text(text,encoding='utf-8-sig')
    print(target)

if __name__=='__main__':main()
