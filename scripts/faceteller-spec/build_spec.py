from pathlib import Path
import json, hashlib, re
from docx import Document
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parents[2]; OUT=ROOT/'outputs/faceteller-spec'; FIG=OUT/'figures'
D=Document(); MD=[]; sections=[]; reqs=[]
sec=D.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297);sec.top_margin=Mm(18);sec.bottom_margin=Mm(17);sec.left_margin=Mm(21);sec.right_margin=Mm(21);sec.header_distance=Mm(8);sec.footer_distance=Mm(8);sec.different_first_page_header_footer=True
styles=D.styles
for name in ['Normal','Body Text','Caption','Heading 1','Heading 2','Heading 3','TOC 1','TOC 2']:
 try:s=styles[name]
 except KeyError:s=styles.add_style(name,1)
 s.font.name='맑은 고딕';s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕');s.font.color.rgb=RGBColor.from_string('111111')
 s.font.size=Pt(10.5);s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=Pt(14)
 grid=OxmlElement('w:snapToGrid');grid.set(qn('w:val'),'0');s._element.get_or_add_pPr().append(grid)
styles['Heading 1'].font.size=Pt(17);styles['Heading 1'].font.bold=True;styles['Heading 1'].paragraph_format.line_spacing=Pt(22)
styles['Heading 2'].font.size=Pt(13);styles['Heading 2'].font.bold=True;styles['Heading 2'].paragraph_format.line_spacing=Pt(18)
styles['Heading 3'].font.size=Pt(11);styles['Heading 3'].font.bold=True
for name in ['Heading 1','Heading 2','Heading 3']:
 styles[name].paragraph_format.space_before=Pt(9);styles[name].paragraph_format.keep_with_next=True
styles['Caption'].font.size=Pt(9);styles['Caption'].font.italic=False
header=sec.header.paragraphs[0];header.text='페이스텔러 (FaceTeller)  |  상세 요구사항 명세서';header.style='Caption'
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
foot.add_run('− ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld);foot.add_run(' −')

def para(t='',style=None):
 p=D.add_paragraph(t,style);MD.append(t+'\n');return p

def heading(t,l=2):
 D.add_heading(t,l);MD.append('#'*l+' '+t+'\n');sections.append((l,t))

def page(t,l=2):
 heading(t,l);D.paragraphs[-1].paragraph_format.page_break_before=True

def sub(t):heading(t,3)

def table(headers,rows,widths=None):
 t=D.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths:
  for c,w in zip(t.columns,widths):c.width=Mm(w)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h)
 for row in rows:
  cells=t.add_row().cells
  for c,v in zip(cells,row):c.text=str(v)
 for ri,row in enumerate(t.rows):
  if widths:
   for c,w in zip(row.cells,widths):c.width=Mm(w)
  trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if ri==0:trPr.append(OxmlElement('w:tblHeader'))
  for c in row.cells:
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'555555');b.append(el)
   pr.append(b);mar=OxmlElement('w:tcMar')
   for side,v in [('top','65'),('bottom','65'),('left','90'),('right','90')]:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),v);el.set(qn('w:type'),'dxa');mar.append(el)
   pr.append(mar)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F7E7DE');pr.append(sh)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=Pt(12)
    for run in p.runs:
     run.font.size=Pt(9.5);run.bold=ri==0
 MD.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+'\n'.join('| '+' | '.join(str(x).replace('\n','<br>') for x in r)+' |' for r in rows)+'\n')
 gap=D.add_paragraph();gap.paragraph_format.space_after=Pt(0);gap.paragraph_format.space_before=Pt(0);gap.paragraph_format.line_spacing=Pt(3);gap.add_run(' ').font.size=Pt(1)
 return t

def kv(rows):return table(['항목','내용'],rows,[30,138])

def fig(name,caption,width=165,maxh=145):
 from PIL import Image
 im=Image.open(FIG/(name+'.png'));w,h=im.size
 ww=min(width,maxh*w/h)
 p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(FIG/(name+'.png')),width=Mm(ww))
 p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.0
 p=para(caption,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 MD.append(f'![{caption}](figures/{name}.png)\n')

def code(t):
 p=D.add_paragraph();p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=1.0
 for i,line in enumerate(t.splitlines()):
  r=p.add_run(('\n' if i else '')+line);r.font.name='Consolas';r.font.size=Pt(8.5)
 MD.append('```\n'+t+'\n```\n')

# cover
p=para('컴퓨터공학 종합설계  [202602-CSE4205-001]');p.paragraph_format.space_before=Pt(55)
p=para('상세 요구사항 명세서');p.runs[0].font.size=Pt(28);p.runs[0].bold=True;p.paragraph_format.line_spacing=1.1
ppr=p._p.get_or_add_pPr();bs=OxmlElement('w:pBdr')
for edge,col in [('top','9A3C3C'),('bottom','D9AAAA')]:
 el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'12');el.set(qn('w:space'),'15');el.set(qn('w:color'),col);bs.append(el)
ppr.append(bs)
para('AI 기반 얼굴 형태 분류 및 전통 관상 문헌 연계 연구').paragraph_format.space_before=Pt(22)
para('시스템 · 데이터 · 화면 · 기능 상세설계').paragraph_format.space_after=Pt(28)
kv([('팀 이름','페이스텔러 (FaceTeller)'),('팀원','12211669 이태용 / 12211563 곽동훈 / 12211658 이기문'),('작성일 / 버전','2026.10.01 / 1.0 (구현 기준 설계안)')])
para('작성 기준', 'Heading 3')
para('「2026년 2학기 컴퓨터공학 종합설계_상세요구사항 명세서.pdf」를 유일한 요구사항 기준(SSOT)으로 한다. 해당 문서의 목차와 FR_001~FR_010을 유지하고, 상세설계서 샘플의 표·구성도·API·DB 설명 형식을 적용하였다. HTML 목업은 화면 구성 참고 자료이다.')
para('이 문서는 구현할 동작과 확인 기준을 정의한다. 기존 프로토타입의 실험 결과는 필요한 부분에서만 구분해 인용하며, 설계 수치를 완료된 성능으로 표시하지 않는다.','Caption')
p=para('목차');p.paragraph_format.page_break_before=True;p.runs[0].font.size=Pt(17);p.runs[0].bold=True;p.paragraph_format.line_spacing=Pt(22)
for name in ['TOC 1','TOC 2']:
 styles[name].font.size=Pt(10);styles[name].paragraph_format.space_after=Pt(1);styles[name].paragraph_format.line_spacing=Pt(13)
p=D.add_paragraph();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u');p._p.append(fld)
MD.append('[목차: Word 자동 생성]\n')
for filename in ['content_1.py','content_2.py','content_3.py','content_4.py','content_5.py']:
 exec(compile((Path(__file__).parent/filename).read_text(encoding='utf-8'),filename,'exec'))
OUT.mkdir(exist_ok=True)
D.save(OUT/'페이스텔러_상세요구사항명세서_1.0.docx')
(OUT/'페이스텔러_상세요구사항명세서_1.0.md').write_text('\n'.join(MD),encoding='utf-8')
(OUT/'문서구성.json').write_text(json.dumps({'sections':sections,'requirements':reqs},ensure_ascii=False,indent=2),encoding='utf-8')
print('DOCX complete',len(D.paragraphs),'paragraphs',len(D.tables),'tables',len(D.inline_shapes),'figures')
