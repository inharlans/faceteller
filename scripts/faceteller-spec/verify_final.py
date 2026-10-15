from pathlib import Path
import json,hashlib,csv,re
from pypdf import PdfReader
from docx import Document
R=Path(__file__).resolve().parents[2];O=R/'outputs/faceteller-spec';D=Document(O/'페이스텔러_상세요구사항명세서_1.0.docx');P=PdfReader(O/'페이스텔러_상세요구사항명세서_1.0.pdf');text='\n'.join(p.extract_text() or '' for p in P.pages)
reqs=[]
for t in D.tables:
 rows=[[c.text for c in r.cells] for r in t.rows]
 if len(rows)>1 and rows[1][0]=='ID / 중요도':reqs.append(rows[1][1].split(' / ')[0])
features=json.loads((R/'data/label_schema_kface_v3.json').read_text(encoding='utf-8'))['features']
headings=[p.text for p in D.paragraphs if p.style.name.startswith('Heading')]
images=[]
for i,p in enumerate(P.pages):
 for im in p.images:images.append({'page':i+1,'name':im.name,'size':list(im.image.size)})
checks={'pages':len(P.pages),'functional_requirements':reqs,'ten_fr_forms':reqs==[f'FR_{i:03d}' for i in range(1,11)],'tables':len(D.tables),'embedded_figures':len(D.inline_shapes),'pdf_images':images,'all_38_features_in_pdf':all(k in text for k in features),'all_six_screens_in_pdf':all('UI_'+str(i).zfill(2) in text for i in range(1,7)),'no_empty_body_pages':all(len(p.extract_text() or '')>250 for p in P.pages),'excluded_draft_not_referenced':all(x not in text for x in ['AI_관상_상세요구사항명세서_FaceTeller.docx','Word 초안','신병석']), 'unresolved_word_fields':any(s in text for s in ['오류!','Error! Reference','Error! Bookmark'])}
assert checks['ten_fr_forms'] and checks['all_38_features_in_pdf'] and checks['all_six_screens_in_pdf'] and checks['no_empty_body_pages'] and not checks['unresolved_word_fields']
assert checks['excluded_draft_not_referenced']
assert len(D.inline_shapes)==10 and len(images)==10
checks['visual_review']='전체 페이지 축소판 및 표지·구성도·DB표·실제 목업·요구사항 표 확대 검수. 최종 변경 페이지 재확인.'
(O/'문서검수결과.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
srcs=[Path('C:/Users/user/Documents/카카오톡 받은 파일/2026년 2학기 컴퓨터공학 종합설계_상세요구사항 명세서.pdf'),Path('C:/Users/user/Downloads/상세설계서_sample.pdf'),Path('C:/Users/user/Documents/카카오톡 받은 파일/페이스텔러 프로토타입.html'),R/'data/label_schema_kface_v3.json',R/'data/interpretation_service_v1.json',R/'outputs/experiment-report-2026-10-01/먼저읽기.md']
(O/'작성근거_SHA256.json').write_text(json.dumps([{'file':str(p),'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in srcs],ensure_ascii=False,indent=2),encoding='utf-8')
with (O/'요구사항_목록.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f);w.writerow(['ID','요구사항 명','중요도','INPUT','OUTPUT']);
 for r in json.loads((O/'문서구성.json').read_text(encoding='utf-8'))['requirements']:w.writerow([r['id'],r['name'],r['importance'],r['input'],r['output']])
(O/'파일안내.md').write_text(f'''# 페이스텔러 상세 요구사항 명세서 1.0

- **페이스텔러_상세요구사항명세서_1.0.pdf**: 검토용 최종본, 표지·목차 포함 {len(P.pages)}쪽.
- **페이스텔러_상세요구사항명세서_1.0.docx**: 편집용 Word. 표와 본문을 직접 수정할 수 있으며 목차는 Word에서 업데이트할 수 있습니다.
- **페이스텔러_상세요구사항명세서_1.0.md**: 검색·비교용 본문. 인쇄 형식은 PDF/Word를 기준으로 봅니다.
- **요구사항_목록.csv**: FR_001~FR_010의 입력·출력 요약.
- **문서검수결과.json / 작성근거_SHA256.json**: 문서 구조 확인과 사용한 자료의 식별 기록.

「2026년 2학기 컴퓨터공학 종합설계_상세요구사항 명세서.pdf」를 유일한 요구사항 기준(SSOT)으로 삼았습니다. Word 초안은 작성 근거에서 제외했습니다. 1~5장 구성과 기능 ID를 유지했습니다. 제공된 HTML의 6개 화면을 직접 캡처해 넣고, 화면별 버튼·입출력·실패 시 동작을 작성했습니다. 그림과 표를 포함하며 원본 자료는 수정하지 않았습니다.

목업의 고정값은 실제 실험 수치와 구분했습니다. 기존 프로토타입 실험, 새 구현의 설계 초기값, 원문·레이블 검수 및 장비 측정이 필요한 항목은 각각 표시했습니다. 본 작업은 명세서 작성이며 실제 서비스 구현이나 추가 모델 학습을 수행한 결과는 아닙니다.
''',encoding='utf-8')
print(json.dumps({k:v for k,v in checks.items() if k not in ['pdf_images','visual_review']},ensure_ascii=True))
