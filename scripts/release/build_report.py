"""CPU-only, fail-closed release report; never imports models or contacts the service.

Default execution requires the complete real 400-person evidence bundle. Synthetic
checks require --synthetic --no-copy and both root/output below the OS temp folder.
Only explicit aggregate fields reach the shared documents, never participant rows.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import datetime, timezone
import hashlib
import json
import math
import os
from pathlib import Path
import shutil
import statistics
import sys
import tempfile
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:  # direct CLI use imports scripts.release.sync_label_metadata
    sys.path.insert(0, str(ROOT))
TITLE = 'FaceTeller_전체구현_실험보고서'
SPLITS = {'train': 270, 'selection': 34, 'calibration': 34, 'test': 62}
SPLIT_KO = {'train': '학습', 'selection': '선택', 'calibration': '보정', 'test': '역사적 시험'}
ARCH = {'resnet18': 'ResNet-18', 'efficientnet_b0': 'EfficientNet-B0', 'convnext_tiny': 'ConvNeXt-Tiny'}
PART_KO = {'eye': '눈', 'brow': '눈썹', 'nose': '코', 'mouth': '입', 'face': '얼굴형',
           'forehead': '이마', 'glabella': '미간', 'under_eye': '눈 밑', 'cheekbone': '광대',
           'ear': '귀', 'philtrum': '인중', 'mole': '점'}
STAGES = ('validation_seconds', 'alignment_seconds', 'preprocessing_seconds', 'cnn_seconds', 'depth_seconds', 'literature_seconds')
STAGE_KO = dict(zip(STAGES, ('입력 검증', '정렬', '전처리', 'CNN 분류', 'VGGT 깊이', '문헌 조회')))
STAGE_KO['recorded_stage_sum'] = '기록된 단계 합'
STAGE_KO['received_to_result_seconds'] = '수신 완료→결과 준비'
REASON_KO = {
    'LOW_CONFIDENCE': '분류 신뢰도 기준 미달', 'LOW_OR_INVALID_CONFIDENCE': '분류 신뢰도 기준 미달/유효하지 않음',
    'CONFIDENCE_NOT_CALIBRATED': '신뢰도 보정 미완료', 'UNCALIBRATED': '미보정',
    'PREDICTED_CLASS_NOT_TRAINED': '예측 클래스 학습 표본 없음', 'CROP_UNAVAILABLE': '부위 영상 확보 불가',
    'CNN_HEAD_FAILED': '특성 분류 실패', 'CNN_PART_FAILED': '부위 분류 실패',
    'DEPTH_QUALITY_FAILED': '깊이 품질 기준 미달', 'DEPTH_CONFIDENCE_LOW': '깊이 출력 품질 낮음',
    'DEPTH_VIEWS_INCONSISTENT': '사진 간 깊이 측정 불일치', 'DEPTH_ESTIMATION_FAILED': '깊이 추정 실패',
    'DEPTH_EXPRESSION_CHANGED': '사진 간 표정 변화', 'DEPTH_VIEW_DIRECTIONS_INVALID': '사선 방향 부적합',
    'DEPTH_SCORE_UNAVAILABLE': '깊이 점수 없음', 'MEASUREMENT_UNAVAILABLE': '측정값 없음',
    'LABEL_UNAVAILABLE': '형태 레이블 없음', 'BASIS_UNAVAILABLE': '판정 근거 없음',
    'FACE_TOO_SMALL': '얼굴 크기 부족', 'FRONTAL_POSE_INVALID': '정면 방향 부적합',
    'EXPRESSION_SMILE': '웃는 표정', 'EXPRESSION_BLINK': '눈 감음/깜빡임', 'EXPRESSION_FROWN': '찡그린 표정',
    'MOUTH_OPEN': '입 벌림', 'ILLUMINATION_TOO_DARK': '밝기 부족', 'LANDMARKS_INVALID': '얼굴 위치 검출 불안정',
    'PARSING_FAILED': '부위 경계 검출 실패', 'EAR_VIEW_DIRECTIONS_INVALID': '양쪽 귀 사선 방향 부적합',
    'FOREHEAD_VISIBILITY_UNAVAILABLE': '이마 가시성 부족', 'HAIRLINE_NOT_SUFFICIENTLY_VISIBLE': '머리선 가시성 부족',
    'SKIN_MARK_VISIBILITY_UNAVAILABLE': '피부 표식 가시성 확인 불가', 'SKIN_MARK_VISIBILITY_INSUFFICIENT': '피부 표식 가시성 부족',
    # Service hold codes added after the first report; wording mirrors frontend/src/Results.tsx.
    'BOUNDARIES_NOT_VALIDATED': '측정 경계 검증 전', 'DEPTH_BOUNDARIES_UNVALIDATED': '깊이 분류 기준을 확인할 수 없음',
    'DEPTH_MEASUREMENT_FAILED': '깊이 측정 실패', 'DEPTH_SCORE_MISSING': '측정 적합도 없음',
    'EAR_GEOMETRY_IMPLAUSIBLE': '귀의 경계를 안정적으로 확인할 수 없음', 'EAR_NOT_VISIBLE': '귀가 충분히 보이지 않음',
    'EAR_VIEW_OUTSIDE_OBLIQUE_RANGE': '귀를 확인할 사선 각도가 맞지 않음', 'FOREHEAD_HEIGHT_IMPLAUSIBLE': '이마 높이를 안정적으로 측정할 수 없음',
    'FOREHEAD_SKIN_NOT_VISIBLE': '이마 피부가 충분히 보이지 않음', 'FOREHEAD_WIDTH_NOT_VISIBLE': '이마의 양쪽 경계가 보이지 않음',
    'HAIRLINE_IRREGULAR_OR_OCCLUDED': '머리선이 가려져 있거나 경계가 불명확함', 'HAT_OCCLUSION': '모자가 이마를 가리고 있음',
    'ILLUMINATION_TOO_BRIGHT': '얼굴이 너무 밝게 찍힘', 'INSUFFICIENT_CLEAR_SKIN_COVERAGE': '가리지 않은 피부 영역이 부족함',
    'LIGHTING_OR_RESOLUTION_UNSUITABLE': '피부 표시를 확인할 조명이나 해상도가 부족함', 'LIGHTING_UNSTABLE': '조명에 따라 결과가 달라짐',
    'LOW_HAIR_EDGE_OR_BANGS_UNCERTAIN': '앞머리와 머리선을 구분하기 어려움', 'MODEL_NOT_TRAINED': '해당 항목의 모델 준비 전',
    'QUALITY_FAILED': '입력 품질 기준 미충족', 'REGRESSION_HEAD_HAS_NO_CLASS_CONFIDENCE': '분류 신뢰도가 없는 측정 항목',
    'TEXTURE_OR_SHADOW_UNCERTAIN': '피부 표시와 그림자를 구분하기 어려움',
}
METRICS = ('accuracy', 'macro_f1', 'balanced_accuracy', 'nll', 'ece', 'brier')
SCORE_PROOF = ('score_formula_verified', 'score_formula_verified_count', 'score_components_verified_count', 'selection_verified')
SOURCES = [
    ('VGGT 공식 저장소: 원 연구 체크포인트는 비상업용', 'https://github.com/facebookresearch/vggt'),
    ('원 VGGT-1B 모델 카드', 'https://huggingface.co/facebook/VGGT-1B'),
    ('Guo 등, 2017: 온도 보정', 'https://proceedings.mlr.press/v70/guo17a.html'),
    ('PyTorch 공식 전이학습 안내', 'https://docs.pytorch.org/tutorials/beginner/transfer_learning_tutorial.html'),
]
LIMITS = [
    '분류 수치는 AI 합의 또는 자동 형태 측정으로 만든 참조 레이블과의 일치도다. 독립적인 사람·전문가 정답에 대한 정확도가 아니다. 성격·지능·건강·운명 예측을 검증하지 않는다.',
    '시험 62명은 앞선 개발에서 사용된 역사적 시험 집합이다. 신규 외부 시험이 아니며 400명 중 학습·선택·보정 집합의 결과를 독립 시험 성능으로 합치지 않는다.',
    'ImageNet 사전학습은 일반 영상 표현을 제공한다. 관상을 이미 아는 모델이 아니며 부위별 관찰 가능한 외형 레이블에 맞춰 미세조정했다.',
    '기존 45회 비교는 3계열×5부위×3시드, 기존 15헤드와 전처리의 선택 실험이다. 이번 배포 모델은 확대된 레이블과 눈 중심 정규화 등 조건이 달라 직접적인 성능 향상·계열 우열의 재검증으로 해석할 수 없다.',
    'macro F1은 많은 유형에만 치우치지 않도록 유형별 분류성능을 평균한 지표, accuracy는 맞힌 비율이다. F1의 평균 클래스는 참조 또는 예측에 등장한 유형의 합집합이며 참조 지지수 0인 유형의 재현율은 평가할 수 없다.',
    '헤드별 평균은 평가 가능한 헤드의 동일 가중 평균이다. 부위별 평균도 같은 방식이며 사람별 모든 특성을 동시에 맞힌 비율이 아니다. 보정 전후 비교는 두 값이 모두 있는 동일 헤드에서만 계산한다.',
    '선택 34명으로 체크포인트를 고정한 뒤 보정 34명으로 배포 조합의 헤드별 온도를 맞춘다. 양의 온도 보정은 분류 argmax를 바꾸지 않는다. NLL·ECE는 낮을수록 좋지만 시험에서의 개선은 보장되지 않는다.',
    '깊이 경계는 400명 전체의 기술적 기준 분포다. 반복 촬영의 산포와 품질 기준은 학습군에서 정한다. 독립적인 실제 3차원 측정 정답이나 깊이 분류 정확도 검증이 아니다.',
    '문헌 순위 G는 CNN 보정 신뢰도와 깊이 경계점수를 포함한 동일 가중 기하평균이다. 종합 일치점수와 경계점수는 정답확률이 아니다. 원문·번역·해설의 AI 주석은 인간 검수와 구분한다.',
    '원 연구용 facebook/VGGT-1B 가중치는 비상업용 제한이 남아 있다. 별도 VGGT-1B-Commercial의 허용 범위를 이 체크포인트에 적용하지 않으며 상업 서비스 준비 완료로 표현하지 않는다.',
    '보고서는 가중치를 로드하거나 자동 배포하지 않는다. 자료에 없는 실제 웹캠·기기별 촬영·외부 모집단·모든 문헌 조건을 검증했다고 주장하지 않는다. 실제 웹캠과 장치별 촬영 검증은 별도로 남아 있다.',
]


class EvidenceError(ValueError):
    pass


def require(ok, message):
    if not ok:
        raise EvidenceError(message)


def numeric(value, name, minimum=0, nullable=False):
    if nullable and value is None:
        return None
    require(type(value) in (int, float) and math.isfinite(value) and value >= minimum,
            f'{name}: 유효한 수치가 필요합니다')
    return value


def integer(value, name, minimum=0):
    require(type(value) is int and value >= minimum, f'{name}: 정수가 필요합니다')
    return value


def read(path):
    require(path.is_file(), f'필수 자료 없음: {path}')
    try:
        return json.loads(path.read_text(encoding='utf-8-sig'))
    except (ValueError, UnicodeError) as exc:
        raise EvidenceError(f'UTF-8 JSON 읽기 실패: {path.name}') from exc


def sha(path):
    digest = hashlib.sha256()
    with path.open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()


def canonical_sha(document):
    encoded = json.dumps(document, ensure_ascii=False, sort_keys=True, separators=(',', ':'), allow_nan=False)
    return hashlib.sha256(encoded.encode('utf-8')).hexdigest()


EAR_TIME_SCOPE = ('학습 시간은 최종 채택 모델 metadata만 합산한다. 폐기된 이전 귀 3모델의 학습과 '
                  '이전 보정·시험은 제외하며, 최종 보정·시험도 모델 학습합과 별도로 표시한다. '
                  '입력 수정 전체 경과와 각 단계 비용은 중첩될 수 있어 합산하지 않는다.')
EAR_COSTS = (
    ('baseline_ear_training_seconds', '폐기된 이전 귀 3모델 학습'),
    ('baseline_calibration_seconds', '이전 배포 보정'),
    ('baseline_test_seconds', '이전 배포 시험'),
    ('correction_wall_seconds', '입력 수정 전체 경과'),
)


def is_sha256(value):
    return isinstance(value, str) and len(value) == 64 and all(c in '0123456789abcdef' for c in value)


def validate_ear_correction(proof, deployment, parts, candidates, active):
    """The correction is historical (3 candidates); the active deployment may use a subset of it."""
    require(isinstance(proof, dict) and proof.get('pass') is True, '귀 입력 수정 검증 실패')
    require(integer(proof.get('unchanged_non_ear_models'), '유지 모델 수') == 33 and
            integer(proof.get('retrained_ear_models'), '귀 재학습 수') == 3, '귀 수정의 33+3 모델 범위 불일치')
    version = proof.get('ear_transform_version')
    require(isinstance(version, str) and version and version == deployment.get('ear_transform_version'),
            '귀 수정/배포 전처리 버전 불일치')
    if proof.get('ear_preprocessing_sha256') is not None or deployment.get('ear_preprocessing_sha256') is not None:
        require(proof.get('ear_preprocessing_sha256') == deployment.get('ear_preprocessing_sha256'),
                '귀 수정/배포 귀 전처리 SHA256 불일치')
    for key, _ in EAR_COSTS:
        numeric(proof.get(key), key)
    def weights(value):
        require(isinstance(value, dict), '귀 수정 가중치 해시 목록 누락')
        normalized = {}
        for key, digest in value.items():
            require(isinstance(key, str), '귀 수정 가중치 경로 오류')
            key = key.replace('\\', '/')
            require(key not in normalized, '귀 수정 가중치 경로 별칭 중복')
            require(is_sha256(digest), '귀 수정 가중치 SHA256 오류')
            normalized[key] = digest
        return normalized
    before = weights(proof.get('before_non_ear_weights'))
    after = weights(proof.get('after_non_ear_weights'))
    history = {f'models/release/{arch}/{part}.pt' for arch in candidates for part in parts if part != 'ear'}
    require(len(candidates) == 3 and before == after and set(after) == history and len(history) == 33,
            '귀 수정 전후의 역사적 3계열 비귀 33모델 해시 불일치')
    deployed = weights(deployment.get('weights'))
    non_ear = {k: v for k, v in deployed.items() if not k.endswith('/ear.pt')}
    ear = set(deployed) - set(non_ear)
    require(set(active) <= set(candidates) and ear == {f'models/release/{arch}/ear.pt' for arch in active},
            '현재 배포 귀 모델이 수정 대상 계열과 다릅니다')
    require(non_ear and all(after.get(k) == v for k, v in non_ear.items()),
            '현재 배포 비귀 모델 해시가 귀 수정 기록의 유지 모델과 다릅니다')
    return {'history_non_ear': len(after), 'history_ear': proof['retrained_ear_models'],
            'active_non_ear': len(non_ear), 'active_ear': len(ear)}


def ear_correction_note(b):
    proof = b.get('ear_correction')
    basis = '기존 귀 bbox 확대·축소가 길이 기준을 지워, 눈 사이 거리·yaw 기준의 고정 창으로 수정했다. '
    if proof is None:
        return basis + '귀 3모델만 재학습하는 수정 검증 자료는 미제공/대기이며 완료·비용을 추정하지 않는다.'
    costs = ' / '.join(f'{label} {proof[key]/60:.2f}분' for key, label in EAR_COSTS)
    scope = b['ear_correction_scope']
    text = basis + f"역사적 수정 기록: 3계열 비귀 33모델 해시 유지·귀 3모델 재학습 검증 통과. 전처리 {proof['ear_transform_version']}. " + costs + '.'
    if scope['active_non_ear'] + scope['active_ear'] < scope['history_non_ear'] + scope['history_ear']:
        text += (f" 현재 배포는 그중 비귀 {scope['active_non_ear']}모델(수정 기록의 해시와 동일)과 수정 귀 {scope['active_ear']}모델만 사용한다. "
                 '이전 배포 보정·시험 비용은 당시 3계열 배포의 기록이며 현재 단일 계열 보정·시험 시간과 별개다.')
    return text


INFERENCE_VERSION = 'release-inference-choice-v1'


def validate_inference_selection(root, deployment, config, hashes, paths):
    """Bind the deployed classifier subset without rewriting the historical training config.

    A legacy manifest without selection fields deploys every trained candidate. A selection
    manifest must match models/release/inference_frozen.json byte-for-byte and that file must
    bind the unchanged training_frozen.json and registry.
    """
    fields = ('inference_architectures', 'inference_selection_sha256')
    path = root / 'models/release/inference_frozen.json'
    if not any(key in deployment for key in fields):
        require(not path.exists(), 'inference_frozen.json이 있으나 배포에 추론 선택 신원이 없습니다: 재보정 필요')
        return list(config['architectures']), None
    require(all(key in deployment for key in fields), '배포 추론 선택 필드 일부 누락')
    training_path = root / 'models/release/training_frozen.json'
    selection, training = read(path), read(training_path)
    for key, item in (('inference_selection', path), ('training_config', training_path)):
        paths[key] = item
        hashes[key] = sha(item)
    require(isinstance(selection, dict) and isinstance(training, dict), '추론 선택/학습 설정 객체 필요')
    require(deployment['inference_selection_sha256'] == hashes['inference_selection'],
            '배포/추론 선택 파일 SHA256 불일치')
    require(selection.get('version') == INFERENCE_VERSION, '추론 선택 schema 불일치')
    architectures = selection.get('architectures')
    require(architectures == ['resnet18'] == deployment['inference_architectures'],
            '현재 추론은 ResNet-18 단일 계열이어야 합니다(앙상블 금지)')
    require(set(architectures) <= set(config['architectures']), '추론 계열이 원 학습 후보에 없습니다')
    require(selection.get('registry_sha256') == hashes['registry'] == config.get('registry_sha256'),
            '추론 선택/registry SHA256 불일치')
    require(selection.get('training_config_sha256') == hashes['training_config'],
            '추론 선택/학습 설정 파일 SHA256 불일치')
    require(training == config, '배포 config가 원 학습 설정(training_frozen.json)과 다릅니다')
    basis = selection.get('selection_basis')
    require((isinstance(basis, str) and basis.strip()) or (isinstance(basis, dict) and basis), '추론 선택 근거 누락')
    return list(architectures), selection


THRESHOLD_VERSION = 'validation-reject-cost-v1'


def validate_threshold_policy(root, deployment, hashes, paths):
    """Bind per-head reject thresholds to the frozen calibration-only policy, or None for legacy fixed .6."""
    path = root / 'models/release/threshold_policy_frozen.json'
    if 'threshold_policy_sha256' not in deployment:
        require(not path.exists(), 'threshold_policy_frozen.json이 있으나 배포에 기준값 정책 해시가 없습니다: 재보정 필요')
        require(not any('threshold_selection' in c for c in deployment.get('features', {}).values()),
                '기준값 선택 기록이 있으나 정책 해시가 없습니다')
        return None
    policy = read(path)
    paths['threshold_policy'] = path
    hashes['threshold_policy'] = sha(path)
    require(deployment['threshold_policy_sha256'] == hashes['threshold_policy'], '배포/기준값 정책 파일 SHA256 불일치')
    require(isinstance(policy, dict) and policy.get('version') == THRESHOLD_VERSION, '기준값 정책 schema 불일치')
    require(policy.get('split') == 'calibration' and policy.get('test_used') is False, '기준값 정책이 보정군 전용이 아닙니다')
    for key in ('wrong_confirmation_cost', 'withheld_cost', 'initial_threshold'):
        numeric(policy.get(key), f'threshold policy {key}')
    require(policy['wrong_confirmation_cost'] > 0 and policy['withheld_cost'] > 0, '기준값 비용은 양수여야 합니다')
    require(policy.get('binary_floor') == .6 and policy.get('multiclass_floor') == .5 and policy['initial_threshold'] == .6,
            '설계 하한(2클래스 .60/다중 .50)·초기값 .60 불일치')
    grid = policy.get('candidates')
    require(isinstance(grid, list) and grid and grid == sorted(set(grid)) and all(numeric(v, 'candidate') <= 1 for v in grid),
            '기준값 후보 격자 오류')
    return policy


def validate_threshold_selection(key, cal, classes, policy):
    """Recompute the calibration-only choice from the recorded candidate counts."""
    choice = cal.get('threshold_selection')
    require(isinstance(choice, dict), f'{key}: 기준값 선택 기록 누락')
    require(choice.get('policy_version') == policy['version'] and choice.get('source_split') == 'calibration'
            and choice.get('test_used') is False, f'{key}: 기준값 선택 출처 불일치')
    require(choice.get('threshold') == cal.get('threshold'), f'{key}: 배포 기준값과 선택 기록 불일치')
    n = integer(choice.get('n_valid'), f'{key}: 기준값 보정 유효 수')
    require(n == cal['n_val'], f'{key}: 기준값/온도 보정 표본 수 불일치')
    if not cal['calibrated'] or n == 0:
        status = 'disabled_uncalibrated' if not cal['calibrated'] else 'no_valid_calibration_labels'
        require(choice.get('status') == status and choice.get('candidates') == [] and
                choice['threshold'] == policy['initial_threshold'], f'{key}: 미선택 기준값 상태 불일치')
        return
    floor = policy['binary_floor'] if classes == 2 else policy['multiclass_floor']
    candidates = choice.get('candidates')
    require(choice.get('status') == 'selected_on_calibration' and isinstance(candidates, list) and
            [c.get('threshold') for c in candidates] == [v for v in policy['candidates'] if v >= floor],
            f'{key}: 기준값 후보가 정책 격자/하한과 다릅니다')
    for c in candidates:
        accepted, held, wrong = (integer(c.get(k), f'{key}: {k}') for k in ('accepted', 'withheld', 'wrong_confirmed'))
        require(accepted + held == n and wrong <= accepted, f'{key}: 기준값 후보 집계 불일치')
        loss = (wrong * policy['wrong_confirmation_cost'] + held * policy['withheld_cost']) / n
        require(math.isclose(numeric(c.get('loss'), 'loss'), loss, abs_tol=1e-12) and
                math.isclose(numeric(c.get('coverage'), 'coverage'), accepted / n, abs_tol=1e-12) and
                (c.get('accepted_error_rate') is None if accepted == 0 else
                 math.isclose(numeric(c.get('accepted_error_rate'), 'error rate'), wrong / accepted, abs_tol=1e-12)),
                f'{key}: 기준값 후보 loss/coverage 재계산 불일치')
    best = min(candidates, key=lambda c: (c['loss'], -c['accepted'], abs(c['threshold'] - policy['initial_threshold']), c['threshold']))
    require(choice.get('selected') == best and choice['threshold'] == best['threshold'], f'{key}: 정책 규칙으로 재계산한 기준값과 다릅니다')


def avg(values):
    values = [v for v in values if v is not None]
    return statistics.mean(values) if values else None


def percentile(values, fraction=.95):
    if not values:
        return None
    values = sorted(values)
    position = (len(values) - 1) * fraction
    lower = int(position)
    upper = min(lower + 1, len(values) - 1)
    return values[lower] + (values[upper] - values[lower]) * (position - lower)


def summary(values):
    return [len(values), avg(values), percentile(values), max(values) if values else None,
            sum(v > 5 for v in values)]


def fmt(value):
    if value is None:
        return '미기록/평가불가'
    if isinstance(value, bool):
        return '예' if value else '아니오'
    if isinstance(value, float):
        return f'{value:.4f}'
    return str(value)


def validate_metrics(m, classes, name):
    require(isinstance(m, dict), f'{name}: metrics 누락')
    for key in ('total_count', 'valid_count', 'missing_count', 'class_count'):
        integer(m.get(key), f'{name}.{key}')
    require(m['class_count'] == classes, f'{name}: 클래스 개수 불일치')
    require(m['total_count'] == m['valid_count'] + m['missing_count'], f'{name}: 평가 분모 불일치')
    for key in METRICS:
        numeric(m.get(key), f'{name}.{key}', nullable=True)
        require(m['valid_count'] == 0 or m.get(key) is not None, f'{name}.{key}: 평가 수치 누락')
        if key in ('accuracy', 'macro_f1', 'balanced_accuracy', 'ece') and m.get(key) is not None:
            require(m[key] <= 1, f'{name}.{key}: 범위 오류')
    for key in ('support', 'predicted_support'):
        require(isinstance(m.get(key), list) and len(m[key]) == classes, f'{name}.{key}: 길이 불일치')
        for value in m[key]:
            integer(value, f'{name}.{key}')
        require(sum(m[key]) == m['valid_count'], f'{name}.{key}: 합계 불일치')
    require(m.get('macro_f1_classes') == [i for i in range(classes) if m['support'][i] or m['predicted_support'][i]],
            f'{name}: macroF1 truth/pred 합집합 정책 불일치')
    threshold = m.get('threshold', {})
    for key in ('value', 'coverage', 'accuracy'):
        numeric(threshold.get(key), f'{name}.threshold.{key}', nullable=key != 'value')
    integer(threshold.get('accepted_count'), f'{name}.accepted_count')
    require(isinstance(m.get('ece_bins'), list), f'{name}: ECE bin 누락')


def validate_ui_evidence(ui, hashes, corpus_digest):
    require(ui.get('pass') is True and ui.get('actual_models') is True and ui.get('mocked') is False
            and ui.get('artifact_identity_verified') is True, '실제 브라우저 모델·응답 검증 증거 누락')
    for key, value in {'deployment_sha256': hashes['deployment'], 'depth_config_sha256': hashes['depth'],
                       'registry_sha256': hashes['registry'], 'corpus_sha256': corpus_digest}.items():
        require(ui.get(key) == value, '실제 브라우저 검증의 배포/문헌 신원이 다릅니다')
    require(ui.get('participants') == 1 and ui.get('split') == 'train', '브라우저 검증 범위 불일치')
    require(ui.get('photos_saved') is False and ui.get('external_requests') == 0 and ui.get('page_errors') == [],
            '브라우저 사진·외부요청·오류 확인 실패')
    integer(ui.get('confirmed'), 'UI confirmed', minimum=1)
    integer(ui.get('held'), 'UI held')
    require(type(ui.get('one_page_button')) is bool, 'UI A4 인쇄 상태는 boolean이어야 합니다')
    integer(ui.get('passages'), 'UI passage count')
    require(ui['passages'] <= 5, 'UI 최대 문헌 수 초과')
    require(ui.get('pdf_passage_text_verified_count') == ui['passages'], 'PDF 전문 검증 조항 수 불일치')
    integer(ui.get('pdf_pages'), 'UI PDF pages', minimum=1)
    require(ui.get('one_page_button') is not True or ui['pdf_pages'] == 1, 'UI A4 한 장 표시 불일치')


def validate_tls_camera(doc, b, frontend_copy=None):
    """Supplemental recorded-camera/TLS run: identity-bound, never a physical webcam claim."""
    require(doc.get('pass') is True and doc.get('actual_models') is True, 'TLS/카메라 검증 미통과')
    require(doc.get('certificate_verified_by_scoped_ca') is True and doc.get('global_trust_modified') is False,
            'TLS 인증서 검증 범위 불일치(전역 신뢰 변경 금지)')
    require(type(doc.get('physical_camera_tested')) is bool, '물리 웹캠 시험 여부 누락')
    health, camera = doc.get('health'), doc.get('camera')
    require(isinstance(health, dict) and isinstance(camera, dict), 'TLS health/카메라 기록 누락')
    require(health.get('ready') is True and health.get('deployment_sha256') == b['hashes']['deployment'] and
            health.get('registry_sha256') == b['hashes']['registry'] and health.get('depth_config_sha256') == b['hashes']['depth'] and
            health.get('corpus_sha256') == b['corpus_canonical_sha256'] and health.get('classifier_architectures') == b['architectures'] and
            health.get('classifier_model_count') == len(b['deployment']['weights']), 'TLS 서비스 신원이 현재 배포/문헌과 다릅니다')
    require(camera.get('pass') is True and camera.get('actual_models') is True and camera.get('detector_mocked') is False and
            camera.get('photos_saved') is False and camera.get('external_requests') == 0 and camera.get('page_errors') == [] and
            camera.get('tracks_stopped') is True and camera.get('deployment_sha256') == b['hashes']['deployment'],
            '카메라 경로 검증 조건 불일치')
    require(camera.get('physical_camera') is doc['physical_camera_tested'] and type(camera.get('recorded_source')) is bool,
            '카메라 입력 출처 표기 불일치')
    require(camera.get('quality_version') == health.get('quality_version') and isinstance(health.get('quality_version'), str),
            '카메라/서비스 품질 게이트 버전 불일치')
    service_quality = b.get('service', {}).get('service', {}).get('quality_version')
    require(service_quality is None or service_quality == health['quality_version'], 'TLS와 HTTP400의 품질 게이트 버전이 다릅니다')
    integer(camera.get('capture_posts'), 'capture_posts')
    integer(camera.get('equivalent_item_count'), 'equivalent_item_count')
    numeric(camera.get('max_score_difference'), 'max_score_difference')
    if frontend_copy is not None:
        require(frontend_copy == camera, 'frontend-camera-verification.json과 TLS 카메라 기록이 다릅니다')


def validate_db_metadata(check, sync, final, b):
    """Production migration004/label sync evidence must match the frozen definitions and corpus."""
    from scripts.release.sync_label_metadata import load_inputs
    rows, revision, inputs = load_inputs({'training_registry': b['paths']['registry'], 'depth_config': b['paths']['depth'],
                                          'landmark_labels': b['root'] / 'data/release/landmark_labels.json'})
    require(check.get('matches') is True and check.get('differing_labels') == [] and check.get('eligible_without_metadata') == [] and
            check.get('metadata_revision') == revision == sync.get('metadata_revision'), 'DB 레이블 메타데이터가 동결 정의와 다릅니다')
    require(sync.get('inputs_sha256') == inputs and sync.get('registry_labels') == len(rows) and sync.get('unmapped_db_labels') == [],
            'DB 레이블 동기화 입력/범위 불일치')
    eligible = sum(p['service_eligible'] for p in b['corpus']['passages'])
    require(final.get('corpus_sha256') == b['corpus_canonical_sha256'] and final.get('eligible') == eligible and
            final.get('passages') == len(b['corpus']['passages']) and final.get('labels_with_metadata') == len(rows) and
            final.get('labels') == len(rows), 'DB 최종 상태가 현재 corpus/레이블 메타데이터와 다릅니다')
    return {'metadata_revision': revision, 'labels': len(rows), 'eligible': eligible}


def load_bundle(root, corpus_arg=None, synthetic=False):
    paths = {
        'registry': root / 'data/release/training_registry.json',
        'deployment': root / 'models/release/deployment.json',
        'test': root / 'outputs/release/test_metrics.json',
        'depth': root / 'data/release/depth_config.json',
        'service': root / 'outputs/release/service-verification.json',
        'corpus': Path(corpus_arg) if corpus_arg else root / 'data/release/literature/reference_corpus_v1.json',
    }
    # Check the whole required list before processing or creating any output.
    missing = [str(p) for p in paths.values() if not p.is_file()]
    require(not missing, '필수 자료 미완료: ' + ', '.join(missing))
    bundle = {key: read(path) for key, path in paths.items()}
    for key, doc in bundle.items():
        require(isinstance(doc, dict), f'{key}: 객체 필요')
        require(bool(doc.get('synthetic', False)) == synthetic, f'{key}: 합성/실측 표식 불일치')
    r, d, t, depth, service, corpus = [bundle[k] for k in paths]
    hashes = {key: sha(path) for key, path in paths.items()}
    require(r.get('split_sizes') == SPLITS, '분할이 270/34/34/62와 다릅니다')
    subjects = r.get('subjects', {})
    require(len(subjects) == 400, 'registry 인물 수가 400이 아닙니다')
    require(Counter(s.get('split') for s in subjects.values()) == Counter(SPLITS), 'registry 분할 합계 불일치')
    features, parts = r.get('features', {}), r.get('parts', {})
    require(features and parts, 'registry features/parts 누락')
    bundle['stage_timings'] = {}
    for stage, timing in [('보정34', d.get('calibration_timings')), ('시험62', t.get('timings'))]:
        if timing is None:
            require(synthetic, f'{stage}: 실제 보정·시험 시간 기록 누락')
            continue
        require(isinstance(timing, dict) and set(timing.get('parts', {})) == set(parts), f'{stage}: 부위별 시간 누락')
        total = numeric(timing.get('total_seconds'), stage + ' total seconds')
        seconds = {part:numeric(value, stage + '/' + part) for part,value in timing['parts'].items()}
        require(total > 0 and sum(seconds.values()) <= total + 1e-6, f'{stage}: 전체/부위 시간 불일치')
        require(isinstance(timing.get('scope'), str) and timing['scope'], f'{stage}: 시간 측정 범위 누락')
        bundle['stage_timings'][stage] = {'total_seconds':total, 'parts':seconds, 'scope':timing['scope']}
    assigned = []
    for part, spec in parts.items():
        require(spec.get('features'), f'{part}: 학습 헤드 없는 부위')
        assigned.extend(spec['features'])
        for key in spec['features']:
            require(key in features and features[key]['part'] == part, '부위/특성 배정 불일치')
    require(len(assigned) == len(set(assigned)) and set(assigned) == set(features), '부위/헤드 집합 불일치')
    for key, feature in features.items():
        require(feature.get('ko') and feature.get('classes'), f'{key}: 특성 이름/클래스 누락')
        for split, size in SPLITS.items():
            support = r.get('class_support', {}).get(split, {}).get(key, {})
            counts = support.get('counts')
            require(isinstance(counts, list) and len(counts) == len(feature['classes']), f'{key}/{split}: support 누락')
            for n in counts:
                integer(n, f'{key}/{split} support')
            integer(support.get('unknown'), f'{key}/{split} unknown')
            require(sum(counts) + support['unknown'] == size, f'{key}/{split}: support 합계 불일치')
    config = d.get('config', {})
    require(config.get('registry_sha256') == hashes['registry'], 'deployment/registry SHA256 불일치')
    require(config.get('split_sizes') == SPLITS, '배포 분할 불일치')
    require(config.get('preprocess_version') == r.get('preprocess_version'), '전처리 버전 불일치')
    require(d.get('test_used') is False, '배포 고정 시 test 사용 여부 확인 불가')
    require(t.get('deployment_sha256') == hashes['deployment'], '시험/실제 배포 SHA256 불일치')
    require(set(d.get('features', {})) == set(features) == set(t.get('features', {})), '배포/시험 전체 헤드 집합 불일치')
    for name, doc in (('registry', r), ('deployment config', config), ('test', t)):
        require(doc.get('human_review_count') == 0, f'{name}: human_review_count 0이 아님(인간 검수 기록 없음과 불일치)')
    candidates = config.get('architectures')
    require(isinstance(candidates, list) and candidates and len(set(candidates)) == len(candidates), '학습 후보 계열 누락/중복')
    architectures, selection = validate_inference_selection(root, d, config, hashes, paths)
    threshold_policy = validate_threshold_policy(root, d, hashes, paths)
    training, historical_training = [], []
    require(isinstance(d.get('weights'), dict), '배포 weights 누락')
    weight_paths = {str(k).replace('\\', '/'): v for k, v in d['weights'].items()}
    expected_weights = {f'models/release/{arch}/{part}.pt' for arch in architectures for part in parts}
    require(set(weight_paths) == expected_weights and len(weight_paths) == len(d['weights']), '배포 weights와 부위/계열 조합 불일치')
    require(all(is_sha256(v) for v in weight_paths.values()), '배포 weights SHA256 형식 오류')
    identity = service.get('service', {})
    if selection is not None or 'classifier_architectures' in identity:
        require(identity.get('classifier_architectures') == architectures and
                identity.get('classifier_model_count') == len(expected_weights), '실서비스 분류 모델 계열/수 불일치')
    for arch in candidates:
        for part, spec in parts.items():
            require('/' not in arch and '\\' not in arch and '/' not in part and '\\' not in part, '잘못된 모델 경로')
            path = root / 'models/release' / arch / f'{part}.json'
            model = read(path)
            require(model.get('config') == config, f'{arch}/{part}: 학습 설정 불일치')
            meta = model.get('meta', {})
            require(meta.get('arch') == arch and meta.get('part') == part, '학습 메타데이터 모델 불일치')
            require(meta.get('registry_sha256') == hashes['registry'], '학습/registry 해시 불일치')
            require(meta.get('classification_features') == spec['features'], f'{arch}/{part}: 헤드 순서 불일치')
            require(meta.get('preprocess_version') == config['preprocess_version'], '학습 전처리 불일치')
            if part == 'ear' and d.get('ear_transform_version') is not None:
                require(meta.get('ear_transform_version') == d['ear_transform_version'], f'{arch}/ear: 귀 입력 전처리 버전 불일치')
            require(meta.get('input_size') == spec.get('input_size') and meta.get('crops') == spec.get('crops'), '학습 crop 설정 불일치')
            for key in spec['features']:
                require(meta.get('classes', {}).get(key) == features[key]['classes'], f'{key}: 클래스 순서 불일치')
                integer(model.get('train_people_by_feature', {}).get(key), f'{key}: 학습 인물 수')
                validate_metrics(model.get('selection', {}).get(key, {}).get('metrics'), len(features[key]['classes']), f'{key}/selection')
            hist = model.get('history')
            require(isinstance(hist, list) and hist, f'{arch}/{part}: 학습 이력 없음')
            require([h.get('epoch') for h in hist] == list(range(1, len(hist) + 1)), '학습 epoch 순서 불일치')
            integer(config.get('epochs_max'), 'epochs_max', 1)
            require(len(hist) <= config['epochs_max'], '학습 budget 초과')
            selected = integer(model.get('selected_epoch'), 'selected_epoch', 1)
            require(selected <= len(hist), '선택 epoch가 실행 이력 밖입니다')
            for h in hist:
                numeric(h.get('train_seconds'), 'train_seconds')
                numeric(h.get('selection_macro_f1'), 'selection_macro_f1', nullable=True)
                if 'validation_seconds' in h:
                    numeric(h['validation_seconds'], 'validation_seconds')
            numeric(model.get('elapsed_seconds'), 'elapsed_seconds')
            integer(model.get('train_people_with_crops'), 'train_people_with_crops')
            integer(model.get('train_crops'), 'train_crops')
            require(model['train_people_with_crops'] <= 270, '학습 인물 수 초과')
            require(all(n <= model['train_people_with_crops'] for n in model['train_people_by_feature'].values()), '특성 학습 인물 수 초과')
            # Non-deployed candidates stay in the report as unchanged training history.
            label = 'train' if arch in architectures else 'history'
            (training if arch in architectures else historical_training).append(model)
            paths[f'{label}/{arch}/{part}'] = path
            hashes[f'{label}/{arch}/{part}'] = sha(path)
    calibrated_count, uncalibrated = 0, set()
    for key, feature in features.items():
        row, cal = t['features'][key], d['features'][key]
        require(row.get('participants') == 62, f'{key}: 시험 인물 수 불일치')
        integer(row.get('photo_available'), f'{key}: 사진 가용 수')
        require(row['photo_available'] <= 62 and row.get('annotation_provenance'), f'{key}: 사진/주석 근거 누락')
        validate_metrics(row.get('raw'), len(feature['classes']), f'{key}/test/raw')
        require(row['raw']['total_count'] == 62, f'{key}: 시험 평가 분모 불일치')
        require(row['raw']['valid_count'] <= row['photo_available'], f'{key}: 사진보다 평가 인물 수가 많습니다')
        require(all(a <= z for a, z in zip(row['raw']['support'], r['class_support']['test'][key]['counts'])), f'{key}: 평가 지지수가 registry보다 큽니다')
        require(cal.get('test_used_for_fit') is False, f'{key}: 보정 test 혼입 확인 불가')
        require(isinstance(cal.get('calibrated'), bool) and cal.get('status'), f'{key}: 보정 상태 누락')
        train_support = cal.get('train_support')
        require(isinstance(train_support, list) and len(train_support) == len(feature['classes']), f'{key}: 배포 유효 학습 지지수 누락')
        for eligible, reference in zip(train_support, r['class_support']['train'][key]['counts']):
            integer(eligible, f'{key}: 유효 학습 클래스 지지수')
            require(eligible <= reference, f'{key}: 유효 학습 지지수가 원본 참조 지지수를 초과합니다')
        validate_metrics(cal.get('raw'), len(feature['classes']), f'{key}/cal/raw')
        require(cal['raw']['total_count'] == 34 and cal.get('n_val') == cal['raw']['valid_count'], f'{key}: 보정 평가 분모 불일치')
        # Single-class calibration data cannot authorize a confidence threshold.
        observed = sum(n > 0 for n in cal.get('support', cal['raw']['support']))
        if cal.get('support') is not None:
            require(cal['support'] == cal['raw']['support'], f'{key}: 보정 지지수 불일치')
        if cal['status'] == 'single_class_calibration_support':
            require(cal['calibrated'] is False and cal.get('temperature') is None and observed <= 1,
                    f'{key}: 단일 클래스 보정 보류 상태 불일치')
        if threshold_policy is not None:
            validate_threshold_selection(key, cal, len(feature['classes']), threshold_policy)
            for name, m in (('보정 raw', cal['raw']), ('보정 T', cal.get('calibrated_metrics')),
                            ('시험 raw', row['raw']), ('시험 T', row.get('calibrated'))):
                require(m is None or (m.get('threshold') or {}).get('value') == cal['threshold'], f'{key}: {name} 지표가 선택 기준값을 쓰지 않았습니다')
        if cal['calibrated']:
            require(observed >= 2, f'{key}: 관측 클래스 1개 이하의 보정으로 신뢰도를 승인할 수 없습니다')
            calibrated_count += 1
        else:
            uncalibrated.add(key)
        if cal['calibrated']:
            require(numeric(cal.get('temperature'), f'{key}: 온도') > 0, '온도는 양수여야 합니다')
            validate_metrics(cal.get('calibrated_metrics'), len(feature['classes']), f'{key}/cal/T')
            validate_metrics(row.get('calibrated'), len(feature['classes']), f'{key}/test/T')
            require(row['calibrated']['support'] == row['raw']['support'], f'{key}: 보정 비교 분모 불일치')
        else:
            require(row.get('calibrated') is None, f'{key}: 미보정 헤드의 보정 수치가 존재합니다')
    require(depth.get('cohort_n') == 400 and depth.get('boundary_scope'), '깊이 400명 경계 근거 없음')
    require(set(depth.get('features', {})) == set(r.get('depth_features', {})) and depth['features'], '깊이 특성 집합 불일치')
    for key, f in depth['features'].items():
        require(f.get('n') == 400 and f.get('quantiles') == [.25, .75], '깊이 경계 표본/분위수 불일치')
        numeric(f.get('lower'), f'{key}.lower', minimum=-math.inf)
        numeric(f.get('upper'), f'{key}.upper', minimum=-math.inf)
        require(f['lower'] < f['upper'], '깊이 경계 순서 불일치')
        require(numeric(f.get('sigma_repeat_train'), 'sigma_repeat_train') > 0, '깊이 sigma 양수 필요')
        integer(f.get('repeat_train_people'), 'repeat_train_people', 5)
    for key in ('min_confidence', 'max_cross_view_error'):
        numeric(depth.get('quality', {}).get(key), f'depth.quality.{key}')
    require(depth['quality'].get('fit_scope'), '깊이 품질 기준 출처 없음')
    require(service.get('schema_version') == 'release-service-verification-v1', '서비스 검증 schema 불일치')
    require(service.get('complete') is True and service.get('all400') is True, '서비스 400명 검증 미완료')
    rows = service.get('participants', [])
    require(len(rows) == 400 and len({v.get('subject_id') for v in rows}) == 400, '서비스 400개 고유 기록 필요')
    require({v.get('subject_id') for v in rows} == set(subjects), '서비스/registry 인물 집합 불일치')
    identity = service.get('service', {})
    require(identity.get('ready') is True, '서비스 준비 상태 확인 불가')
    if 'calibrated_head_count' in identity or 'uncalibrated_heads' in identity:
        require(identity.get('calibrated_head_count') == calibrated_count and
                isinstance(identity.get('uncalibrated_heads'), dict) and set(identity['uncalibrated_heads']) == uncalibrated,
                '실서비스 보정/미보정 헤드 신원 불일치')
    for key, digest in [('deployment_sha256', hashes['deployment']), ('depth_config_sha256', hashes['depth']), ('registry_sha256', hashes['registry'])]:
        require(identity.get(key) == digest, f'실서비스 {key} 불일치')
    require(identity.get('score_mode') == 'include_measurement_scores', '혼합 점수 모드 불일치')
    require(isinstance(service.get('warmup'), list) and isinstance(service.get('errors_tested'), list), '예열/오류검증 목록 누락')
    require(type(service.get('error_checks_passed')) is bool, '오류검증 결과 누락')
    for error in service['errors_tested']:
        require(error.get('case') and type(error.get('passed')) is bool, '오류검증 항목 불완전')
    require(service['error_checks_passed'] == bool(service['errors_tested'] and all(e['passed'] for e in service['errors_tested'])), '오류검증 총괄 결과 불일치')
    resume_events = service.get('resume_events', [])
    require(isinstance(resume_events, list), '재개 기록 형식 오류')
    previous = 0
    for event in resume_events:
        done = integer(event.get('completed_before_resume'), '재개 전 완료 수') if isinstance(event, dict) else None
        require(done is not None and previous <= done < 400 and isinstance(event.get('reason'), str) and event['reason'].strip(),
                '재개 기록 불완전/순서 오류')
        previous = done
    known_items = set(features) | set(depth['features'])
    for row in rows + service['warmup']:
        numeric(row.get('seconds'), 'HTTP seconds')
        integer(row.get('http_status'), 'HTTP status', 100)
        require(row.get('subject_id') in subjects, '서비스/예열 인물 집합 불일치')
        require(row.get('split') == subjects[row['subject_id']]['split'], '서비스/registry 분할 불일치')
        require(row.get('status') in ('COMPLETE', 'PARTIAL', 'REINPUT_REQUIRED', 'error'), '서비스 terminal status 누락')
        require(isinstance(row.get('progress_event_seconds'), list), 'progress 시각 목록 누락')
        for seconds in row['progress_event_seconds']:
            numeric(seconds, 'progress seconds')
            require(seconds <= row['seconds'], 'progress 시각이 요청 완료 이후입니다')
        require(row['progress_event_seconds'] == sorted(row['progress_event_seconds']), 'progress 시각 순서 오류')
        require(all(key in row for key in SCORE_PROOF), '점수식/선정 검증 필드 누락: 최종 보고서 생성 불가')
        if all(key in row for key in SCORE_PROOF):
            require(all(key in row for key in SCORE_PROOF), '점수식/선정 검증 필드 일부 누락')
            require(type(row['score_formula_verified']) is bool and type(row['selection_verified']) is bool, '점수식/선정 검증 상태 오류')
            for key in SCORE_PROOF[1:3]:
                integer(row[key], key)
            selected = row.get('passages', [])
            components = sum(p.get('score_feature_count', 0) for p in selected)
            require(row['score_formula_verified_count'] <= len(selected) and row['score_components_verified_count'] <= components, '점수 검증 개수가 실제 선정 개수를 초과합니다')
            if row['score_formula_verified']:
                require(row['score_formula_verified_count'] == len(selected) and row['score_components_verified_count'] == components, '점수식 검증 통과의 개수 불일치')
            if row['status'] == 'error':
                require(not row['score_formula_verified'] and not row['selection_verified'], '오류 응답을 점수/선정 검증 통과로 표시할 수 없습니다')
        if row['status'] == 'error':
            require(row.get('error_code'), '오류 코드 누락')
            continue  # Error terminals do not carry runtime hashes in the verifier schema.
        require(row.get('deployment_sha256') == hashes['deployment'] and row.get('depth_config_sha256') == hashes['depth'], '응답/실제 배포·깊이 해시 불일치')
        items = row.get('items', [])
        require(len(items) == row.get('features') and len({i.get('item_id') for i in items}) == len(items), '응답 특성 수/중복 불일치')
        require({i.get('item_id') for i in items} == known_items, '응답 전체 특성 집합 불일치')
        require(row.get('confirmed') == sum(i.get('status') == 'CONFIRMED' for i in items), '확정 합계 불일치')
        require(row.get('held') == sum(i.get('status') == 'HELD' for i in items), '보류 합계 불일치')
        require(row['confirmed'] + row['held'] == len(items), '응답 특성 상태 오류')
        for item in items:
            require(item.get('kind') in ('classification', 'depth'), '특성 kind 오류')
            require(isinstance(item.get('hold_reasons'), list), '보류 사유 목록 누락')
            for field in ('confidence', 'measurement_score'):
                value = numeric(item.get(field), field, nullable=True)
                require(value is None or value <= 1, '특성 점수 범위 오류')
        for stage in STAGES:
            numeric(row.get('timings', {}).get(stage), f'실서비스 {stage}')
        numeric(row.get('timings', {}).get('received_to_result_seconds'), '실서비스 수신 완료→결과 준비 시간')
        passages = row.get('passages', [])
        require(row.get('passage_count') == len(passages) <= 5, '문헌 수 불일치/5 초과')
        require([p.get('rank') for p in passages] == list(range(1, len(passages) + 1)), '문헌 순위 오류')
        for p in passages:
            require(numeric(p.get('G'), 'G') <= 1, 'G 범위 오류')
            integer(p.get('score_feature_count'), 'score_feature_count', 1)
            require(set(p.get('evidence_item_ids', [])) <= known_items, '문헌 근거 특성 오류')
        require(type(row.get('db_source_verified')) is bool, 'DB 출처 검증 상태 없음')
        integer(row.get('db_verified_passages'), 'db_verified_passages')
        integer(row.get('db_candidates'), 'db_candidates')
        require(row['db_verified_passages'] <= len(passages) <= row['db_candidates'], 'DB 검증/후보 수 불일치')
    require(type(service.get('score_formula_verified')) is bool, '점수식 총괄 검증 상태 누락')
    require(service['score_formula_verified'] == all(v['score_formula_verified'] for v in rows), '점수식 총괄 상태와 개별 기록 불일치')
    for key in SCORE_PROOF[1:3]:
        integer(service.get(key), f'총괄 {key}')
        require(service[key] == sum(v[key] for v in rows), '점수식 총괄/개별 검증 개수 불일치')
    integer(service.get('selection_verified_participants'), 'selection_verified_participants')
    require(service['selection_verified_participants'] == sum(v['selection_verified'] for v in rows), '선정 검증 총괄 개수 불일치')
    sources = corpus.get('sources', [])
    passages = corpus.get('passages', [])
    require(sources and passages and corpus.get('schema_version'), '문헌 원자료 없음')
    source_ids = {s.get('literature_id') for s in sources}
    require(len(source_ids) == len(sources) and None not in source_ids, '문헌 원자료 source 중복/누락')
    require(len({p.get('passage_id') for p in passages}) == len(passages), '문헌 조항 중복')
    for p in passages:
        require(p.get('passage_id') and p.get('literature_id') in source_ids, '문헌 출처 연결 오류')
        require(type(p.get('service_eligible')) is bool and isinstance(p.get('exclusion_reasons'), list), '문헌 적격/제외 근거 없음')
        require(p.get('annotation_status'), '문헌 검수 상태 없음')
        require(p['service_eligible'] or p['exclusion_reasons'], '제외 문헌 사유 없음')
        if p['service_eligible']:
            require(p.get('original_text') and p.get('translation') and p.get('source') and p.get('required_labels'), '적격 문헌 필수 내용 누락')
    for row in rows:
        require(all(p.get('passage_id') in {p['passage_id'] for p in passages if p['service_eligible']} for p in row.get('passages', [])), '실서비스 문헌이 현재 corpus 적격 목록에 없습니다')
    eligible = sum(p['service_eligible'] for p in passages)
    require(identity.get('repository') == 'postgresql_readonly', '실서비스 PostgreSQL repository 확인 불가')
    require(identity.get('eligible_passages') == eligible, '실서비스/현재 corpus 적격 조항 수 불일치')
    corpus_digest = canonical_sha(corpus)
    require(service.get('corpus_sha256') is not None and identity.get('corpus_sha256') is not None,
            '활성 corpus canonical 검증 해시 누락: 최종 보고서 생성 불가')
    if service.get('corpus_sha256') is not None:
        require(service.get('corpus_sha256') == corpus_digest and identity.get('corpus_sha256') == corpus_digest,
                '실서비스/corpus canonical JSON SHA256 불일치')
        for row in rows + service['warmup']:
            if row['status'] != 'error':
                require(row.get('corpus_sha256') == corpus_digest, '응답/corpus canonical JSON SHA256 불일치')
    bundle.update(training=training, historical_training=historical_training, hashes=hashes, paths=paths, root=root,
                  synthetic=synthetic, corpus_canonical_sha256=corpus_digest, architectures=architectures,
                  candidates=list(candidates), inference_selection=selection, threshold_policy=threshold_policy)
    correction_path = root / 'outputs/release/ear-input-correction/verification.json'
    bundle['ear_correction'] = bundle['ear_correction_scope'] = None
    if correction_path.is_file():
        correction = read(correction_path)
        bundle['ear_correction_scope'] = validate_ear_correction(correction, d, parts, candidates, architectures)
        bundle['ear_correction'] = correction
        paths['ear_correction'] = correction_path
        hashes['ear_correction'] = sha(correction_path)
    influence_path = root / 'outputs/release/score-influence.json'
    bundle['score_influence'] = None
    if influence_path.is_file():
        influence = read(influence_path)
        require(bool(influence.get('synthetic', False)) == synthetic, '점수 영향 감사 합성/실측 표식 불일치')
        require(influence.get('version') == 'actual-release-score-influence-v1', '점수 영향 감사 schema 불일치')
        require(influence.get('service_report_sha256') == hashes['service'] and influence.get('corpus_sha256') == corpus_digest,
                '점수 영향 감사의 서비스/corpus 신원이 현재 입력과 다릅니다')
        require(influence.get('all400', {}).get('counts', {}).get('participants') == 400, '점수 영향 감사 400명 누락')
        for split, n in SPLITS.items():
            require(influence.get('splits', {}).get(split, {}).get('counts', {}).get('participants') == n, '점수 영향 감사 분할 불일치')
        bundle['score_influence'] = influence
        paths['score_influence'] = influence_path
        hashes['score_influence'] = sha(influence_path)
    stability_path = root / 'outputs/release/depth-stability.json'
    bundle['depth_stability'] = None
    if not synthetic:
        require(stability_path.is_file(), '실제 깊이 반복 안정성 감사가 누락됐습니다')
    if stability_path.is_file():
        stability = read(stability_path)
        require(bool(stability.get('synthetic', False)) == synthetic, '깊이 반복 감사 합성/실측 표식 불일치')
        require(stability.get('version') == 'release-depth-repeatability-v1', '깊이 반복 감사 schema 불일치')
        require(stability.get('depth_config_sha256') == hashes['depth'] and stability.get('records_sha256') == bundle['depth']['records_sha256'], '깊이 반복 감사의 구성/측정 파일 신원 불일치')
        require(stability.get('base_participants') == 400 and set(stability.get('features', {})) == set(bundle['depth']['features']), '깊이 반복 감사의 인물/특성 누락')
        bundle['depth_stability'] = stability
        paths['depth_stability'] = stability_path
        hashes['depth_stability'] = sha(stability_path)
    ui_path = root / 'outputs/release/frontend-actual-verification.json'
    bundle['actual_ui'] = None
    if not synthetic:
        require(ui_path.is_file(), '최종 실제 브라우저 검증이 누락됐습니다')
    if ui_path.is_file():
        ui = read(ui_path)
        require(bool(ui.get('synthetic', False)) == synthetic, '실제 브라우저 합성/실측 표식 불일치')
        validate_ui_evidence(ui, hashes, corpus_digest)
        bundle['actual_ui'] = ui
        paths['actual_ui'] = ui_path
        hashes['actual_ui'] = sha(ui_path)
    tls_path = root / 'outputs/release/tls-camera-verification.json'
    bundle['tls_camera'] = None
    if tls_path.is_file():
        tls = read(tls_path)
        require(bool(tls.get('synthetic', False)) == synthetic, 'TLS/카메라 합성/실측 표식 불일치')
        copy_path = root / 'outputs/release/frontend-camera-verification.json'
        validate_tls_camera(tls, bundle, read(copy_path) if copy_path.is_file() else None)
        bundle['tls_camera'] = tls
        paths['tls_camera'] = tls_path
        hashes['tls_camera'] = sha(tls_path)
    db_names = {'db_check': 'database-label-metadata-check.json', 'db_sync': 'database-label-metadata-sync.json',
                'db_status': 'database-final-status.json'}
    bundle['db_metadata'] = None
    if any((root / 'outputs/release' / name).is_file() for name in db_names.values()):
        docs = {key: read(root / 'outputs/release' / name) for key, name in db_names.items()}
        for key, doc in docs.items():
            require(bool(doc.get('synthetic', False)) == synthetic, f'{key}: 합성/실측 표식 불일치')
        bundle['db_metadata'] = {**validate_db_metadata(docs['db_check'], docs['db_sync'], docs['db_status'], bundle), **docs}
        for key, name in db_names.items():
            paths[key] = root / 'outputs/release' / name
            hashes[key] = sha(paths[key])
    recovery_path = root / 'outputs/release/windows-report-write-recovery/recovery.json'
    bundle['write_recovery'] = None
    if recovery_path.is_file():
        recovery = read(recovery_path)
        require(recovery.get('prefix_identical') is True and
                integer(recovery.get('preserved_committed'), 'preserved_committed') <= integer(recovery.get('recovered_actual_response_count'), 'recovered') <= 400,
                '응답 파일 복구 기록이 기존 응답 보존을 증명하지 않습니다')
        require(service.get('resume_events'), '응답 파일 복구 기록이 있으나 서비스 재개 기록이 없습니다')
        bundle['write_recovery'] = recovery
        paths['write_recovery'] = recovery_path
        hashes['write_recovery'] = sha(recovery_path)
    # Optional evidence is identified and hashed, never silently converted to a pass.
    optional = []
    for name in ('docs/release/frontend-report.txt', 'docs/release/database-report.txt', 'docs/release/tls-report.txt',
                 'data/release/literature/source_registry.json', 'data/release/literature/validation_report.json'):
        path = root / name
        if path.is_file():
            optional.append((name, sha(path)))
    bundle['optional'] = optional
    bundle['crop_evidence'] = []
    for path in sorted((root / 'outputs/release').glob('crop_parity*.json')):
        doc = read(path)
        require(bool(doc.get('synthetic', False)) == synthetic, 'crop parity 합성 표식 불일치')
        bundle['crop_evidence'].append((path.name, doc, sha(path)))
    return bundle


def make_tables(b):
    """Every shared cell is deliberately selected; participant rows stay private."""
    r, d, test, service = [b[k] for k in ('registry', 'deployment', 'test', 'service')]
    tables = {}
    def table(name, headers, rows):
        tables[name] = (headers, list(rows))
    table('분할', ['분할', '인물 수', '역할'], [
        [SPLIT_KO[s], n, {'train': '학습·깊이 반복 산포/품질 기준', 'selection': '체크포인트 선택',
                         'calibration': '헤드별 온도 보정' + ('·확정 기준값 선택' if b.get('threshold_policy') else ''),
                         'test': '배포 고정 후 역사적 시험'}[s]] for s, n in SPLITS.items()])
    stage_time_rows = []
    for stage, timing in b.get('stage_timings', {}).items():
        stage_time_rows.append([stage, '전체', timing['total_seconds'], timing['total_seconds']/60])
        stage_time_rows.extend([stage, PART_KO.get(part,part), value, value/60] for part,value in timing['parts'].items())
    table('보정시험시간', ['단계(인물 수)', '부위', '경과 초', '경과 분'], stage_time_rows)
    correction = b.get('ear_correction')
    table('귀입력수정', ['항목', '설명', '초', '분'], [
        ['시간 집계 범위', EAR_TIME_SCOPE, None, None],
        ['수정 검증', ear_correction_note(b), None, None],
        *[[label, '실측 검증 자료' if correction else '미제공/대기',
           correction[key] if correction else None, correction[key]/60 if correction else None]
          for key, label in EAR_COSTS],
    ])
    feature_rows, metric_rows, supports, calibration, bins, provenance, evaluation_supports = [], [], [], [], [], [], []
    for key, f in r['features'].items():
        row, cal = test['features'][key], d['features'][key]
        raw, corrected = row['raw'], row['calibrated']
        feature_rows.append([key, f['ko'], PART_KO.get(f['part'], f['part']), len(f['classes']),
                             row['photo_available'], raw['valid_count'], raw['missing_count'], raw['accuracy'],
                             raw['macro_f1'], corrected['nll'] if corrected else None, corrected['ece'] if corrected else None])
        cut = f.get('thresholds', {})
        provenance.append([key, f['ko'], f.get('source'), f.get('annotation_provenance'), row['annotation_provenance'],
                           r.get('human_review_count'), test.get('human_review_count'), f.get('metric'), cut.get('method'),
                           cut.get('physical_centre'), json.dumps(cut['cuts'], ensure_ascii=False) if 'cuts' in cut else None,
                           cut.get('fit_split'), cut.get('train_subjects'), cut.get('boundary_guard')])
        for split in SPLITS:
            item = r['class_support'][split][key]
            for i, cls in enumerate(f['classes']):
                supports.append([SPLIT_KO[split], key, f['ko'], cls,
                                 (f.get('classes_ko') or f['classes'])[i], item['counts'][i], item['unknown'],
                                 cal['train_support'][i] if split == 'train' else None,
                                 '참조 지지수 0: 재현율 평가불가' if item['counts'][i] == 0 else '참조 레이블 있음'])
        for split, variants in [('test', {'raw': raw, 'calibrated': corrected}),
                                ('calibration', {'raw': cal['raw'], 'calibrated': cal.get('calibrated_metrics')})]:
            for variant, m in variants.items():
                metric_rows.append([SPLIT_KO[split], key, variant] + ([m.get(k) for k in ('valid_count', 'missing_count', *METRICS)] if m else [None] * 8)
                                   + ([m['threshold'].get(k) for k in ('value', 'accepted_count', 'coverage', 'accuracy')] if m else [None] * 4))
                if m:
                    for i, cls in enumerate(f['classes']):
                        evaluation_supports.append([SPLIT_KO[split], key, variant, cls, m['support'][i], m['predicted_support'][i],
                                                    i in m['macro_f1_classes'], '재현율 평가불가' if m['support'][i] == 0 else '평가 가능'])
                    for binrow in m['ece_bins']:
                        bins.append([SPLIT_KO[split], key, variant] + [binrow.get(k) for k in ('lower', 'upper', 'count', 'accuracy', 'mean_confidence')])
        calibration.append([key, f['ko'], cal['calibrated'], cal['status'], cal.get('temperature'), cal.get('n_val'),
                            cal.get('at_bound'), cal.get('participant_separation_checked'),
                            raw['nll'], corrected['nll'] if corrected else None,
                            raw['ece'], corrected['ece'] if corrected else None,
                            corrected['nll'] - raw['nll'] if corrected and raw['nll'] is not None else None,
                            corrected['ece'] - raw['ece'] if corrected and raw['ece'] is not None else None])
    table('전체분류특성', ['특성 키', '특성', '부위', '클래스 수', '사진 가용/62', '평가 수', '미평가 수', '맞힌 비율', 'macroF1', '보정 NLL', '보정 ECE'], feature_rows)
    table('전체지표', ['분할', '특성 키', '원점수/보정', '평가 수', '미평가 수', *METRICS, '보류 기준', '수용 수', 'coverage', '수용 정확도'], metric_rows)
    table('클래스지지수', ['분할', '특성 키', '특성', '클래스 키', '클래스', 'registry 참조 인물 수', '특성 unknown(중복합산 금지)', '배포 유효 학습 인물(학습행만)', '해석'], supports)
    table('평가클래스지지수', ['분할', '특성 키', '원점수/보정', '클래스', '평가 참조 수', '예측 수', 'macroF1 평균 포함', '해석'], evaluation_supports)
    table('온도보정', ['특성 키', '특성', '보정됨', '상태', '온도T', '보정 유효 수', '경계해', '인물 분리 확인', '시험 raw NLL', '시험 T NLL', '시험 raw ECE', '시험 T ECE', 'NLL 차이(T-raw)', 'ECE 차이(T-raw)'], calibration)
    table('ECE구간', ['분할', '특성 키', '원점수/보정', '하한', '상한', '수', '정확도', '평균 신뢰도'], bins)
    if b.get('threshold_policy'):
        policy = b['threshold_policy']
        table('기준값정책', ['항목', '값'], [[k, ' / '.join(map(str, v)) if isinstance(v, list) else v]
                                       for k, v in policy.items() if isinstance(v, (str, int, float, bool, list))]
              + [['threshold_policy_frozen.json SHA256', b['hashes']['threshold_policy']]])
        threshold_rows = []
        for key, f in r['features'].items():
            cal, row = d['features'][key], test['features'][key]
            choice = cal['threshold_selection']
            picked = choice.get('selected') or {}
            test_cut = (row.get('calibrated') or row['raw'])['threshold']
            threshold_rows.append([key, f['ko'], len(f['classes']), cal['calibrated'], choice['status'], cal['threshold'],
                                   policy['binary_floor'] if len(f['classes']) == 2 else policy['multiclass_floor'],
                                   choice['n_valid'], picked.get('accepted'), picked.get('withheld'), picked.get('wrong_confirmed'),
                                   picked.get('coverage'), picked.get('accepted_error_rate'), picked.get('loss'),
                                   test_cut.get('coverage'), test_cut.get('accuracy')])
        table('헤드기준값', ['특성 키', '특성', '클래스 수', '보정됨', '선택 상태', '기준값', '정책 하한', '보정 유효 수',
                        '보정 수용', '보정 보류', '보정 오확정', '보정 coverage', '보정 수용 오류율', '보정 loss',
                        '시험 coverage(T 우선)', '시험 수용 정확도(T 우선)'], threshold_rows)
    table('주석출처', ['특성 키', '특성', '주석 원천', 'registry 주석 근거', '시험 주석 근거', 'registry 인간 검수 수', '시험 인간 검수 수', '측정 metric', '경계 설정 방법', '물리 중심', '경계값', '경계 적합 분할', '경계 적합 유효 인물', '경계 보류 규칙'], provenance)
    train_rows, history_rows, epochs, train_features, selection_rows = [], [], [], [], []
    for scope, models in (('현재 배포', b['training']), ('역사적 후보(미배포)', b.get('historical_training', []))):
        for model in models:
            meta, hist = model['meta'], model['history']
            arch, part = meta['arch'], meta['part']
            validation = sum(h['validation_seconds'] for h in hist) if all('validation_seconds' in h for h in hist) else None
            train = sum(h['train_seconds'] for h in hist)
            (train_rows if scope == '현재 배포' else history_rows).append([
                ARCH.get(arch, arch), PART_KO.get(part, part), len(meta['classification_features']),
                model['train_people_with_crops'], model['train_crops'], len(hist), model['selected_epoch'],
                len(hist) == d['config']['epochs_max'], model['selected_epoch'] == d['config']['epochs_max'],
                train, validation, train + validation if validation is not None else None, model['elapsed_seconds'], scope])
            for h in hist:
                epochs.append([ARCH.get(arch, arch), PART_KO.get(part, part)] + [h.get(k) for k in ('epoch', 'loss', 'selection_macro_f1', 'train_seconds', 'validation_seconds', 'amp_skipped_steps')] + [scope])
            for key in meta['classification_features']:
                train_features.append([ARCH.get(arch, arch), PART_KO.get(part, part), key, r['features'][key]['ko'], model['train_people_by_feature'][key], scope])
                m = model['selection'][key]['metrics']
                selection_rows.append([ARCH.get(arch, arch), key, m['valid_count'], *[m[k] for k in METRICS], scope])
    train_headers = ['모델', '부위', '헤드 수', 'crop 가용 학습 인물', '학습 crop 수', '실행 epoch', '선택 epoch', 'budget까지 실행', 'budget epoch 선택', '학습합 초', '검증합 초', '학습+검증합 초', '전체 경과 초', '구분']
    table('모델학습시간', train_headers, train_rows)
    if history_rows:
        table('역사적후보학습시간', train_headers, history_rows)
    table('학습epoch', ['모델', '부위', 'epoch', 'loss', '선택 macroF1', '학습 초', '검증 초', 'AMP skip', '구분'], epochs)
    table('유효학습인물', ['모델', '부위', '특성 키', '특성', 'crop+레이블 가용 학습 인물', '구분'], train_features)
    table('모델선택지표', ['모델', '특성 키', '선택 평가 수', *METRICS, '구분'], selection_rows)
    if b.get('inference_selection'):
        s = b['inference_selection']
        table('추론선택', ['항목', '값'], [
            ['schema', s['version']], ['현재 추론 계열', ' + '.join(ARCH.get(a, a) for a in b['architectures'])],
            ['현재 배포 가중치 수', len(d['weights'])], ['앙상블', '아니오(단일 계열)'],
            ['원 학습 후보 계열(역사적 config)', ' + '.join(ARCH.get(a, a) for a in b['candidates'])],
            ['원 학습 후보 모델 수(역사적)', len(b['training']) + len(b['historical_training'])],
            ['선택 근거(원문)', s['selection_basis'] if isinstance(s['selection_basis'], str) else json.dumps(s['selection_basis'], ensure_ascii=False, sort_keys=True)],
            ['결정 시각 UTC', s.get('decided_utc')], ['registry SHA256', s['registry_sha256']],
            ['training_frozen.json SHA256', s['training_config_sha256']],
            ['inference_frozen.json SHA256', b['hashes']['inference_selection']],
            ['보정 적용', '선택 계열 출력에만 헤드별 온도를 다시 맞춤(config의 "평균" 문구는 원 3계열 설정 원문)'],
        ])
    part_rows = []
    for part, spec in r['parts'].items():
        heads = [test['features'][key]['raw'] for key in spec['features']]
        models = [m for m in b['training'] if m['meta']['part'] == part]
        part_rows.append([PART_KO.get(part, part), len(heads), sum(m['valid_count'] > 0 for m in heads),
                          avg([m['macro_f1'] for m in heads]), avg([m['accuracy'] for m in heads]),
                          sum(sum(h['train_seconds'] for h in m['history']) for m in models) / 60,
                          sum(m['elapsed_seconds'] for m in models) / 60])
    table('부위요약', ['부위', '헤드 수', '시험 평가가능 헤드', '시험 macroF1 평균', '시험 맞힌 비율 평균', '배포모델 학습 분', '배포모델 전체 경과 분'], part_rows)
    request_rows, latency, stages, statuses, holds = service['participants'], [], [], [], Counter()
    feature_counts = defaultdict(Counter)
    part_counts = defaultdict(Counter)
    dbrows, rankrows = [], []
    for split in ['전체', *SPLITS]:
        rows = request_rows if split == '전체' else [v for v in request_rows if v['split'] == split]
        name = SPLIT_KO.get(split, split)
        latency.append([name, *summary([v['seconds'] for v in rows])])
        statuses.extend([name, status, count] for status, count in sorted(Counter(v['status'] for v in rows).items()))
        good = [v for v in rows if v['status'] != 'error']
        proof = [v for v in rows if all(k in v for k in SCORE_PROOF)]
        dbrows.append([name, len(rows), len(good), sum(v['db_source_verified'] for v in good),
                       sum(v['passage_count'] > 0 for v in good), sum(v['db_verified_passages'] for v in good),
                       sum(v['passage_count'] for v in good), sum(v['db_candidates'] for v in good),
                       len(proof), sum(v['score_formula_verified'] for v in proof) if proof else None,
                       sum(v['score_formula_verified_count'] for v in proof) if proof else None,
                       sum(v['score_components_verified_count'] for v in proof) if proof else None,
                       sum(v['selection_verified'] for v in proof) if proof else None,
                       sum(v['score_formula_verified'] and v['passage_count'] == 0 for v in proof) if proof else None])
        for stage in (*STAGES, 'recorded_stage_sum', 'received_to_result_seconds'):
            values = [sum(v['timings'][k] for k in STAGES) if stage == 'recorded_stage_sum' else v['timings'][stage] for v in good]
            stages.append([name, stage, *summary(values)])
    for row in request_rows:
        for item in row.get('items', []):
            key = item['item_id']
            for scope in ('전체', row['split']):
                feature_counts[(scope, key, item['region'], item['kind'])][item['status']] += 1
                part_counts[(scope, item['region'])][item['status']] += 1
                for reason in item['hold_reasons']:
                    holds[(scope, key, reason)] += 1
    table('실서비스시간', ['분할', '요청 수', '평균 초', 'p95 초', '최대 초', '5초 초과 수'], latency)
    table('단계시간', ['분할', '단계', '기록 수', '평균 초', 'p95 초', '최대 초', '5초 초과 수'], stages)
    table('상태', ['분할', '응답 상태', '요청 수'], statuses)
    table('보류사유', ['분할', '특성 키', '보류 사유 코드', '설명', '항목 발생 수(복수사유 중복)'], [[SPLIT_KO.get(s, s), k, reason, REASON_KO.get(reason, '추가 코드: 원자료 정의 참조'), n] for (s, k, reason), n in sorted(holds.items())])
    table('실서비스특성', ['분할', '특성 키', '부위', 'kind', '확정', '보류'], [[SPLIT_KO.get(s, s), k, PART_KO.get(p, p), kind, c['CONFIRMED'], c['HELD']] for (s, k, p, kind), c in sorted(feature_counts.items())])
    table('실서비스부위', ['분할', '부위', '확정', '보류'], [[SPLIT_KO.get(s, s), PART_KO.get(p, p), c['CONFIRMED'], c['HELD']] for (s, p), c in sorted(part_counts.items())])
    table('DB응답검증', ['분할', '요청', '결과 응답', 'DB원문일치 응답', '문헌있는 응답', 'DB검증 반환조항(중복포함)', '반환조항(중복포함)', 'AND후보(중복포함)', '점수검증 필드 기록 요청', '점수식 검증 참 요청', '점수식 검증 조항', '검증 구성 특성', '선정 일치 요청', '문헌0개 점수검증 참 요청'], dbrows)
    table('예열', ['측정 구분', '요청 수', '평균 초', 'p95 초', '최대 초', '5초 초과 수'], [['예열(본400 제외)', *summary([v['seconds'] for v in service['warmup']])]])
    table('오류검증', ['실서버 합성 오류 입력', 'HTTP', '오류 코드', '초', '검증 통과'], [[v['case'], v.get('http_status'), v.get('error_code'), v.get('seconds'), v['passed']] for v in service['errors_tested']])
    if service.get('resume_events') or b.get('write_recovery'):
        resume_rows = [['재개', i + 1, e['completed_before_resume'], e['reason']] for i, e in enumerate(service.get('resume_events', []))]
        if b.get('write_recovery'):
            w = b['write_recovery']
            resume_rows.append(['Windows 파일 교체 복구', None, w['preserved_committed'], f"복구 {w['recovered_actual_response_count']}건, 기존 응답 접두 동일={w['prefix_identical']}; {w.get('reason', '')}"])
        table('응답기록재개', ['구분', '순번', '재개 전 완료 수', '내용'], resume_rows)
    table('요청오류집계', ['분할', '오류 코드', '수'], [[SPLIT_KO.get(s, s), e, n] for (s, e), n in sorted(Counter((v['split'], v['error_code']) for v in request_rows if v['status'] == 'error').items())])
    for rank in range(1, 6):
        entries = [p for v in request_rows for p in v.get('passages', []) if p['rank'] == rank]
        if entries:
            rankrows.append([rank, len(entries), avg([p['G'] for p in entries]), avg([p['score_feature_count'] for p in entries])])
    table('문헌순위집계', ['순위', '반환 수', 'G 평균(정답확률 아님)', '근거 특성 수 평균'], rankrows)
    depth_rows = []
    for key, f in b['depth']['features'].items():
        depth_rows.append([key, r['depth_features'][key].get('ko', key), f['n'], f['lower'], f['upper'], f['sigma_repeat_train'],
                           f['repeat_train_people'], f.get('population_sd')])
    table('깊이기준', ['특성 키', '특성', '경계 표본 수', '25% 경계', '75% 경계', '학습 반복 sigma', '학습 반복 인물', '집단 SD'], depth_rows)
    table('깊이품질', ['항목', '값/해석'], [[k, v] for k, v in b['depth']['quality'].items() if isinstance(v, (str, int, float, bool))])
    if b.get('depth_stability'):
        distribution, repeat = [], []
        for key, stats in b['depth_stability']['features'].items():
            distribution.append([key, stats['participants'], stats['mean'], stats['population_sd'], stats['skewness'], stats['excess_kurtosis'], stats['shapiro_w'], stats['shapiro_p'], stats['mean_assigned_score'], stats['normal_distribution_required']])
            for split, values in stats['repeat'].items():
                repeat.append([key, SPLIT_KO.get(split, '전체' if split == 'all' else split), values['people'], values['pairs']] + [values.get(k) for k in ('mean_absolute_change_from_L1','mean_shift_L2','mean_shift_L3','label_agreement_L1_L2','label_agreement_L1_L3','all_three_labels_agree','pooled_repeat_sd','icc_1_1')])
        table('깊이분포진단', ['특성', '인물', '평균', '집단 SD', '왜도', '초과첨도', 'Shapiro W', 'Shapiro p', '지정 구간점수 평균', '정규분포 필수 여부'], distribution)
        table('깊이반복안정성', ['특성', '분할', '반복 인물', 'L1 대비 쌍', '평균 절대 변화', 'L2 평균 이동', 'L3 평균 이동', 'L1-L2 분류일치', 'L1-L3 분류일치', '세 조명 모두 일치', '반복 SD', 'ICC(1,1)'], repeat)
    corpus = b['corpus']
    if 'source_sections' in corpus:
        archive_rows = []
        for source in corpus['sources']:
            sections = [a for a in corpus['source_sections'] if a['literature_id'] == source['literature_id']]
            archive_rows.append([source.get('title_ko', source['title']), len(sections), sum(len(a['original_text']) for a in sections), '확보 판본별 원문; 번역·조건 완료와 다름; 교감본 중복 포함'])
        table('원문보관범위', ['서명', '원문 절 수', '원문 문자 수', '범위'], archive_rows)
    score_kinds = Counter({'CNN 조건만':0, '깊이 조건만':0, '한 조항에 CNN+깊이 조건':0})
    for passage in corpus['passages']:
        if not passage['service_eligible']: continue
        kinds = {label.split(':')[0] in r['depth_features'] for label in passage['required_labels']}
        score_kinds['한 조항에 CNN+깊이 조건' if len(kinds)==2 else '깊이 조건만' if True in kinds else 'CNN 조건만'] += 1
    table('문헌조건구성', ['적격 조항 구성', '조항 수'], list(score_kinds.items()))
    coverage, exclusion = [], Counter()
    for source in corpus['sources']:
        entries = [p for p in corpus['passages'] if p['literature_id'] == source['literature_id']]
        coverage.append([source['literature_id'], source.get('title_ko', source['title']), source['title'], len(entries),
                         sum(p['service_eligible'] for p in entries), sum(not p['service_eligible'] for p in entries),
                         source.get('text_availability', '수록 판본 설명 참조'), source.get('text_witness'),
                         ' | '.join(source.get('source_urls', []))])
    for p in corpus['passages']:
        for reason in p['exclusion_reasons']:
            exclusion[reason] += 1
    extra = corpus.get('literature_coverage')
    if extra is None:
        registry_path = b['root'] / 'data/release/literature/source_registry.json'
        require(registry_path.is_file(), '서명별 미수록 범위 근거(literature_coverage/source_registry) 없음')
        extra = read(registry_path).get('coverage_only')
    require(isinstance(extra, list), '문헌 미수록 범위 목록 누락')
    table('문헌서명별범위', ['문헌 키', '서명', '원서명', '수록 조항', '서비스 적격', '제외', '원문 접근 상태', '판본/전재 범위', '출처 URL'], coverage)
    table('미수록문헌', ['서명', '원서명', '설계 지정', '전체 원문 상태'], [[s.get('title_ko'), s.get('work'), s.get('named_in_requirements'), s.get('fulltext_status')] for s in extra])
    table('문헌제외사유', ['사유 코드', '조항 발생 수(복수사유 중복)'], sorted(exclusion.items()))
    table('문헌주석상태', ['주석 상태', '조항 수'], sorted(Counter(p['annotation_status'] for p in corpus['passages']).items()))
    table('문헌조항범위', ['조항 키', '문헌 키', '부위', '서비스 적격', '주석 상태', '조건 수', '정확 조건', '제외 사유'],
          [[p['passage_id'], p['literature_id'], p.get('region'), p['service_eligible'], p['annotation_status'], len(p.get('required_labels', [])),
            ' | '.join(p.get('required_labels', [])), ' | '.join(p['exclusion_reasons'])] for p in corpus['passages']])
    evidence = [
        ['400명 HTTP', '실측' if not b['synthetic'] else '합성 TEST ONLY', '400 고유 인물/분할 및 실제 배포·깊이 해시 대조', '상태·실패·지연을 분리 집계'],
        ['문헌 응답', '실측 HTTP + PostgreSQL 대조' if not b['synthetic'] else '합성 TEST ONLY', '반환 조항의 원문·직역·해설·출처 대조', '빈 반환과 미호출 조건은 전체 DB 내용 검증의 근거가 아님'],
        ['프런트엔드', '별도 문서의 합성 UI·로컬 검출 검증', 'docs/release/frontend-report.txt' if any('frontend-report' in p for p, _ in b['optional']) else '검증 문서 미제공', '실제 인물 카메라·추론 성공으로 일반화하지 않음; 최신 상세 범위는 해당 문서'],
        ['DB 구현', '별도 문서의 PostgreSQL/합성 corpus 검증', 'docs/release/database-report.txt' if any('database-report' in p for p, _ in b['optional']) else '검증 문서 미제공', '단위 검증과 실제 corpus·400 응답 대조를 구분'],
        ['오류 경로', '실서버에 합성 오류 입력',
         f"{len(service['errors_tested'])}회 / {len({e['case'] for e in service['errors_tested']})}종 / 통과 {sum(e['passed'] for e in service['errors_tested'])}회",
         '오류 처리 검증이며 모델 정확도 시험이 아님. 재시작마다 같은 사례를 반복 기록'],
    ]
    if b.get('inference_selection'):
        evidence.append(['추론 계열 선택', 'inference_frozen.json 바이트 SHA256 + training_frozen.json + registry 대조',
                         f"{' + '.join(b['architectures'])} 단일 {len(d['weights'])}가중치; 원 학습 후보 {len(b['candidates'])}계열",
                         '사용자 지정 단일 계열 배포. 계열 우열·앙상블 대비 우열의 새 검증이 아님'])
    if b.get('tls_camera'):
        tls = b['tls_camera']
        cam = tls['camera']
        source = '물리 웹캠' if cam['physical_camera'] else '녹화 영상 입력(물리 웹캠 아님)'
        evidence.append(['HTTPS·카메라 경로', '동결 배포 HTTPS 서비스 + 브라우저 자동 촬영 대조',
                         f"{source}; 촬영 POST {cam['capture_posts']}회; 업로드 동일성 {cam['equivalent_item_count']}항목 최대차 {cam['max_score_difference']}",
                         '학습군1명 보조 검증. 범위 한정 인증서 신뢰이며 공개 배포·물리 카메라·정확도 시험이 아님'])
        table('TLS카메라검증', ['항목', '값'], [
            ['HTTPS 주소', tls.get('url')], ['범위 한정 CA로 인증서 검증', tls['certificate_verified_by_scoped_ca']],
            ['브라우저 인증서 SPKI 고정', tls.get('browser_certificate_spki_pinned')], ['OS/전역 신뢰 변경', tls['global_trust_modified']],
            ['자체 서명 개발 인증서', tls.get('self_signed_development')], ['물리 웹캠 시험', tls['physical_camera_tested']],
            ['녹화 입력 사용', cam['recorded_source']], ['검출기 mock', cam['detector_mocked']], ['실제 모델', cam['actual_models']],
            ['검증 인물/분할', f"{cam.get('participants')} / {cam.get('split')}"], ['촬영 POST / 전체 POST', f"{cam['capture_posts']} / {cam.get('total_posts')}"],
            ['카메라·업로드 동일 항목 수', cam['equivalent_item_count']], ['점수 최대 차이', cam['max_score_difference']],
            ['카메라 트랙 종료', cam['tracks_stopped']], ['사진 저장', cam['photos_saved']], ['외부 요청 수', cam['external_requests']],
            ['품질 게이트 버전', cam['quality_version']], ['배포 SHA256', cam['deployment_sha256']]])
    if b.get('db_metadata'):
        db_meta = b['db_metadata']
        evidence.append(['DB 설계 메타데이터', 'migration004 + 동결 정의 동기화 + 읽기 계정 대조',
                         f"레이블 {db_meta['labels']}개 메타데이터 일치; 적격 {db_meta['eligible']}조항·corpus revision 불변",
                         '설계 4.3 컬럼 추가. TEXT 자연키 유지, passage_no는 기존 locator. 레이블 정의의 전문가 검수가 아님'])
        table('DB설계메타데이터', ['항목', '값'], [
            ['metadata_revision', db_meta['metadata_revision']], ['레이블(메타데이터 완비)', db_meta['db_status']['labels_with_metadata']],
            ['신규/갱신/무변경', f"{db_meta['db_sync']['inserted']} / {db_meta['db_sync']['updated']} / {db_meta['db_sync']['unchanged']}"],
            ['적격 조항 / 전체 조항', f"{db_meta['db_status']['eligible']} / {db_meta['db_status']['passages']}"],
            ['corpus revision', db_meta['db_status']['corpus_sha256']],
            ['식별자 표현', 'TEXT 자연키(feature:class, corpus passage_id) 의도적 유지; BIGINT 재번호 없음'],
            ['passage chapter/passage_no', 'source.chapter / source.locator에서 생성; 원문 조번호를 새로 만들지 않음']])
    if b.get('threshold_policy'):
        evidence.append(['헤드별 기준값', 'threshold_policy_frozen.json 바이트 SHA256 + 후보 집계 재계산',
                         f"{b['threshold_policy']['version']}; 보정34 전용, test_used=false",
                         '설계 5.4의 검증자료 조정. 같은 보정군의 개발 기준이며 시험 성능 개선의 보장이 아님'])
    if b.get('actual_ui'):
        ui = b['actual_ui']
        evidence.append(['실제 브라우저', '동결 모델·PostgreSQL HTTP 응답과 파일 해시 대조',
                         f"학습군1명 원본3장; 확정{ui['confirmed']}항목; PDF{ui['pdf_pages']}쪽; 전문검증{ui['pdf_passage_text_verified_count']}조항",
                         '업로드·결과·모바일·PDF·사진정리의 통합 검증. 실제 웹캠/독립 정확도 시험 아님'])
        ui_fields = {'participants':'검증 인물 수', 'split':'분할', 'confirmed':'확정 항목', 'held':'보류 항목',
                     'passages':'반환 조항', 'pdf_passage_text_verified_count':'PDF 전문 검증 조항',
                     'pdf_pages':'PDF 쪽 수', 'one_page_button':'A4 한 장 인쇄 가능',
                     'artifact_identity_verified':'실제 배포 신원 대조', 'photos_saved':'사진 저장',
                     'external_requests':'외부 요청 수'}
        table('실제브라우저검증', ['항목', '값'], [[name, ui[key]] for key,name in ui_fields.items()])
    for name, doc, digest in b['crop_evidence']:
        evidence.append(['crop 일치', '실제 입력/캐시 대조' if not b['synthetic'] else '합성 TEST ONLY', name,
                         f"기록 pass={doc.get('pass')}; 비교 {doc.get('image_crops_compared', '미기록')}; 범위: {doc.get('scope', '미기록')}"])
    table('검증근거', ['대상', '근거 유형', '근거', '해석 범위'], evidence)
    table('방법과한계', ['항목', '내용'], [[i + 1, text] for i, text in enumerate(LIMITS)])
    if service.get('corpus_sha256') is None:
        tables['방법과한계'][1].append([len(LIMITS) + 1, 'HTTP 검증은 corpus canonical SHA256을 기록하지 않았다. 현재 JSON과 DB의 동일 리비전을 확증하지 않으며 적격 수·반환 조항 ID·당시 DB본문 대조 기록을 확인한 범위다.'])
    table('파일해시', ['자료', 'SHA256'], [[key, digest] for key, digest in b['hashes'].items()] + b['optional'] + [[name, digest] for name, _, digest in b['crop_evidence']])
    tables['파일해시'][1].append(['corpus canonical JSON (sort_keys, compact, UTF-8)', b['corpus_canonical_sha256']])
    config_rows = [[f'{k} (원 학습 설정)' if b.get('inference_selection') else k, ' + '.join(v) if isinstance(v, list) else v]
                   for k, v in d['config'].items() if isinstance(v, (str, float, int, bool, list))]
    if b.get('inference_selection'):
        config_rows += [['inference_architectures (현재 추론)', ' + '.join(b['architectures'])],
                        ['inference_selection_sha256', d['inference_selection_sha256']]]
    table('배포설정', ['항목', '값'], config_rows
          + [['depth framing', b['depth'].get('framing_version')], ['depth checkpoint', b['depth'].get('checkpoint')],
             ['depth checkpoint SHA256', b['depth'].get('checkpoint_sha256')], ['depth code commit', b['depth'].get('code_commit')],
             ['HTTP max_side', service.get('max_side')], ['HTTP input', service.get('input')]])
    table('부위입력설정', ['부위', '높이', '너비', 'crop', '분류 헤드'],
          [[PART_KO.get(part, part), spec['input_size'][0], spec['input_size'][1], ' | '.join(spec['crops']),
            ' | '.join(spec['features'])] for part, spec in r['parts'].items()])
    table('공식자료', ['자료', 'URL'], SOURCES)
    if b['score_influence']:
        influence_rows = []
        count_keys = {'participants', 'request_errors', 'selected_mixed', 'selected_depth_only', 'selected_cnn_only',
                      'order_changed_by_depth', 'selected_set_changed_by_depth', 'people_with_passages'}
        stats = ('n', 'mean', 'median', 'p10', 'p90', 'ge_099', 'zero_count', 'mean_negative_log_clipped_1e15')
        groups = {'전체': b['score_influence']['all400'], **b['score_influence']['splits']}
        for split, group in groups.items():
            name = SPLIT_KO.get(split, split)
            for key in sorted(count_keys & group['counts'].keys()):
                influence_rows.append([name, '실측 응답 재생 건수', key, integer(group['counts'][key], key)])
            for key in ('paired_people', 'mean_depth_minus_cnn_confirmed', 'mean_pairwise_depth_higher_rate'):
                influence_rows.append([name, '같은 응답 내 비교', key, numeric(group.get(key), key, minimum=-1, nullable=True)])
            for key, values in group.get('score_distributions', {}).items():
                require(key in {'classification_all', 'classification_confirmed', 'depth_all', 'depth_confirmed'}, '점수 영향 감사 집계 유형 오류')
                for stat in stats:
                    if stat in values:
                        influence_rows.append([name, key, stat, numeric(values[stat], stat)])
            for key, values in group.get('features', {}).items():
                require(key in r['features'] or key in r['depth_features'], '점수 영향 감사의 특성 집합 불일치')
                for stat in stats:
                    if stat in values:
                        influence_rows.append([name, key, stat, numeric(values[stat], stat)])
        table('점수영향감사', ['분할', '대상', '지표', '값'], influence_rows)
    return tables


def brief(b):
    rows = b['service']['participants']
    metrics = list(b['test']['features'].values())
    paired = [r for r in metrics if r['calibrated'] and r['raw']['nll'] is not None]
    f1 = avg([r['raw']['macro_f1'] for r in metrics])
    acc = avg([r['raw']['accuracy'] for r in metrics])
    if b.get('inference_selection'):
        deployed = (f"배포 분류 모델: {' + '.join(ARCH.get(a, a) for a in b['architectures'])} 단일 계열 {len(b['deployment']['weights'])}모델(앙상블 아님). "
                    f"원 학습 설정의 후보 {len(b['candidates'])}계열 학습 기록은 역사적 자료로 그대로 둔다. ")
    else:
        deployed = f"배포 계열: {' + '.join(ARCH.get(a, a) for a in b['architectures'])}. "
    return [
        ('합성 TEST ONLY — 실제 실험 결과가 아님. 공유 금지.' if b['synthetic'] else '실제 산출물에 근거한 전체 구현 실험 보고'),
        deployed +
        f"{len(b['registry']['parts'])}부위 / CNN {len(metrics)}특성 / 깊이 {len(b['depth']['features'])}특성. "
        f"배포 학습 seed {b['deployment']['config'].get('seed', '미기록')}; 과거 3시드 비교와 별도 실험이다. 인간 검수 0건.",
        f"역사적 시험62: 헤드 평균 macro F1 {fmt(f1)}, 맞힌 비율 {fmt(acc)}. "
        f"평가 가능 {sum(r['raw']['valid_count'] > 0 for r in metrics)}/{len(metrics)}헤드. 자동 참조 레이블 일치도다.",
        f"실제 HTTP400: 평균 {fmt(avg([r['seconds'] for r in rows]))}초, p95 {fmt(percentile([r['seconds'] for r in rows]))}초, "
        f"최대 {fmt(max(r['seconds'] for r in rows))}초, 5초 초과 {sum(r['seconds'] > 5 for r in rows)}건, "
        f"오류 응답 {sum(r['status'] == 'error' for r in rows)}건. 예열 {len(b['service']['warmup'])}건은 별도다.",
        f"보정 비교 가능한 {len(paired)}헤드: 시험 NLL raw {fmt(avg([r['raw']['nll'] for r in paired]))} → T {fmt(avg([r['calibrated']['nll'] for r in paired]))}; "
        f"ECE raw {fmt(avg([r['raw']['ece'] for r in paired]))} → T {fmt(avg([r['calibrated']['ece'] for r in paired]))}. "
        f"미보정 {sum(not v['calibrated'] for v in b['deployment']['features'].values())}헤드. 개선·악화 모두 원값을 보고한다.",
        f"문헌: 수록 {len(b['corpus']['passages'])}조항, 서비스 적격 {sum(p['service_eligible'] for p in b['corpus']['passages'])}조항, "
        f"제외 {sum(not p['service_eligible'] for p in b['corpus']['passages'])}조항. 확보 판본의 수록 범위이며 지정 고전 전권의 완전 구축을 뜻하지 않는다.",
        LIMITS[0], LIMITS[4], LIMITS[9],
    ]


def write_txt(path, b, tables):
    lines = [TITLE, *brief(b), '', '빈 수치는 미기록 또는 평가불가이며 0으로 대체하지 않는다.']
    for name, (headers, rows) in tables.items():
        lines.extend(['', f'[{name}]', '\t'.join(headers)])
        lines.extend('\t'.join(fmt(v).replace('\n', ' ') for v in row) for row in rows)
    path.write_text('\n'.join(lines) + '\n', encoding='utf-8')


def write_xlsx(path, b, tables):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    all_tables = {'읽는방법': (['내용'], [[v] for v in brief(b)] + [
        ['빈 셀은 미기록/평가불가다. 수치 0과 구별한다. 인물별 자료·사진·좌표는 포함하지 않는다.'],
        ['NLL/ECE/Brier 낮을수록 좋음. 평균은 유효 헤드 동일 가중. 보정 전후는 동일 헤드 짝 비교.'],
        ['HTTP는 전송 전부터 응답 종료까지, 단계합은 서버 기록된 여섯 단계의 합이며 전체 수신 후 지연과 동일하지 않다.'],
        ['클래스지지수 unknown은 특성별 값이 클래스마다 반복되므로 열 전체를 합산하지 않는다.'],
        ['원문·번역·주석 본문을 복제하지 않고 서명별/조항별 출처 범위를 제공한다.']]), **tables}
    for name, (headers, rows) in all_tables.items():
        ws = wb.create_sheet(name[:31])
        ws.append(headers)
        for row in rows:
            ws.append(list(row))
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.font = Font(name='맑은 고딕', color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='213A52')
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Do not execute source text as an Excel formula.
                cell.font = Font(name='맑은 고딕', size=10)
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                if isinstance(cell.value, float):
                    cell.number_format = '0.0000'
        from openpyxl.utils import get_column_letter
        for i, header in enumerate(headers, 1):
            width = min(65, max(14, len(header) * 1.7, max((len(str(row[i - 1])) for row in rows[:150] if len(row) >= i and row[i - 1] is not None), default=0) * 1.05))
            ws.column_dimensions[get_column_letter(i)].width = 100 if len(headers) == 1 else width
        ws.print_title_rows = '1:1'
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    wb.save(path)


def write_pdf(path, b, tables, font_dir):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    require((font_dir / 'malgun.ttf').is_file() and (font_dir / 'malgunbd.ttf').is_file(), 'Malgun 글꼴 없음: --font-dir 지정')
    pdfmetrics.registerFont(TTFont('Malgun', str(font_dir / 'malgun.ttf')))
    pdfmetrics.registerFont(TTFont('MalgunBold', str(font_dir / 'malgunbd.ttf')))
    pdfmetrics.registerFontFamily('Malgun', normal='Malgun', bold='MalgunBold', italic='Malgun', boldItalic='MalgunBold')
    styles = {
        'title': ParagraphStyle('title', fontName='MalgunBold', fontSize=20, leading=29, textColor=colors.HexColor('#213A52'), spaceAfter=15, wordWrap='CJK'),
        'head': ParagraphStyle('head', fontName='MalgunBold', fontSize=14, leading=21, spaceAfter=10, wordWrap='CJK'),
        'body': ParagraphStyle('body', fontName='Malgun', fontSize=9.4, leading=15.5, spaceAfter=8, wordWrap='CJK'),
        'cell': ParagraphStyle('cell', fontName='Malgun', fontSize=8, leading=11.8, wordWrap='CJK'),
        'small': ParagraphStyle('small', fontName='Malgun', fontSize=8.2, leading=12.5, spaceAfter=6, wordWrap='CJK'),
    }
    story = []
    def p(text, kind='body'):
        return Paragraph(escape(fmt(text)).replace('\n', '<br/>'), styles[kind])
    def para(text, kind='body'):
        story.append(p(text, kind))
    def page(title):
        if story:
            story.append(PageBreak())
        para(title, 'head')
    def tab(headers, rows, widths=None):
        if not rows:
            para('해당 기록 없음.', 'small')
            return
        if widths is None:
            widths = [174 * mm / len(headers)] * len(headers)
        data = [[p(h, 'cell') for h in headers]] + [[p(v, 'cell') for v in row] for row in rows]
        t = Table(data, colWidths=widths, repeatRows=1, hAlign='LEFT')
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E8EEF4')),
            ('LINEBELOW', (0, 0), (-1, 0), .7, colors.HexColor('#8294A5')),
            ('LINEBELOW', (0, 1), (-1, -1), .25, colors.HexColor('#D7DEE5')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        story.extend([t, Spacer(1, 9)])
    para('FaceTeller\n전체 구현 실험 보고서', 'title')
    for text in brief(b):
        para(text)
    para('분류·깊이·문헌·촬영은 서로 다른 검증 범위를 가진다. 전체 항목의 일괄 통과 대신 실제 수치와 남은 한계를 아래에 기록한다.')
    para('부록 XLSX/TXT에는 전체 특성·클래스·학습 epoch·보류 사유·출처 범위를 수록한다. 이 PDF는 인물 식별자나 사진 없이 핵심 실험을 요약한다.', 'small')
    page('1. 실험 구성과 비교 범위')
    tab(*tables['분할'])
    for i in (1, 2, 3, 5, 6):
        para(LIMITS[i])
    para(f"배포 전처리: {b['deployment']['config']['preprocess_version']}. 모델 입력 크기와 부위별 crop는 실제 registry/meta 일치를 검사했다.")
    if any(f.get('thresholds', {}).get('physical_centre') is not None for f in b['registry']['features'].values()):
        para('방향형 측정은 0, 입술 상하 비는 1을 물리 중심으로 삼는다. 해당 방향을 인구 분위수로 바꾸지 않는다. 눈썹 꼬리는 양 끝점의 전체 기울기(brow_outer_inner_tilt)다. 다른 측정형 경계의 학습군 분위수와 구분하며 특성별 실제 방법·경계·유효 인물 수는 XLSX에 수록했다.', 'small')
    if b.get('inference_selection'):
        s = b['inference_selection']
        basis = s['selection_basis'] if isinstance(s['selection_basis'], str) else json.dumps(s['selection_basis'], ensure_ascii=False, sort_keys=True)
        para(f"원 학습은 {' + '.join(ARCH.get(a, a) for a in b['candidates'])} 후보 {len(b['training']) + len(b['historical_training'])}모델로 수행했다(원 config 그대로). "
             f"현재 서비스는 {' + '.join(ARCH.get(a, a) for a in b['architectures'])} 단일 계열 {len(b['training'])}모델만 로드하며 앙상블하지 않는다. "
             f"보정34·시험62는 이 단일 계열로 다시 산출한 값이다. 선택 기록({s['version']}) 근거 원문: {basis}")
        para('이번 기록만으로 단일 계열이 앙상블보다 낫다거나 계열 우열, 신규 모든 특성의 사람 기준 검증을 결론 내리지 않는다. 미배포 후보의 학습 기록은 역사적 자료로 XLSX에 별도 표시한다.', 'small')
    else:
        para('기존 비교에서 선정한 계열을 확대 범위에 재학습한 결과다. 이번 기록만으로 단일 계열보다 앙상블이 낫다거나 신규 모든 특성이 사람 기준으로 검증되었다고 결론 내리지 않는다.')
    page('2. 부위별 시험 결과와 학습 시간')
    tab(*tables['부위요약'], widths=[22*mm, 13*mm, 18*mm, 30*mm, 30*mm, 30*mm, 31*mm])
    para('학습 분은 history.train_seconds 합이다. 전체 경과에는 캐시 준비·선택 평가·저장 등의 시간이 포함될 수 있어 별도로 제시한다. validation_seconds가 없으면 학습+검증 시간을 추정하지 않는다.')
    if b.get('stage_timings'):
        para(' / '.join(f"{stage}: {timing['total_seconds']/60:.2f}분" for stage,timing in b['stage_timings'].items()) + '. 배포 조합의 기록된 경과 시간이며 측정 범위 원문은 '
             + ' / '.join(f"{stage}: {timing['scope']}" for stage,timing in b['stage_timings'].items()) + '. 부위별 시간은 XLSX에 수록했다.', 'small')
    models = tables['모델학습시간'][1]
    para(f"현재 배포 {len(models)}모델: 기록된 학습합 {fmt(sum(row[9] for row in models)/60)}분, 전체 경과합 {fmt(sum(row[12] for row in models)/60)}분. "
         f"예산 epoch까지 실행 {sum(row[7] for row in models)}개, 예산 epoch가 최종 선택된 모델 {sum(row[8] for row in models)}개. 두 수는 별도다.")
    if '역사적후보학습시간' in tables:
        history = tables['역사적후보학습시간'][1]
        para(f"역사적 미배포 후보 {len(history)}모델: 학습합 {fmt(sum(row[9] for row in history)/60)}분, 전체 경과합 {fmt(sum(row[12] for row in history)/60)}분. "
             f"원 {len(models) + len(history)}모델 학습 기록의 일부로 보존하며 현재 배포 비용에 합산하지 않는다.", 'small')
        scope = b.get('ear_correction_scope')
        if scope:
            para(f"현재 배포 비귀 {scope['active_non_ear']}모델의 가중치 해시는 3계열 배포 당시 귀 수정 기록의 해시와 같아, 단일 계열 선택 뒤 재학습한 모델이 아니다. "
                 f"귀 {scope['active_ear']}모델은 귀 입력 수정 때 재학습한 모델이다.", 'small')
    para(EAR_TIME_SCOPE, 'small')
    para(ear_correction_note(b), 'small')
    feature_rows = tables['전체분류특성'][1]
    chunk_size = math.ceil(len(feature_rows) / math.ceil(len(feature_rows) / 23))
    chunks = [feature_rows[i:i+chunk_size] for i in range(0, len(feature_rows), chunk_size)]
    for index, chunk in enumerate(chunks, 1):
        page(f'3. 전체 CNN 특성 시험 ({index}/{len(chunks)})')
        tab(['특성', '부위', '평가/62', '맞힌 비율', 'macro F1', 'T NLL', 'T ECE'],
            [[v[1], v[2], v[5], *v[7:11]] for v in chunk], widths=[49*mm, 18*mm, 17*mm, 23*mm, 23*mm, 22*mm, 22*mm])
        para('유효 참조와 사진이 있는 사람만 헤드별 평가에 포함된다. 분모가 다른 특성을 같은 62명 완전 평가로 읽지 않는다. 참조 지지수 0 유형·unknown·사진 가용 수·주석 출처는 XLSX에서 확인할 수 있다.', 'small')
    page('4. 보정 전후와 보류')
    pair = [v for v in b['test']['features'].values() if v['calibrated'] and v['raw']['nll'] is not None]
    tab(['동일 헤드 짝 평균', 'raw', '온도 보정 T'], [[metric, avg([v['raw'][metric] for v in pair]), avg([v['calibrated'][metric] for v in pair])] for metric in ('nll', 'ece', 'accuracy', 'macro_f1')])
    para(f'비교 가능 헤드 {len(pair)}개. 보정 실패 또는 유효 표본 없는 헤드는 짝 평균에서 제외하며 빈 수치를 0으로 채우지 않는다.')
    single = sum(cal['status'] == 'single_class_calibration_support' for cal in b['deployment']['features'].values())
    para(f'보정34에서 관측 클래스가 1개뿐인 헤드 {single}개는 온도 보정을 승인하지 않고 확정 신뢰도를 막았다(원점수 지표는 유지). 인간 검수 수는 0이다.', 'small')
    if b.get('threshold_policy'):
        policy = b['threshold_policy']
        cuts = Counter(v['threshold'] for k, v in b['deployment']['features'].items() if v['threshold_selection']['status'] == 'selected_on_calibration')
        para(f"헤드별 기준값({policy['version']})은 시험62를 보기 전에 정책을 고정하고 같은 보정34에서만 골랐다. "
             f"비용은 오확정 {policy['wrong_confirmation_cost']}·보류 {policy['withheld_cost']}이며 하한은 2클래스 {policy['binary_floor']}·다중 {policy['multiclass_floor']}이다. "
             f"선택된 {sum(cuts.values())}헤드 분포: " + ', '.join(f'{k} {n}개' for k, n in sorted(cuts.items())) +
             '. 같은 보정군을 온도와 기준값에 함께 쓴 개발 기준이며 독립 정확도가 아니다. 헤드별 후보 결과는 XLSX/TXT에 있다.', 'small')
    zero_classes = sum(n == 0 for cal in b['deployment']['features'].values() for n in cal['train_support'])
    para(f'배포 유효 학습 지지수가 0인 특성·클래스 조합은 {zero_classes}개다. 이 유형의 독립적인 양성 검출 능력을 확인했다고 주장하지 않는다. 원본 참조 지지수와 crop 가시성을 반영한 유효 지지수는 별도 열로 제공한다.', 'small')
    for metric in ('nll', 'ece'):
        changes = [v['calibrated'][metric] - v['raw'][metric] for v in pair]
        para(f"{metric.upper()}: 개선 {sum(v < -1e-12 for v in changes)} / 동일 {sum(abs(v) <= 1e-12 for v in changes)} / 악화 {sum(v > 1e-12 for v in changes)}헤드. 낮을수록 좋은 지표다.")
    held = Counter()
    for row in b['service']['participants']:
        for item in row.get('items', []):
            held.update(item['hold_reasons'])
    tab(['보류 사유(상위 8종)', '발생 항목 수'], [[REASON_KO.get(k, k), n] for k, n in held.most_common(8)], widths=[140*mm, 34*mm])
    para('사유 코드는 실서비스 원값이다. 한 항목에 복수 사유가 있을 수 있으므로 사유 합계를 인물 수나 보류 항목 수로 읽지 않는다. 모든 코드는 XLSX/TXT에 수록했다.')
    para(LIMITS[6])
    page('5. 실제 HTTP400 처리시간과 오류')
    tab(*tables['실서비스시간'])
    tab(*tables['예열'])
    stage_rows = [[STAGE_KO[row[1]], *row[2:]] for row in tables['단계시간'][1] if row[0] == '전체']
    tab(['서버 기록 단계', '수', '평균 초', 'p95 초', '최대 초', '>5초'], stage_rows, widths=[54*mm, 16*mm, 26*mm, 26*mm, 26*mm, 26*mm])
    para('HTTP 시간은 클라이언트가 전송을 시작하기 전부터 NDJSON 종료까지다. 업로드와 응답 전송이 포함된다. 서버 여섯 단계 합은 계측된 부분의 합이며 요청 수신 완료부터 결과까지의 전체 벽시계 시간은 아니다. 누락 단계를 0으로 채우지 않는다.', 'small')
    para('수신 완료→결과 준비 시간은 이미지 디코딩·검증·DB·단계 사이 대기를 포함한 연속 계측이다. 마지막 NDJSON 직렬화·전송·정리는 제외하므로 클라이언트 결과 수신 시간과 구분한다.', 'small')
    para('p95는 정렬된 측정값의 (n-1)×0.95 위치를 선형 보간했다. 5초 초과는 엄격히 >5초다. 실패 요청도 HTTP 집계에 포함하며 서버 단계는 결과 응답만 집계한다.', 'small')
    statuses = Counter(v['status'] for v in b['service']['participants'])
    para('응답 상태: ' + ', '.join(f'{k} {v}건' for k, v in sorted(statuses.items())) + '. 오류 사례별 검증은 XLSX 참조.', 'small')
    events = b['service'].get('resume_events', [])
    if events:
        para(f"이 400명 실행은 {len(events)}회 재개됐다(재개 전 완료 " + ', '.join(str(e['completed_before_resume']) for e in events) +
             "명). 이미 기록된 실제 응답은 그대로 두고 같은 서비스 신원으로 나머지만 이어서 측정했다. 예열·오류 검증은 재시작마다 다시 기록되어 "
             f"예열 {len(b['service']['warmup'])}건·오류 검증 {len(b['service']['errors_tested'])}회로 집계된다. 모든 결과 응답의 배포 SHA256이 같아 재개 전후 모델·보정·기준값은 동일하다.", 'small')
    page('6. VGGT 깊이 기준과 품질 한계')
    tab(['특성', '인물', '25% 경계', '75% 경계', '반복 sigma', '반복 인물'], [[v[1], *v[2:7]] for v in tables['깊이기준'][1]], widths=[44*mm, 16*mm, 28*mm, 28*mm, 30*mm, 28*mm])
    quality_names = {'min_confidence':'깊이 출력 신뢰도 하한', 'max_cross_view_error':'시점 간 상대 오차 상한', 'fit_scope':'품질 기준 계산 범위'}
    quality_rows = [[quality_names.get(k,k), '학습270명의 신뢰도1백분위수·시점 오차99백분위수. 실제3D 정답 기준은 아님' if k == 'fit_scope' and v == 'train270 empirical 1st/99th percentiles; not metric 3D ground truth' else v] for k,v in tables['깊이품질'][1]]
    tab(['검사 항목', '기준'], quality_rows, widths=[54*mm, 120*mm])
    para(LIMITS[7])
    para('두 경계에 대한 정규 누적분포 함수(CDF)를 사용한다. 낮음/보통/높음의 경계점수는 각각 1-P1, P1-P2, P2이다. sigma는 학습 인물 반복 측정 산포이며 단순 집단 표준편차로 대신하지 않는다.')
    para(LIMITS[8])
    para(LIMITS[9])
    if b.get('depth_stability'):
        tab(['특성', '반복 인물', 'L1-L2 일치', 'L1-L3 일치', 'ICC(1,1)'], [[b['registry']['depth_features'].get(v[0],{}).get('ko',v[0]), v[2], v[7], v[8], v[11]] for v in tables['깊이반복안정성'][1] if v[1] == '전체'], widths=[44*mm, 25*mm, 35*mm, 35*mm, 35*mm])
        para('같은 사람을 조명3종에서 다시 계산한 안정성이다. 실제3D 정답 대비 정확도가 아니며, 정규성 검정과 분할별 결과는 XLSX/TXT에 수록했다. 분포 진단으로 경계나 모델을 다시 선택하지 않았다.', 'small')
    page('7. 문헌 수록 범위와 실제 DB 연결')
    text_status = {'standalone_fulltext':'독립 판본 원문', 'fulltext_reproduced_in_compilation':'편찬서에 전재된 원문', 'image_only_scan_ai_transcribed':'스캔 원문과 AI 전사'}
    tab(['서명', '수록', '적격', '제외', '원문 접근 상태'], [[v[1], v[3], v[4], v[5], text_status.get(v[6],v[6])] for v in tables['문헌서명별범위'][1]], widths=[35*mm, 17*mm, 17*mm, 17*mm, 88*mm])
    if '원문보관범위' in tables:
        archive = tables['원문보관범위'][1]
        para(f"별도 원문 보관: {sum(v[1] for v in archive)}절, {sum(v[2] for v in archive):,}자. 확보한 판본의 본문 전체를 보관하며 교감본의 중복도 포함한다. 표의 수록 조항 수는 번역·조건 주석을 만든 조항 수로, 원문 보관량과 다르다.", 'small')
    if tables['미수록문헌'][1]:
        tab(*tables['미수록문헌'], widths=[30*mm, 30*mm, 23*mm, 91*mm])
    tab(['문헌 제외 사유(상위 6종)', '발생 수'], sorted(tables['문헌제외사유'][1], key=lambda v: -v[1])[:6], widths=[145*mm, 29*mm])
    db = tables['DB응답검증'][1][0]
    para(f"400 요청 중 결과 응답 {db[2]}건, DB 출처 일치 {db[3]}건, 문헌이 있는 응답 {db[4]}건. 반환 {db[6]}조항 중 DB 대조 {db[5]}조항(반복 반환 포함).")
    if db[8]:
        para(f'점수식·선정 검증 기록 {db[8]}/400요청: 점수식 검증 참 {db[9]}, 실제 검증 조항 {db[10]}, 구성 특성 {db[11]}, 선정 일치 {db[12]}건. 이 중 문헌0개에서 점수검증 참인 요청은 {db[13]}건이며 실조항 점수 검증 건수와 구분한다. 예열은 제외했다.', 'small')
    else:
        para('점수식·선정 결과의 독립 검증 필드는 미기록이다. G·순위를 표시했다는 사실을 계산/선정 검증 통과로 대신하지 않는다.', 'small')
    if b['score_influence']:
        para('동일 서비스 파일·corpus 해시의 score-influence.json도 대조했다. CNN/깊이 평균·포화 비율·깊이 제외 시 순위 변화를 XLSX/TXT에 제공한다. 관측된 영향이며 공정성이나 정답 정확도의 증거는 아니다.', 'small')
    else:
        para('실측 score-influence.json은 미제공이다. 수학적 동일 가중만으로 실제 CNN/깊이 영향이 같다고 주장하지 않는다.', 'small')
    kind_counts = dict(tables['문헌조건구성'][1])
    para(f"적격 문헌의 조건 구성: CNN만 {kind_counts['CNN 조건만']}개, 깊이만 {kind_counts['깊이 조건만']}개, 한 조항에 둘을 함께 요구하는 경우 {kind_counts['한 조항에 CNN+깊이 조건']}개.", 'small')
    if kind_counts['한 조항에 CNN+깊이 조건'] == 0:
        para('현재 문헌에서는 두 종류의 후보를 함께 순위에 넣을 수 있지만, 한 조항 안에서 CNN과 깊이를 결합한 실제 선정 사례는 없다. 조합 계산 기능의 검증과 실제 문헌 적용 범위를 구분한다.', 'small')
    para('수록·적격 수는 corpus의 실제 조항을 세었다. 제외 사유가 여러 개인 조항은 사유별 중복 집계한다. 원문 전권 접근, 일부 장/전재문 확보, 조건 주석 완료, 서비스 적격은 서로 다른 범위다.')
    para('실서비스 반환 문헌과 PostgreSQL의 원문·직역·해설·출처가 같은지 확인한 기록이다. 빈 후보·빈 반환의 성공을 전 DB 모든 조항의 검증으로 세지 않는다. 원문 AI 번역·주석은 인간 전문가 검수 완료가 아니다.', 'small')
    if b['service'].get('corpus_sha256') is None:
        para('HTTP 검증에 corpus canonical 해시가 없어 현재 JSON과 DB의 동일 리비전까지 확증하지는 않는다. 현재 적격 조항 수와 반환 조항 ID를 추가 대조한 범위다.', 'small')
    else:
        para('현재 corpus의 canonical JSON SHA256과 서비스 health·검증 총괄·모든 결과 응답의 활성 DB corpus 해시를 대조했다. raw 파일 SHA256과 구별한다.', 'small')
    page('8. 검증 근거와 남은 범위')
    tab(*tables['검증근거'], widths=[23*mm, 34*mm, 51*mm, 66*mm])
    para(LIMITS[10])
    para('사진·원시 좌표·개인별 예측·인물 식별자는 이 공유 보고서에 포함하지 않는다. 지표와 보류 사유는 부위·특성·분할별 집계로만 공개한다.')
    page('9. 재현 정보와 공식 자료')
    for key in ('registry', 'deployment', 'inference_selection', 'training_config', 'threshold_policy', 'test', 'depth', 'service', 'corpus'):
        if key in b['hashes']:
            para(key + ' SHA256: ' + b['hashes'][key], 'small')
    para('corpus canonical SHA256: ' + b['corpus_canonical_sha256'], 'small')
    para('실서비스·시험·학습 메타데이터의 배포/registry 해시와 모든 결과 응답의 깊이 설정 해시를 대조했다. 가중치 파일을 역직렬화하거나 GPU를 사용하지 않았다. weight 바이너리를 다시 해시한 검증은 아니며 배포 manifest와 런타임 보고 신원을 대조한 범위다.', 'small')
    para('확인 시점: ' + datetime.now(timezone.utc).isoformat(timespec='seconds') + ' (UTC)', 'small')
    for title, url in SOURCES:
        para(title + '\n' + url, 'small')
    for source in b['corpus']['sources']:
        urls = source.get('source_urls', [])
        if urls:
            para(source.get('title_ko', source['title']) + ': ' + urls[0], 'small')
    para('모든 문헌 URL·파일 해시·주석 상태는 XLSX/TXT에 제공한다. 원본 상세설계 PDF는 변경하지 않는다.', 'small')
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('Malgun', 8)
        canvas.setFillColor(colors.HexColor('#536477'))
        label = '합성 TEST ONLY · 공유 금지' if b['synthetic'] else 'FaceTeller · 실제 산출물 집계 · 인물별 자료 제외'
        canvas.drawString(18*mm, 12*mm, label)
        canvas.drawRightString(192*mm, 12*mm, str(doc.page))
        canvas.restoreState()
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=18*mm, leftMargin=18*mm,
                            topMargin=17*mm, bottomMargin=20*mm, title=TITLE,
                            author='FaceTeller', pageCompression=1)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    require(8 <= doc.page <= 12, f'PDF {doc.page}쪽: 8~12쪽 범위 초과. 공개 전 레이아웃 조정 필요')
    return doc.page


def privacy_check(b, tables):
    # No person ID may cross the aggregate allowlist, even through source strings.
    import re
    strings = '\n'.join(str(v) for headers, rows in tables.values() for row in [headers, *rows] for v in row if isinstance(v, str))
    for subject in b['registry']['subjects']:
        require(not re.search(r'(?<![A-Za-z0-9])' + re.escape(subject) + r'(?![A-Za-z0-9])', strings), '집계에 인물 식별자 포함: 공유 중단')


def run(args):
    root = args.root.resolve()
    output = (args.output_dir or root / 'outputs/release').resolve()
    if args.synthetic:
        temp = Path(tempfile.gettempdir()).resolve()
        require(root.is_relative_to(temp) and output.is_relative_to(temp), '합성 root/output은 OS TEMP 안이어야 합니다')
        require(args.no_copy, '합성 검증은 --no-copy 필수')
    bundle = load_bundle(root, args.corpus, args.synthetic)
    tables = make_tables(bundle)
    privacy_check(bundle, tables)
    if args.validate_only:
        print('검증 완료: 필수 자료·400명·해시·집계/개인정보 검사. 보고서는 생성하지 않았습니다.')
        return
    # Stage all three files first; generation failure cannot replace published reports.
    with tempfile.TemporaryDirectory(prefix='faceteller-report-stage-') as temporary:
        stage = Path(temporary)
        write_txt(stage / f'{TITLE}.txt', bundle, tables)
        write_xlsx(stage / f'{TITLE}.xlsx', bundle, tables)
        pages = write_pdf(stage / f'{TITLE}.pdf', bundle, tables, args.font_dir)
        # Inputs must not have changed while rendering.
        require(all(sha(bundle['paths'][key]) == value for key, value in bundle['hashes'].items()), '생성 중 입력 변경: 재실행 필요')
        destinations = [output]
        if not args.no_copy:
            destinations.append(args.share_dir.resolve())
        for destination in destinations:
            destination.mkdir(parents=True, exist_ok=True)
            for suffix in ('pdf', 'xlsx', 'txt'):
                name = f'{TITLE}.{suffix}'
                pending = destination / (name + '.partial')
                try:
                    shutil.copyfile(stage / name, pending)
                    os.replace(pending, destination / name)
                finally:
                    pending.unlink(missing_ok=True)
    print(f"{'합성 TEST ONLY' if args.synthetic else '실제 보고서'}: PDF {pages}쪽, XLSX {len(tables) + 1}시트, TXT 완료 → {output}")
    if not args.no_copy:
        print(f'공유 폴더 복사 완료 → {args.share_dir}')


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, default=ROOT)
    parser.add_argument('--corpus', type=Path, help='명시적 corpus 경로; 기본 reference_corpus_v1.json')
    parser.add_argument('--output-dir', type=Path)
    parser.add_argument('--share-dir', type=Path, default=Path('C:/Users/user/Downloads/FaceTeller_완성프로젝트_공유자료') if os.name == 'nt' else Path('/mnt/c/Users/user/Downloads/FaceTeller_완성프로젝트_공유자료'))
    parser.add_argument('--font-dir', type=Path, default=Path('C:/Windows/Fonts') if os.name == 'nt' else Path('/mnt/c/Windows/Fonts'))
    parser.add_argument('--no-copy', action='store_true')
    parser.add_argument('--synthetic', action='store_true')
    parser.add_argument('--validate-only', action='store_true')
    args = parser.parse_args(argv)
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8')
    try:
        run(args)
    except (EvidenceError, KeyError, TypeError, OSError, ImportError) as exc:
        print(f'보고서 생성 중단: {exc}', file=sys.stderr)
        return 2
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
