from pathlib import Path
import json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'outputs/experiment-report-2026-10-01'
x=json.loads((OUT/'보고서_구조화원본.json').read_text(encoding='utf8'))
d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.7);s.left_margin=Cm(1.8);s.right_margin=Cm(1.8)
s.header_distance=Cm(.8);s.footer_distance=Cm(.8);s.different_first_page_header_footer=True
def font(st,size,bold=False,col='243746'):
 st.font.name='맑은 고딕';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor.from_string(col)
 st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
font(d.styles['Normal'],10.5);f=d.styles['Normal'].paragraph_format;f.line_spacing=Pt(16);f.space_after=Pt(8)
for name,size,col in [('Title',26,'19394E'),('Subtitle',15,'087D82'),('Heading 1',17,'19394E'),('Heading 2',11,'087D82')]:
 font(d.styles[name],size,True,col);f=d.styles[name].paragraph_format;f.keep_with_next=True;f.space_after=Pt(8)
for name,size in [('Caption',8.5),('Header',8),('Footer',8),('TOC 1',10)]:
 if name not in d.styles:d.styles.add_style(name,1)
 font(d.styles[name],size,col='526371');d.styles[name].paragraph_format.line_spacing=Pt(14);d.styles[name].paragraph_format.space_after=Pt(3)
def field(p,code):
 r=p.add_run()
 for tag,value in [('begin',None),('instr',code),('separate',None),('end',None)]:
  e=OxmlElement('w:instrText' if tag=='instr' else 'w:fldChar')
  if tag=='instr':e.text=value;e.set(qn('xml:space'),'preserve')
  else:e.set(qn('w:fldCharType'),tag)
  r._r.append(e)
h=s.header.paragraphs[0];h.text='EXP-FACE-20261001  /  실제 실험 결과';h.style='Header'
f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;f.add_run('2026-10-01  |  ');field(f,'PAGE');f.add_run(' / ');field(f,'NUMPAGES')
def p(t,style=None):return d.add_paragraph(t,style)
def table(head,rows):
 n=len(head);weights={2:[.28,.72],3:[.32,.30,.38],4:[.23,.23,.21,.33],5:[.30,.175,.175,.175,.175],6:[.24,.13,.13,.14,.20,.16]}.get(n,[1/n]*n)
 dense=head==['피처','클래스','학습','검증','시험']
 if dense:weights=[.27,.40,.11,.11,.11]
 t=d.add_table(rows=1,cols=n);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for c,w in zip(t.columns,weights):c.width=Cm(17.4*w)
 for c,v in zip(t.rows[0].cells,head):c.text=str(v)
 rh=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rh)
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for ri,row in enumerate(t.rows):
  e=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(e)
  for ci,c in enumerate(row.cells):
   c.width=Cm(17.4*weights[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'19394E' if ri==0 else ('F0F5F7' if ri%2 else 'FFFFFF'));c._tc.get_or_add_tcPr().append(sh)
   margins=OxmlElement('w:tcMar')
   for side,val in [('top',40 if dense else 75),('bottom',40 if dense else 75),('left',85),('right',85)]:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
   c._tc.get_or_add_tcPr().append(margins)
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(1);par.paragraph_format.line_spacing=Pt(12.7)
    for r in par.runs:
     r.font.name='맑은 고딕';r.font.size=Pt(8.8);r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
     if ri==0:r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255)
 p('').paragraph_format.space_after=Pt(1)
p('EXPERIMENT REPORT / 2026.10','Subtitle');p('K-FACE 얼굴 형태 분석','Title');p('실제 실험 결과 보고서','Title')
p('학습·검증·시험 인원부터 38개 피처의 정확도까지, 보관된 결과와 모델 메타데이터를 대조해 정리했습니다.')
table(['검증 대상','규모'],[['서로 다른 얼굴','400명'],['최종 개발 / 시험','338명 / 62명'],['교차학습','5-fold, 학습 270~271명 / 검증 67~68명'],['저장된 형태 모델','5부위 × 2백본 × 5-fold = 50개'],['분석 피처','38개 — 수치·깊이 23개, AI 시각 합의 15개'],['기준일 / 판본','2026-10-01 / v1.0']])
p('확인 범위','Heading 2');p('2026-09-29까지 저장된 평가 결과와 체크포인트를 2026-10-01에 재집계했습니다. 새 학습이나 얼굴 추론을 실행한 보고서는 아닙니다.')
p('원본 얼굴 사진과 참여자별 명단은 첨부하지 않았습니다.','Caption')
d.add_page_break();p('목차','Title');field(p(''),'TOC \\o "1-1" \\h \\z \\u')
for i,pg in enumerate(x['pages'],1):
 d.add_page_break();p(f'{i:02d}. '+pg['title'],'Heading 1');p(pg['lead'])
 for b in pg['blocks']:
  if b['kind']=='p':p(b['text'])
  elif b['kind']=='table':table(b['head'],b['rows'])
  elif b['kind']=='bullets':
   for v in b['items']:p('• '+v)
d.core_properties.title=x['title'];d.core_properties.subject='보관된 학습·평가 실험의 근거 기반 재집계'
d.save(OUT/'실제_실험결과_보고서.docx')
print('DOCX_READY')
