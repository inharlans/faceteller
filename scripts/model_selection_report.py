"""CPU-only reports from completed selection results; never imports training code.

Run with Windows venv Python after model_selection_evaluate.py summary.
Synthetic checks require --synthetic --no-copy and an explicit temporary output.
"""
from __future__ import annotations

import argparse
import hashlib
import itertools
import json
import os
from pathlib import Path
import shutil
import statistics
import sys
import tempfile
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / 'outputs/model-selection-20261014/results.json'
TITLE = 'FaceTeller_CNN_모델선정_검증보고서'
ARCH = {'resnet18': 'ResNet-18', 'efficientnet_b0': 'EfficientNet-B0',
        'convnext_tiny': 'ConvNeXt-Tiny'}
SHORT = {'resnet18': 'R18', 'efficientnet_b0': 'E0', 'convnext_tiny': 'CN'}
HEADS = {'eyelid': '눈꺼풀', 'sclera_show': '흰자 노출', 'eye_depth': '눈 깊이(외형)',
         'brow_shape': '눈썹 모양', 'brow_thickness': '눈썹 굵기',
         'brow_density': '눈썹 밀도', 'brow_tail_clarity': '눈썹 꼬리 선명도',
         'nose_tip': '코끝 모양', 'nostril_show': '콧구멍 노출',
         'alar_fullness': '콧방울 풍만함', 'lip_bow': '윗입술 활 윤곽',
         'face_shape': '얼굴 윤곽', 'chin_shape': '턱끝 모양',
         'cheekbone': '광대 돌출(얼굴형)', 'cheek_fullness': '볼살'}
METRICS = ['macro_f1', 'accuracy', 'balanced_accuracy', 'nll', 'brier', 'ece',
           'threshold_coverage', 'threshold_accuracy']
SOURCES = [
    ('ResNet-18 공식 모델·가중치', 'https://docs.pytorch.org/vision/stable/models/generated/torchvision.models.resnet18.html'),
    ('EfficientNet-B0 공식 모델·가중치', 'https://docs.pytorch.org/vision/stable/models/generated/torchvision.models.efficientnet_b0.html'),
    ('ConvNeXt-Tiny 공식 모델·가중치', 'https://docs.pytorch.org/vision/stable/models/generated/torchvision.models.convnext_tiny.html'),
    ('PyTorch 전이학습 안내', 'https://docs.pytorch.org/tutorials/beginner/transfer_learning_tutorial.html'),
    ('Guo 등 (2017), 온도 보정 방법', 'https://proceedings.mlr.press/v70/guo17a.html'),
]
PROPOSAL_SCOPE = ('기존 5부위·15개 외형 분류 헤드의 모델 선택 부분에 한정한 문안 제안이다. '
                  '원래 설계의 신규 7영역 범위를 변경하지 않으며 원본 PDF는 수정하지 않았다.')
LIMITS = [
    '참조 레이블은 AI 3투표 외형 합의이며 독립 전문가 정답이 아니다. 외형 레이블 일치도이지 성격·운명·관상 예측 정확도가 아니다.',
    '시험 62명은 이전 실험에서 사용된 역사적 집합이다. 새로운 외부 모집단 검증이 아니다. 선택 34명은 조기 종료와 계열 선택에 함께 사용되어 선택 지표가 낙관적일 수 있다.',
    '3시드는 동일 참여자 분할을 공유한다. 시드 표준편차·최솟값·최댓값은 일반화 신뢰구간이 아니다. 첫 시드 paired bootstrap도 전체 학습 절차의 불확실성을 포함하지 않는다.',
    '범위는 기존 5부위의 15개 외형 분류 헤드뿐이다. 신규 7영역, 회귀 보조 출력, VGGT 깊이 정확도 또는 설계 전체 항목을 검증한 결과가 아니다.',
    '학습·평가는 얼굴 폭 기준 기존 crop cache를 사용한다. 설계의 눈 중심·눈 사이 거리 정규화와 다르며 코·얼굴 고정 resize도 완전한 종횡비 보존이 아니다. 실제 API의 안내 crop·640 JPEG·재검출 경로와도 차이가 있다.',
    '서로 다른 공식 ImageNet 사전학습 레시피를 포함한 실용 모델 비교이다. 구조만의 인과 우열이나 사전학습 자체의 효과를 분리한 실험이 아니다.',
    '비교 학습의 FP16 GradScaler 실행 중 optimizer 갱신 생략 관련 scheduler 경고가 기록됐다. 생략 횟수는 수집하지 않았으며 고정한 비교 레시피는 중간에 바꾸지 않았다. 최종 전체 모델 학습에서는 실제 optimizer 갱신 때만 scheduler를 진행하고 생략 횟수를 기록한다.',
    '지연은 예열된 CNN 단계만 측정한다. 업로드·디코딩·MediaPipe·VGGT·HTTP·콜드 스타트를 제외하므로 전체 서비스 5초 달성 근거로 쓰지 않는다.',
    '보고서 생성은 자동 배포를 수행하지 않는다. 실서비스 적용 전 로더·클래스 순서·전처리·보정·보류·문헌 순위의 통합검증이 필요하다.',
]


def read(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def number(value, digits=3):
    return '—' if value is None else f'{value:.{digits}f}'


def mean(values):
    values = [v for v in values if v is not None]
    return statistics.mean(values) if values else None


def label(name, short=False):
    lookup = SHORT if short else ARCH
    return ' + '.join(lookup.get(a, a) for a in name.split('+'))


def flat(value, prefix=''):
    """Preserve all scalar/nested metric fields without importing ML libraries."""
    if isinstance(value, dict):
        out = {}
        for k, v in value.items():
            out.update(flat(v, f'{prefix}.{k}' if prefix else str(k)))
        return out
    if isinstance(value, (list, tuple)):
        return {prefix: json.dumps(value, ensure_ascii=False, allow_nan=False)}
    return {prefix: value}


def validate(data, dev, dev_path):
    p, choice = data['protocol'], data['choice']
    arches, seeds, parts = p['architectures'], p['seeds'], p['parts']
    expected = {'+'.join(c) for n in (1, 2, 3) for c in itertools.combinations(arches, n)}
    if len(arches) != 3 or len(seeds) != 3 or len(parts) != 5 or len(expected) != 7:
        raise ValueError('Expected frozen 3 architectures × 3 seeds × 5 parts protocol')
    if p['participants'] != {'train': 270, 'selection': 34, 'calibration': 34, 'test': 62}:
        raise ValueError('Participant split differs from frozen report protocol')
    features = [key for part in parts for key in p['heads'][part]]
    if len(features) != 15 or len(set(features)) != 15:
        raise ValueError('Expected 15 unique visual heads')
    if set(data['families']) != expected or set(dev['families']) != expected:
        raise ValueError('All seven candidate families are required')
    if choice.get('test_predictions_present_at_freeze') is not False:
        raise ValueError('Choice was not frozen before test predictions')
    if dev['protocol_sha256'] != choice['protocol_sha256']:
        raise ValueError('Development protocol hash mismatch')
    if sha(dev_path) != choice['development_sha256']:
        raise ValueError('Development file differs from frozen choice')
    expected_runs = set(itertools.product(arches, seeds, parts))
    actual_runs = [(r['arch'], r['seed'], r['part']) for r in data['training_runs']]
    if len(actual_runs) != 45 or set(actual_runs) != expected_runs:
        raise ValueError('Expected 45 unique completed training runs')
    for r in data['training_runs']:
        if r['protocol_sha256'] != choice['protocol_sha256'] or not r['history']:
            raise ValueError('Training metadata incomplete or protocol mismatch')
        for epoch in r['history']:
            for key in ('train_seconds', 'validation_seconds'):
                if not isinstance(epoch[key], (int, float)) or epoch[key] < 0:
                    raise ValueError('Invalid core training time')
    for name, family in data['families'].items():
        for runs in (family['runs'], dev['families'][name]['runs']):
            if set(runs) != {str(s) for s in seeds}:
                raise ValueError('Incomplete seed results')
            for run in runs.values():
                for mode in ('raw', 'calibrated'):
                    if set(run[mode]['features']) != set(features):
                        raise ValueError('Incomplete per-head results')
        for run in dev['families'][name]['runs'].values():
            if set(run['temperatures']) != set(features):
                raise ValueError('Incomplete calibration temperatures')
        for split in ('selection', 'test'):
            for metric in METRICS:
                if len(family[split][metric]['values']) != 3:
                    raise ValueError('Incomplete three-seed summary')
    scores = {n: f['selection']['macro_f1']['mean'] for n, f in data['families'].items()}
    eligible = [n for n, score in scores.items()
                if score >= max(scores.values()) - choice['tolerance_absolute'] - 1e-12]
    winner = min(eligible, key=lambda n: data['families'][n]['latency']['mean_ms'])
    if winner != choice['selected_family'] or set(eligible) != set(choice['eligible']):
        raise ValueError('Frozen choice disagrees with recorded selection/latency rule')
    return features


def core_seconds(run):
    return sum(e['train_seconds'] + e['validation_seconds'] for e in run['history'])


def family_cost(data, name):
    arches = data['families'][name]['architectures']
    return [sum(core_seconds(r) for r in data['training_runs']
                if r['arch'] in arches and r['seed'] == seed) / 60
            for seed in data['protocol']['seeds']]


def replacements(data):
    chosen = data['choice']['selected_family']
    return [
        ('§5.1 교체 제안',
         f'{label(chosen)} — ImageNet 사전학습 모델을 부위별 외형 분류에 맞게 미세조정.'),
        ('§5.3 교체 제안',
         f'선정한 {label(chosen)} 구성으로 눈·눈썹·코·입·얼굴형 5부위의 외형을 분류한다. '
         '15개 항목별 형태 레이블·신뢰도·근거를 출력한다. '
         '신뢰도에는 해당 모델 구성과 항목에 맞게 적합한 온도 보정을 적용한다.'),
    ]


def proposal_reason(data):
    f = data['families'][data['choice']['selected_family']]
    return (f'선정 이유: 선택 macro-F1 {number(f["selection"]["macro_f1"]["mean"])}'
            f'(역사적 시험 {number(f["test"]["macro_f1"]["mean"])})이며, '
            f'CNN 평균 지연 {number(f["latency"]["mean_ms"], 2)} ms'
            f'(p95 {number(f["latency"]["p95_ms"], 2)} ms)로 '
            '선택 성능 허용 범위 내 후보 중 평균 지연이 가장 짧았다.')


def report_sections(data, features, provenance):
    """Seven bounded sections shared verbatim by PDF and readable UTF-8 TXT."""
    chosen = data['choice']['selected_family']
    families = data['families']
    f = families[chosen]
    first = str(data['choice']['deployment_seed'])
    best = max(families, key=lambda n: families[n]['selection']['macro_f1']['mean'])
    singles = [n for n in families if len(families[n]['architectures']) == 1]
    best_single = max(singles, key=lambda n: families[n]['selection']['macro_f1']['mean'])
    ens = max((n for n in families if n not in singles), key=lambda n: families[n]['selection']['macro_f1']['mean'])
    gain = families[ens]['selection']['macro_f1']['mean'] - families[best_single]['selection']['macro_f1']['mean']
    sections = []
    def section(title, paragraphs, tables=()):
        sections.append({'title': title, 'paragraphs': paragraphs, 'tables': list(tables)})
    section('1. 실제 선정 결과와 판단', [
        f'선정: {label(chosen)}. 구성은 5부위 × {len(f["architectures"])}계열 = {5 * len(f["architectures"])}개 모델이다. '
        f'배포 대표 시드 {first}는 통합검증용 대표이며 자동 배포되지 않는다.',
        f'선택 최고 후보는 {label(best)}({number(families[best]["selection"]["macro_f1"]["mean"])}). '
        f'허용 차이는 절대 macro-F1 {number(data["choice"]["tolerance_absolute"], 2)}이며 '
        f'{len(data["choice"]["eligible"])}개 적격 후보 중 warm CNN 평균 지연이 가장 짧은 구성을 선택했다. '
        'p95는 운영 비교 지표이며 선택 기준은 평균 지연이다.',
        'macro-F1은 많은 유형에만 치우치지 않도록 유형별 분류 성능을 평균한 지표이고, 정확도(accuracy)는 맞힌 비율이다. '
        '이 보고서는 각 항목의 지표를 다시 동일 가중치로 평균한다. '
        '모든 점수는 0–1 척도이며 성격·운명·문헌 정답 확률이 아니다.',
        LIMITS[0], LIMITS[6], LIMITS[7],
    ], [(['측정 항목', '3시드 평균 / 실측'], [
        ['선택 macro-F1 (평균 ± 시드 SD)', f'{number(f["selection"]["macro_f1"]["mean"])} ± {number(f["selection"]["macro_f1"]["std"])}'],
        ['역사적 시험 macro-F1 (평균 ± 시드 SD)', f'{number(f["test"]["macro_f1"]["mean"])} ± {number(f["test"]["macro_f1"]["std"])}'],
        ['선택 평균 정확도 / 시험 평균 정확도', f'{number(f["selection"]["accuracy"]["mean"])} / {number(f["test"]["accuracy"]["mean"])}'],
        ['CNN 평균 / p95 지연 (ms)', f'{number(f["latency"]["mean_ms"], 2)} / {number(f["latency"]["p95_ms"], 2)}'],
        ['첫 시드 시험 macro-F1 / 평균 정확도', f'{number(f["runs"][first]["calibrated"]["primary_score"])} / {number(f["runs"][first]["calibrated"]["mean_per_head"]["accuracy"])}'],
    ])])
    section('2. 7개 후보의 성능·운영 비용', [
        'R18 = ResNet-18, E0 = EfficientNet-B0, CN = ConvNeXt-Tiny. 모든 성능 열은 보정 후 3시드 평균이다. 지연은 대표 첫 시드 구성의 실측이다.',
        f'최고 단일 계열 {label(best_single)} 대비 최고 앙상블 {label(ens)}의 선택 macro-F1 차이는 {gain:+.4f}이다. '
        f'평균 지연은 각각 {number(families[best_single]["latency"]["mean_ms"], 2)} / {number(families[ens]["latency"]["mean_ms"], 2)} ms이다. '
        '앙상블 이득을 전제하지 않고 이 관측 차이와 지연을 함께 판단한다. 시험 결과로 선정 후보를 바꾸지 않는다.',
        'ImageNet 사전학습은 일반 시각 특징을 활용하는 초기화이다. 이 실험에는 무작위 초기화 대조군이 없으므로 사전학습 자체의 개선 폭을 수치로 주장하지 않는다.',
        LIMITS[5],
    ], [(['후보', '선택 F1', '시험 F1', '선택 정확도', '시험 정확도', '평균 ms', 'p95 ms'],
          [[label(n, True) + (' *' if n == chosen else ''), *[number(row[s][m]['mean']) for s, m in [('selection', 'macro_f1'), ('test', 'macro_f1'), ('selection', 'accuracy'), ('test', 'accuracy')]], number(row['latency']['mean_ms'], 2), number(row['latency']['p95_ms'], 2)] for n, row in families.items()]),
         (['후보', '매개변수 (백만)', '파일 합계 MiB', 'GPU 최대 할당 GiB'],
          [[label(n, True), number(row['latency']['parameters'] / 1e6, 2), number(row['latency']['checkpoint_bytes'] / 2**20, 1), number(row['latency']['gpu_peak_allocated_gib'], 2)] for n, row in families.items()])])
    costs = []
    for n in families:
        vals = family_cost(data, n)
        runs = [r for r in data['training_runs'] if r['arch'] in families[n]['architectures']]
        costs.append([label(n, True), number(min(vals), 2), number(mean(vals), 2), number(max(vals), 2),
                      f'{sum(len(r["history"]) == data["protocol"]["epochs_max"] for r in runs)}/{len(runs)}',
                      f'{sum(r["selected_epoch"] == data["protocol"]["epochs_max"] for r in runs)}/{len(runs)}'])
    section('3. 실험 절차와 실제 학습 비용', [
        '참여자 분할: 학습 270명 / 선택 34명 / 보정 34명 / 역사적 시험 62명. '
        '3계열 × 5부위 × 3시드 = 45회 학습이다. 7후보는 이 모델들의 확률 평균 조합으로 구성하며 별도 7배 재학습하지 않는다.',
        '순서: 학습·선택집합으로 체크포인트 고정 → 보정 34명으로 각 후보·시드·헤드의 양의 온도 적합 → '
        '보정된 선택 지표로 계열 선택·동결 → 시험 평가. 양의 온도는 argmax를 바꾸지 않는다.',
        f'시드: {", ".join(map(str, data["protocol"]["seeds"]))}. '
        f'최대 {data["protocol"]["epochs_max"]} epoch, 최소 {data["protocol"]["epochs_min"]} epoch, '
        f'조기 종료 patience {data["protocol"]["early_stop_patience"]}. '
        f'{data["protocol"]["optimizer"]}; batch {data["protocol"]["batch_size"]}. 분류 손실만 사용하며 회귀 보조 손실은 없다.',
        '학습 비용(분)은 시드마다 구성 계열의 5부위 전체 history.train_seconds + validation_seconds를 합산했다. '
        '캐시 준비 등이 섞이는 elapsed_seconds는 비용 비교에서 제외하고 XLSX에 별도 보존한다. '
        '3시드 실제 최솟값·평균·최댓값이며 새 환경 전체 파이프라인 시간으로 해석하지 않는다.',
        '실행12는 len(history)==12, 최적12는 selected_epoch==12이다. at_epoch_budget 원본 필드는 후자 의미뿐이다. '
        '앙상블 비용은 이미 학습한 구성 모델 비용의 합이며 7행 합계는 중복 집계이다.',
        f'지연 장비: {data["latency_method"]["gpu"]}, {data["latency_method"]["precision"]}, {data["latency_method"]["layout"]}. '
        f'반복당 예열 {data["latency_method"]["warmup_per_repeat"]}회 후 {data["latency_method"]["trials_per_repeat"]}회 × {data["latency_method"]["repeats"]}반복. '
        '미리 resize한 7개 uint8 crop부터 전송·정규화·순전파·양측/모델 평균·보정·CPU 확률 반환까지이다.',
    ], [(['후보 (5부위/계열)', '최소 분', '평균 분', '최대 분', '실행12', '최적12'], costs)])
    headrows = []
    absent = []
    for key in features:
        firstrow = f['runs'][first]['calibrated']['features'][key]
        values = [r['calibrated']['features'][key] for r in f['runs'].values()]
        support = data['class_support']['test'][key]
        zero = [c for c, n in zip(support['classes'], support['counts']) if n == 0]
        if zero:
            absent.append(f'{HEADS.get(key, key)}: {", ".join(zero)}')
        headrows.append([HEADS.get(key, key), number(firstrow['macro_f1']), number(firstrow['accuracy']),
                         number(mean([r['macro_f1'] for r in values])), number(mean([r['accuracy'] for r in values])),
                         '/'.join(map(str, support['counts']))])
    section('4. 선정 계열의 15개 외형 항목', [
        f'역사적 시험 결과. 대표는 첫 시드 {first}, 평균은 동일 62명에 대한 3시드 평균이다. '
        '지지 수는 클래스 순서별 참조 레이블 수이며 시드마다 같으므로 3배로 합산하지 않는다. '
        '클래스명·모든 분할의 지지 수·누락 레이블은 XLSX class_support에 수록했다.',
        'macro-F1은 정답 또는 예측에 출현한 클래스의 합집합을 사용한다. '
        '정답 지지 수 0인 클래스의 재현율은 평가 불가이다. 예측만 나타난 클래스는 F1 계산에 포함될 수 있지만 해당 클래스 인식 능력을 검증한 뜻이 아니다.',
        '시험에서 지지 수 0인 클래스: ' + ('; '.join(absent) if absent else '없음') + '.',
        LIMITS[3],
    ], [(['항목', '대표 F1', '대표 정확도', '평균 F1', '평균 정확도', '클래스별 n'], headrows)])
    boots = []
    for n, b in data['chosen_vs_other_first_seed_paired_bootstrap'].items():
        ci = b.get('ci95')
        boots.append([label(n, True), number(b.get('mean_macro_f1_difference'), 4),
                      '—' if not ci else f'{number(ci[0], 4)} ~ {number(ci[1], 4)}',
                      str(b.get('valid_bootstrap_samples', '—'))])
    section('5. 보정·불확실성·보류 지표', [
        '온도는 체크포인트 고정 후 보정 34명에서 후보·시드·항목별로 적합했다. '
        '첫 표는 선정 계열의 역사적 시험 결과를 보정 전후 3시드 평균으로 비교한다. '
        'NLL·Brier·ECE는 낮을수록 좋으며, 보정 후 증가하면 해당 시험 지표는 악화된 것이다. '
        'XLSX temperatures는 development.json의 실제 온도·적합 상태·NLL 전후·경계 도달 여부를 보존한다.',
        '아래 차이는 선정 후보 - 비교 후보의 첫 시드 raw 확률 argmax 기반 시험 macro-F1이다. '
        '양의 온도는 argmax를 유지하므로 이 F1 차이는 보정 후에도 같다. '
        '동일 참여자를 묶은 percentile bootstrap이며 3시드 통합 신뢰구간이나 통계적 동등성 입증이 아니다.',
        '보류 관련 수치는 평가기의 신뢰도 0.60 기준이다. coverage는 항목별 통과율 평균, 통과 정확도는 유효 항목 평균이다. '
        '서비스의 품질 보류·VGGT 경계점수·문헌 종합 일치점수와 구별한다.',
        LIMITS[2],
    ], [(['시험 지표 (3시드 평균)', '보정 전 (raw)', '온도 보정 후'],
          [[k, *[number(mean([run[mode]['mean_per_head'][k] for run in f['runs'].values()]))
                 for mode in ('raw', 'calibrated')]]
           for k in ['nll', 'brier', 'ece', 'threshold_coverage', 'threshold_accuracy']]),
         (['비교 후보', '선정-비교 F1', '95% 구간', '유효 재표집'], boots)])
    section('6. 설계서 교체 문안 제안', [
        PROPOSAL_SCOPE,
        *[title + '\n' + body for title, body in replacements(data)],
        proposal_reason(data),
    ])
    section('7. 적용 범위·재현 자료·공식 출처', [
        LIMITS[1], LIMITS[0], LIMITS[4], LIMITS[7],
        'XLSX에는 7후보 전체 집계, 선택/시험의 raw·보정 후 시드별 지표, 항목별 지표, '
        '클래스 지지 수, 45회 학습과 모든 epoch, 보정 온도, 첫 시드 paired bootstrap, 지연 원자료와 방법을 수록했다. '
        '입력 사진·crop·참여자 ID는 보고서에 저장하지 않는다.',
        '로컬 근거: scripts/model_selection_evaluate.py; scripts/model_selection_metrics.py; '
        'scripts/model_selection_train.py; models/selection-20261014/protocol.json; '
        'tmp/pilot-20261014/design.txt §5.1/5.2/5.3; outputs/model-selection-20261014/design-audit.txt.',
        *[f'{key}: {value}' for key, value in provenance.items()],
        '공식 자료는 모델·사전학습·전이학습 방법의 근거이다. FaceTeller 수치는 위 로컬 실험 결과에서만 산출했다.',
        *[name + '\n' + url for name, url in SOURCES],
    ])
    return sections


def write_xlsx(path, data, dev, features, provenance, synthetic):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name, rows):
        ws = wb.create_sheet(name)
        rows = list(rows)
        keys = list(dict.fromkeys(k for row in rows for k in row))
        ws.append(keys)
        for row in rows:
            values = []
            for k in keys:
                value = row.get(k)
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, ensure_ascii=False, allow_nan=False)
                if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
                    value = "'" + value
                values.append(value)
            ws.append(values)
        ws.freeze_panes = 'B2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(name='맑은 고딕', bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='203B59')
        for col in ws.columns:
            width = min(65, max(14, max(len(str(c.value or '')) for c in col[:80]) + 2))
            ws.column_dimensions[get_column_letter(col[0].column)].width = width
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name='맑은 고딕', size=10)
                c.alignment = Alignment(vertical='top', wrap_text=True)
                if isinstance(c.value, float):
                    c.number_format = '0.0000'
        return ws

    chosen = data['choice']['selected_family']
    sheet('readme', [{'항목': k, '내용': v} for k, v in [
        ('문서', '합성 검증 데이터 · 실제 성능 아님' if synthetic else '완료된 실측 결과 보고서'),
        ('선정', label(chosen)), ('대표 시드', data['choice']['deployment_seed']),
        ('정확도 해석', '항목별 정확도 평균; 0–1; 관상·운명 예측 정확도 아님'),
        ('훈련 비용', 'history.train_seconds+validation_seconds; elapsed_seconds 별도 원본'),
        ('epoch 구분', 'ran_budget=len(history)==12; at_epoch_budget=best epoch==12'),
        ('클래스', '지지 수 0인 클래스 재현율 평가 불가; macroF1 truth/pred 합집합'),
        *[(f'제한 {i+1}', v) for i, v in enumerate(LIMITS)], *provenance.items()]])
    sheet('aggregate', [dict(family=label(n), family_id=n, selected=n == chosen, **flat({k: v for k, v in r.items() if k != 'runs'})) for n, r in data['families'].items()])
    agg, perhead, temperatures = [], [], []
    for n, family in data['families'].items():
        for split, runs in [('selection', dev['families'][n]['runs']), ('test', family['runs'])]:
            for seed, run in runs.items():
                for mode in ('raw', 'calibrated'):
                    meta = dict(family=label(n), split=split, seed=seed, mode=mode)
                    agg.append({**meta, **flat({k: v for k, v in run[mode].items() if k != 'features'})})
                    for key, metrics in run[mode]['features'].items():
                        perhead.append({**meta, 'feature': key, '항목': HEADS.get(key, key), **flat(metrics)})
                if split == 'selection':
                    for key, temp in run['temperatures'].items():
                        temperatures.append(dict(family=label(n), seed=seed, feature=key, **flat(temp)))
    sheet('aggregate_runs', agg)
    sheet('perfeature_runs', perhead)
    summary = []
    for split in ('selection', 'test'):
        for n, family in data['families'].items():
            runs = dev['families'][n]['runs'] if split == 'selection' else family['runs']
            for key in features:
                row = dict(family=label(n), selected=n == chosen, split=split, feature=key, 항목=HEADS.get(key, key))
                for metric in ('macro_f1', 'accuracy', 'balanced_accuracy', 'nll', 'brier', 'ece'):
                    vals = [r['calibrated']['features'][key][metric] for r in runs.values()]
                    valid = [v for v in vals if v is not None]
                    row.update({metric + '_mean': mean(vals), metric + '_seed_sd': statistics.stdev(valid) if len(valid) > 1 else None,
                                metric + '_valid_seeds': len(valid), metric + '_first_seed': runs[str(data['choice']['deployment_seed'])]['calibrated']['features'][key][metric]})
                row.update(flat(data['class_support'][split][key], 'reference'))
                summary.append(row)
    sheet('perfeature_summary', summary)
    sheet('class_support', [dict(split=split, feature=key, 항목=HEADS.get(key, key), class_index=i,
                                 class_name=c, support=n, recall_assessable=n > 0, missing_labels=row['missing'])
                            for split, rows in data['class_support'].items() for key, row in rows.items()
                            for i, (c, n) in enumerate(zip(row['classes'], row['counts']))])
    sheet('training45', [dict(**flat({k: v for k, v in r.items() if k != 'history'}),
                            core_seconds=core_seconds(r), epochs_run=len(r['history']),
                            ran_budget=len(r['history']) == data['protocol']['epochs_max'],
                            best_epoch_at_budget=r['selected_epoch'] == data['protocol']['epochs_max'])
                       for r in data['training_runs']])
    sheet('epochs', [dict(arch=r['arch'], part=r['part'], seed=r['seed'], **flat(e))
                     for r in data['training_runs'] for e in r['history']])
    sheet('family_training', [dict(family=label(n), seed=s, core_minutes=v, scope='구성 계열 × 5부위; 캐시 준비 제외')
                              for n in data['families'] for s, v in zip(data['protocol']['seeds'], family_cost(data, n))])
    sheet('family_training_summary', [dict(family=label(n), seed_count=3,
                                         min_core_minutes=min(family_cost(data, n)),
                                         mean_core_minutes=mean(family_cost(data, n)),
                                         max_core_minutes=max(family_cost(data, n)),
                                         total_core_minutes=sum(family_cost(data, n)))
                                      for n in data['families']])
    sheet('temperatures', temperatures)
    sheet('paired_bootstrap', [dict(comparator=label(n), chosen=label(chosen), probabilities='first-seed raw; positive T preserves argmax', **flat(b))
                                for n, b in data['chosen_vs_other_first_seed_paired_bootstrap'].items()])
    sheet('latency_samples', [dict(family=label(n), trial=i + 1, milliseconds=v)
                              for n, row in data['families'].items() for i, v in enumerate(row['latency']['milliseconds'])])
    protocol = {k: v for k, v in data['protocol'].items() if k != 'groups'}
    sheet('method_sources', [{'key': k, 'value': v} for k, v in flat(dict(protocol=protocol, choice=data['choice'], latency_method=data['latency_method'], limitations=data['limitations'])).items()]
          + [{'key': n, 'value': u} for n, u in SOURCES])
    sheet('design_proposal', [{'절': '제안 범위', '교체 제안': PROPOSAL_SCOPE}]
          + [{'절': title, '교체 제안': body} for title, body in replacements(data)]
          + [{'절': '선정 이유 (요구사항 밖)', '교체 제안': proposal_reason(data)}])
    wb.save(path)


def font_paths():
    for base in (Path('C:/Windows/Fonts'), Path('/mnt/c/Windows/Fonts')):
        if (base / 'malgun.ttf').exists():
            return base / 'malgun.ttf', base / 'malgunbd.ttf'
    raise FileNotFoundError('Malgun font required: C:/Windows/Fonts/malgun.ttf')


def write_pdf(path, sections, synthetic):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    regular, bold = font_paths()
    pdfmetrics.registerFont(TTFont('Malgun', str(regular)))
    pdfmetrics.registerFont(TTFont('MalgunBold', str(bold)))
    body = ParagraphStyle('body', fontName='Malgun', fontSize=9, leading=13.6, spaceAfter=8, wordWrap='CJK', alignment=TA_LEFT)
    title = ParagraphStyle('title', parent=body, fontName='MalgunBold', fontSize=17, leading=23, spaceAfter=14)
    small = ParagraphStyle('small', parent=body, fontSize=7.8, leading=10.8, spaceAfter=0)
    header = ParagraphStyle('header', parent=small, fontName='MalgunBold', textColor=colors.white)
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=17*mm, leftMargin=17*mm,
                            topMargin=19*mm, bottomMargin=18*mm, title=TITLE,
                            author='FaceTeller', pageCompression=1)
    def para(text, style=body):
        return Paragraph(escape(str(text)).replace('\n', '<br/>'), style)
    story = []
    for i, section in enumerate(sections):
        if i:
            story.append(PageBreak())
        story.append(para(section['title'], title))
        if synthetic:
            story.append(para('합성 검증 데이터 · 실제 성능 아님'))
        # Place tables after introductory paragraphs, before detailed caveats.
        intro = 1 if section['tables'] else len(section['paragraphs'])
        story.extend(para(p) for p in section['paragraphs'][:intro])
        for headings, rows in section['tables']:
            count = len(headings)
            first_width = doc.width * (0.22 if count > 4 else 0.40)
            widths = [first_width] + [(doc.width-first_width)/(count-1)]*(count-1)
            table = Table([[para(h, header) for h in headings]] + [[para(c, small) for c in row] for row in rows], colWidths=widths, repeatRows=1, hAlign='LEFT')
            table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#203B59')),
                ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F0F4F7')]),
                ('VALIGN', (0,0), (-1,-1), 'TOP'), ('LEFTPADDING', (0,0), (-1,-1), 5),
                ('RIGHTPADDING', (0,0), (-1,-1), 5), ('TOPPADDING', (0,0), (-1,-1), 5),
                ('BOTTOMPADDING', (0,0), (-1,-1), 5),
                ('LINEBELOW', (0,-1), (-1,-1), .5, colors.HexColor('#CBD5E1'))]))
            story.extend([table, Spacer(1, 10)])
        story.extend(para(p) for p in section['paragraphs'][intro:])
    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont('Malgun', 8)
        canvas.setFillColor(colors.HexColor('#536476'))
        canvas.drawString(17*mm, 10*mm, 'FaceTeller · CNN 모델 선정' + (' · 합성 검증용' if synthetic else ' · 실측 결과'))
        canvas.drawRightString(A4[0]-17*mm, 10*mm, str(document.page))
        canvas.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    if not 5 <= doc.page <= 8:
        raise ValueError(f'PDF layout needs review: {doc.page} pages (expected 5–8)')
    return doc.page


def write_txt(path, sections, synthetic):
    lines = [TITLE, '합성 검증 데이터 · 실제 성능 아님' if synthetic else '완료된 실측 결과', '']
    for s in sections:
        lines.extend([s['title'], '', *s['paragraphs'], ''])
        for headers, rows in s['tables']:
            lines.append(' | '.join(map(str, headers)))
            lines.extend(' | '.join(map(str, row)) for row in rows)
            lines.append('')
    path.write_text('\n'.join(lines) + '\n', encoding='utf-8')


def default_share():
    prefix = Path('C:/Users/user/Downloads') if os.name == 'nt' else Path('/mnt/c/Users/user/Downloads')
    return prefix / 'FaceTeller_파일럿_공유자료'


def main(argv=None):
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, default=DEFAULT_INPUT)
    parser.add_argument('--development', type=Path)
    parser.add_argument('--output-dir', type=Path)
    parser.add_argument('--share-dir', type=Path, default=default_share())
    parser.add_argument('--no-copy', action='store_true')
    parser.add_argument('--synthetic', action='store_true')
    args = parser.parse_args(argv)
    dev_path = args.development or args.input.with_name('development.json')
    out = args.output_dir or args.input.parent
    if args.synthetic:
        temp = Path(tempfile.gettempdir()).resolve()
        if not args.no_copy or args.output_dir is None or not out.resolve().is_relative_to(temp):
            parser.error('Synthetic checks require --no-copy and --output-dir inside OS TEMP')
    data, dev = read(args.input), read(dev_path)
    if bool(data.get('synthetic', False)) != args.synthetic:
        parser.error('Synthetic fixture marker and --synthetic must agree')
    features = validate(data, dev, dev_path)
    provenance = {'results SHA-256': sha(args.input), 'development SHA-256': sha(dev_path),
                  'protocol SHA-256': data['choice']['protocol_sha256'],
                  '배포 대표 시드': str(data['choice']['deployment_seed'])}
    sections = report_sections(data, features, provenance)
    out.mkdir(parents=True, exist_ok=True)
    # Build all artifacts in staging before replacing any finished reports.
    with tempfile.TemporaryDirectory(prefix='faceteller-report-') as stage:
        stage = Path(stage)
        write_xlsx(stage / f'{TITLE}.xlsx', data, dev, features, provenance, args.synthetic)
        pages = write_pdf(stage / f'{TITLE}.pdf', sections, args.synthetic)
        write_txt(stage / f'{TITLE}.txt', sections, args.synthetic)
        destinations = [out]
        if not args.no_copy and args.share_dir.resolve() != out.resolve():
            destinations.append(args.share_dir)
        for dest in destinations:
            dest.mkdir(parents=True, exist_ok=True)
            for suffix in ('.xlsx', '.pdf', '.txt'):
                source = stage / (TITLE + suffix)
                pending = dest / (TITLE + suffix + '.pending')
                shutil.copyfile(source, pending)
                pending.replace(dest / source.name)
    print(json.dumps({'status': 'ok', 'synthetic': args.synthetic, 'pdf_pages': pages,
                      'selected_family': data['choice']['selected_family'],
                      'outputs': [str(d / (TITLE + s)) for d in destinations for s in ('.xlsx', '.pdf', '.txt')]}, ensure_ascii=False))


if __name__ == '__main__':
    main()
