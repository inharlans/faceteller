from pathlib import Path
import json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'outputs/requirements-v1'
x=json.loads((OUT/'설계서_구조화원본.json').read_text(encoding='utf8'))
doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(1.9);sec.right_margin=Cm(1.9)
sec.header_distance=Cm(.8);sec.footer_distance=Cm(.8);sec.different_first_page_header_footer=True
NAVY='19394E';TEAL='087D82';GRAY='526371';LIGHT='F0F5F7'
def style_font(style,size,bold=False,color=None):
 style.font.name='맑은 고딕';style.font.size=Pt(size);style.font.bold=bold
 style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
 if color:style.font.color.rgb=RGBColor.from_string(color)
style_font(doc.styles['Normal'],10.3)
f=doc.styles['Normal'].paragraph_format;f.line_spacing=Pt(14.8);f.space_after=Pt(6);f.widow_control=True
for st,size,col in [('Title',27,NAVY),('Subtitle',17,TEAL),('Heading 1',18,NAVY),('Heading 2',11,TEAL),('Heading 3',10.5,NAVY)]:
 style_font(doc.styles[st],size,True,col);doc.styles[st].paragraph_format.keep_with_next=True
 doc.styles[st].paragraph_format.space_before=Pt(10 if st!='Heading 1' else 0)
 doc.styles[st].paragraph_format.space_after=Pt(6)
for name,size in [('Caption',8.5),('Header',8),('Footer',8),('TOC 1',9.3)]:
 if name not in doc.styles:doc.styles.add_style(name,1)
 style_font(doc.styles[name],size,color=GRAY)
 doc.styles[name].paragraph_format.line_spacing=Pt(12.3)
 doc.styles[name].paragraph_format.space_after=Pt(2)
# Single-level automatic TOC uses real Word pagination.
def field(p,code):
 r=p.add_run();b=OxmlElement('w:fldChar');b.set(qn('w:fldCharType'),'begin');r._r.append(b)
 t=OxmlElement('w:instrText');t.set(qn('xml:space'),'preserve');t.text=code;r._r.append(t)
 sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r._r.append(sep)
 txt=OxmlElement('w:t');txt.text='';r._r.append(txt)
 end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r._r.append(end)
header=sec.header.paragraphs[0];header.text='SRS-FACE-001  /  얼굴 형태·전통 문헌 서비스';header.style='Header'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('v0.9 · 검토용 초안     |     ');field(footer,'PAGE');footer.add_run(' / ');field(footer,'NUMPAGES')

def p(text,style=None):return doc.add_paragraph(text,style)
def shade(cell,color):
 tcPr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),color);tcPr.append(shd)
def tab(headers,rows):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 n=len(headers)
 fractions={2:[.29,.71],3:[.24,.35,.41],4:[.18,.25,.27,.30],5:[.18,.31,.15,.22,.14]}.get(n,[1/n]*n)
 for c,w in zip(t.columns,fractions):c.width=Cm(17.2*w)
 for cell,text in zip(t.rows[0].cells,headers):cell.text=str(text);shade(cell,NAVY)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for cell,text in zip(cells,row):
   cell.text=str(text)
   if i%2==0:shade(cell,LIGHT)
 for ri,row in enumerate(t.rows):
  cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
  for ci,cell in enumerate(row.cells):
   cell.width=Cm(17.2*fractions[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
   pr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
   for name,v in [('top',65),('bottom',65),('left',90),('right',90)]:
    e=OxmlElement('w:'+name);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');marg.append(e)
   pr.append(marg)
   for par in cell.paragraphs:
    par.paragraph_format.space_after=Pt(1);par.paragraph_format.space_before=Pt(0);par.paragraph_format.line_spacing=Pt(12.2)
    for r in par.runs:
     r.font.name='맑은 고딕';r.font.size=Pt(8.8);r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
     if ri==0:r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255)
 doc.add_paragraph().paragraph_format.space_after=Pt(1)
 return t

# Cover is informative, with no simulated signatures or institutional affiliation.
p('PROJECT REQUIREMENTS  /  2026.09', 'Subtitle')
p('AI 기반 얼굴 형태 분석 및\n전통 관상 문헌 서비스','Title')
p('상세 요구사항 설계서','Subtitle')
p('프로토타입 검증 이후, 데이터와 제품을 순서대로 다시 설계하기 위한 개발 기준선 초안입니다.')
p('')
tab(['문서 정보','내용'],[['문서 ID','SRS-FACE-001'],['버전 / 기준일','v0.9 검토용 초안 / 2026-09-30'],['범위','제품 기능 · 데이터 · 문헌 · 레이블 · 모델 · API · 검증 · 운영'],['상세 구성','104개 요구사항 / 38개 특징 후보 / 12개 결정 항목 / 104개 수용시험 계획'],['작성 상태','팀 검토 전. 실제 승인자·서명·기관 이용 허가를 대신 작성하지 않음']])
p('이번 설계의 중심','Heading 2')
for s in ['관찰 가능한 형태와 역사 문헌의 해석을 출처·조건으로 연결한다.','CNN 분류·기하 측정·문헌 규칙의 역할과 수용 기준을 명확히 한다.','사람 수와 반복 사진 수를 구별하고 독립 검수·외부 촬영 검증을 단계에 포함한다.','연구 시제품, 비공개 베타, 상업 운영을 별도 릴리스로 관리한다.']:p('• '+s)
p('원본 제공 얼굴 사진은 이 문서에 포함하지 않았습니다. 기존 실험의 수치는 참고 기록이며 신규 서비스의 품질 보증이 아닙니다.','Caption')
doc.add_page_break();p('목차','Title');p('각 절은 독립적으로 검토할 수 있습니다. 요구 ID와 동봉 CSV를 사용하면 개발 작업·수용시험을 추적할 수 있습니다.')
field(p(''),'TOC \\o "1-1" \\h \\z \\u')
for i,pg in enumerate(x['pages'],1):
 doc.add_page_break();p(f"{i:02d}. {pg['title']}",'Heading 1');p(pg['lead'])
 for b in pg['blocks']:
  k=b['kind']
  if k=='p':p(b['text'])
  elif k=='note':
   a=p('검토 메모  '+b['text']);a.paragraph_format.left_indent=Cm(.18)
   for r in a.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string(GRAY)
  elif k=='bullets':
   for v in b['items']:p('• '+v)
  elif k=='table':tab(b['head'],b['rows'])
  elif k=='requirement':
   p(f"{b['id']}  {b['title']}  [{b['priority']}]",'Heading 2')
   a=p(f"담당: {b['actor']}  |  연계/게이트: {b['depends']}",'Caption');a.paragraph_format.keep_with_next=True
   p(b['rule']);a=p('수용 기준  '+b['acceptance'])
   for r in a.runs:r.font.size=Pt(9.8)
props=doc.core_properties;props.title=x['title']+' — '+x['subtitle'];props.subject='신규 개발을 위한 상세 요구사항 기준선 초안';props.author='프로젝트 검토용';props.version='0.9'
# Word recalculates fields during controlled export.
settings=doc.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
target=OUT/'상세_요구사항_설계서_v0.9.docx';doc.save(target);print(target)
